# =============================================================================
# crisis_gui.py  —  CRISIS Landslide Model GUI
# =============================================================================
# Data-driven: zero hardcoded parameter names.
# All widgets are generated from crisis_params.PARAMS at startup.
# Add a param to crisis_params.py → it appears in the GUI automatically.
# =============================================================================

import importlib
import threading
import json
import traceback
import os
import sys
import webbrowser
from pathlib import Path

# Detect whether this module is being imported by a spawned worker process.
# On Windows, ProcessPoolExecutor re-imports __main__ (this file) in every
# worker.  Skipping the GUI-only imports keeps worker startup fast.
import multiprocessing as _mp
_IN_WORKER = _mp.current_process().name != 'MainProcess'

_HAS_PIL = False
_HAS_MPL = False
np = None  # forward declaration; set below if matplotlib is available

if not _IN_WORKER:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    try:
        from PIL import Image, ImageTk
        _HAS_PIL = True
    except ImportError:
        pass

    try:
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_tkagg import (
            FigureCanvasTkAgg, NavigationToolbar2Tk as _NavTB2)
        import matplotlib.cm as _mcm
        import matplotlib as _mpl
        import numpy as np
        _HAS_MPL = True
    except ImportError:
        pass

# =============================================================================
# Design tokens
# =============================================================================

C = {
    'bg':           '#f0f4f8',   # window background
    'surface':      '#ffffff',   # card / tab interior
    'surface_alt':  '#f8fafc',   # alternate row tint
    'border':       '#dde3ea',   # subtle borders
    'accent':       '#1a56db',   # primary blue
    'accent_dark':  '#1648c0',   # hover blue
    'accent_light': '#e8eefb',   # tinted blue background
    'success':      '#057a55',   # run-button green
    'success_dark': '#046040',   # hover green
    'danger':       '#c81e1e',   # error red
    'text':         '#111928',   # primary text
    'text_muted':   '#6b7280',   # secondary / label text
    'text_on_dark': '#ffffff',   # text on coloured bg
    'header_bg':    '#1e3a5f',   # dark-navy header
    'header_sub':   '#93c5fd',   # light-blue subtitle
    'log_bg':       '#0d1117',   # log area background
    'log_fg':       '#c9d1d9',   # log area text
    'log_accent':   '#58a6ff',   # log highlights
    'tooltip_bg':   '#1f2937',   # tooltip background
    'tooltip_fg':   '#f9fafb',   # tooltip text
}

F = {
    'ui':        ('Segoe UI', 10),
    'ui_bold':   ('Segoe UI', 10, 'bold'),
    'ui_sm':     ('Segoe UI', 9),
    'label':     ('Segoe UI', 10),
    'section':   ('Segoe UI', 9,  'bold'),
    'title':     ('Segoe UI', 26, 'bold'),
    'subtitle':  ('Segoe UI', 13),
    'tab':       ('Segoe UI', 10, 'bold'),
    'btn':       ('Segoe UI', 10, 'bold'),
    'btn_run':   ('Segoe UI', 11, 'bold'),
    'log':       ('Consolas', 10),
    'log_hdr':   ('Consolas',  9),
}


# =============================================================================
# Logo helpers
# =============================================================================

_LOGO_PATH  = Path(__file__).parent / 'assets' / 'crisis_logo.jpg'
_LOGO_PHOTO = None
_GRAD_PHOTO = None


def _load_logo(height_px: int = 90):
    """Return a PhotoImage of the custom logo scaled to height_px."""
    global _LOGO_PHOTO
    if not _HAS_PIL or not _LOGO_PATH.exists():
        return None
    try:
        img = Image.open(_LOGO_PATH).convert('RGB')
        ratio = height_px / img.height
        new_w = max(1, int(img.width * ratio))
        img = img.resize((new_w, height_px), Image.LANCZOS)
        _LOGO_PHOTO = ImageTk.PhotoImage(img)
        return _LOGO_PHOTO
    except Exception:
        return None


def _make_gradient(width: int, height: int):
    """Return a PhotoImage: navy (#1e3a5f) on the left → white on the right."""
    global _GRAD_PHOTO
    if not _HAS_PIL or width < 1 or height < 1:
        return None
    try:
        import numpy as np
        t   = np.linspace(0, 1, width, dtype=np.float32)
        arr = np.zeros((height, width, 3), dtype=np.uint8)
        arr[:, :, 0] = (30  + (255 - 30)  * t).astype(np.uint8)
        arr[:, :, 1] = (58  + (255 - 58)  * t).astype(np.uint8)
        arr[:, :, 2] = (95  + (255 - 95)  * t).astype(np.uint8)
        _GRAD_PHOTO = ImageTk.PhotoImage(Image.fromarray(arr, 'RGB'))
        return _GRAD_PHOTO
    except Exception:
        return None


# =============================================================================
# Apply ttk styles
# =============================================================================

def _apply_styles(root):
    s = ttk.Style(root)
    s.theme_use('clam')   # clam exposes the most colour knobs

    # --- global defaults ---
    s.configure('.',
                background=C['bg'],
                foreground=C['text'],
                font=F['ui'],
                borderwidth=0,
                relief='flat')

    # --- frames ---
    s.configure('TFrame',   background=C['bg'])
    s.configure('Card.TFrame', background=C['surface'],
                relief='flat', borderwidth=1)

    # --- labels ---
    s.configure('TLabel',   background=C['bg'], foreground=C['text'], font=F['label'])
    s.configure('Muted.TLabel', foreground=C['text_muted'], font=F['ui_sm'])
    s.configure('Section.TLabel',
                background=C['surface'],
                foreground=C['accent'],
                font=F['section'])

    # --- entries ---
    s.configure('TEntry',
                fieldbackground=C['surface'],
                foreground=C['text'],
                insertcolor=C['text'],
                bordercolor=C['border'],
                lightcolor=C['border'],
                darkcolor=C['border'],
                padding=(6, 4),
                font=F['ui'])
    s.map('TEntry',
          bordercolor=[('focus', C['accent'])],
          lightcolor=[('focus', C['accent'])],
          darkcolor=[('focus', C['accent'])],
          fieldbackground=[('readonly', C['surface_alt']), ('disabled', C['border'])])

    # --- combobox ---
    s.configure('TCombobox',
                fieldbackground=C['surface'],
                background=C['surface'],
                foreground=C['text'],
                arrowcolor=C['text_muted'],
                bordercolor=C['border'],
                lightcolor=C['border'],
                darkcolor=C['border'],
                padding=(6, 4),
                font=F['ui'])
    s.map('TCombobox',
          fieldbackground=[('readonly', C['surface'])],
          bordercolor=[('focus', C['accent'])])

    # --- standard buttons (Browse / secondary) ---
    s.configure('TButton',
                background=C['surface'],
                foreground=C['text'],
                bordercolor=C['border'],
                lightcolor=C['border'],
                darkcolor=C['border'],
                focuscolor=C['accent_light'],
                padding=(10, 5),
                font=F['btn'])
    s.map('TButton',
          background=[('active', C['accent_light']), ('pressed', C['border'])],
          foreground=[('active', C['accent'])])

    # --- notebook ---
    s.configure('TNotebook',
                background=C['bg'],
                bordercolor=C['border'],
                tabmargins=(2, 4, 0, 0))
    s.configure('TNotebook.Tab',
                background=C['bg'],
                foreground=C['text_muted'],
                padding=(14, 7),
                font=F['tab'])
    s.map('TNotebook.Tab',
          background=[('selected', C['surface'])],
          foreground=[('selected', C['accent'])],
          expand=[('selected', (1, 1, 1, 0))])

    # --- scrollbar ---
    s.configure('TScrollbar',
                background=C['border'],
                troughcolor=C['bg'],
                arrowcolor=C['text_muted'],
                borderwidth=0,
                arrowsize=10)
    s.map('TScrollbar',
          background=[('active', C['text_muted'])])

    # --- separator ---
    s.configure('TSeparator', background=C['border'])

    root.configure(background=C['bg'])


# =============================================================================
# Tooltip
# =============================================================================

class _Tooltip:
    def __init__(self, widget, text):
        self._w   = widget
        self._txt = text
        self._tip = None
        widget.bind('<Enter>', self._show)
        widget.bind('<Leave>', self._hide)

    def _show(self, _e):
        if not self._txt:
            return
        x = self._w.winfo_rootx() + 28
        y = self._w.winfo_rooty() + 28
        self._tip = tk.Toplevel(self._w)
        self._tip.wm_overrideredirect(True)
        self._tip.wm_geometry(f'+{x}+{y}')
        # Dark, rounded-feeling tooltip
        outer = tk.Frame(self._tip, bg=C['accent'], padx=1, pady=1)
        outer.pack()
        inner = tk.Frame(outer, bg=C['tooltip_bg'], padx=10, pady=6)
        inner.pack()
        tk.Label(inner,
                 text=self._txt,
                 bg=C['tooltip_bg'],
                 fg=C['tooltip_fg'],
                 font=F['ui_sm'],
                 justify=tk.LEFT,
                 wraplength=460).pack()

    def _hide(self, _e):
        if self._tip:
            self._tip.destroy()
            self._tip = None


# =============================================================================
# Scrollable frame helper
# =============================================================================

def _scrollable_frame(parent, bg=None):
    bg = bg or C['surface']
    outer = tk.Frame(parent, bg=bg)
    outer.pack(fill=tk.BOTH, expand=True)

    canvas = tk.Canvas(outer, bg=bg, highlightthickness=0, borderwidth=0)
    sb     = ttk.Scrollbar(outer, orient='vertical', command=canvas.yview)
    inner  = tk.Frame(canvas, bg=bg)

    inner.bind('<Configure>',
               lambda e: canvas.configure(scrollregion=canvas.bbox('all')))
    canvas.create_window((0, 0), window=inner, anchor='nw')
    canvas.configure(yscrollcommand=sb.set)

    canvas.pack(side='left', fill='both', expand=True)
    sb.pack(side='right', fill='y')

    def _wheel(event):
        canvas.yview_scroll(int(-1 * (event.delta / 120)), 'units')
    canvas.bind_all('<MouseWheel>', _wheel)

    return canvas, inner


# =============================================================================
# Flat coloured button (tk.Button styled to look modern)
# =============================================================================

def _flat_btn(parent, text, command, bg, fg, hover_bg,
              font=None, padx=14, pady=6, width=None):
    btn = tk.Button(parent,
                    text=text,
                    command=command,
                    bg=bg, fg=fg,
                    activebackground=hover_bg,
                    activeforeground=fg,
                    font=font or F['btn'],
                    relief='flat',
                    borderwidth=0,
                    cursor='hand2',
                    padx=padx, pady=pady)
    if width:
        btn.config(width=width)
    btn.bind('<Enter>', lambda e: btn.config(bg=hover_bg))
    btn.bind('<Leave>', lambda e: btn.config(bg=bg))
    return btn


# =============================================================================
# Email popup
# =============================================================================

def _show_email_popup(root, widget, addr):
    popup = tk.Toplevel(root)
    popup.wm_overrideredirect(True)
    popup.attributes('-topmost', True)

    popup.update_idletasks()
    pw = popup.winfo_reqwidth()
    x = widget.winfo_rootx() - pw - 20
    y = widget.winfo_rooty() - 4
    popup.wm_geometry(f'+{x}+{y}')

    outer = tk.Frame(popup, bg=C['accent'], padx=1, pady=1)
    outer.pack()
    inner = tk.Frame(outer, bg=C['tooltip_bg'], padx=10, pady=8)
    inner.pack()

    tk.Label(inner, text=addr,
             bg=C['tooltip_bg'], fg=C['tooltip_fg'],
             font=F['ui']).pack(side=tk.LEFT, padx=(0, 10))

    copied = {'done': False}

    def _copy():
        root.clipboard_clear()
        root.clipboard_append(addr)
        root.update()
        copied['done'] = True
        copy_icon.config(fg='#3fb950', text='✓')
        copy_label.config(fg='#3fb950', text='Copied!')
        popup.after(1200, popup.destroy)

    copy_btn_frame = tk.Frame(inner, bg=C['tooltip_bg'], cursor='hand2')
    copy_btn_frame.pack(side=tk.LEFT, padx=(6, 0))

    copy_icon = tk.Label(copy_btn_frame, text='⎘',
                         bg=C['tooltip_bg'], fg='#60a5fa',
                         font=('Segoe UI', 13), cursor='hand2')
    copy_icon.pack(side=tk.LEFT)

    copy_label = tk.Label(copy_btn_frame, text='Copy',
                          bg=C['tooltip_bg'], fg='#60a5fa',
                          font=F['ui_sm'], cursor='hand2')
    copy_label.pack(side=tk.LEFT, padx=(2, 0))

    for w in (copy_btn_frame, copy_icon, copy_label):
        w.bind('<Button-1>', lambda e: _copy())
        w.bind('<Enter>', lambda e: (
            copy_icon.config(fg='#ffffff'),
            copy_label.config(fg='#ffffff')))
        w.bind('<Leave>', lambda e: (
            copy_icon.config(fg='#3fb950' if copied['done'] else '#60a5fa'),
            copy_label.config(fg='#3fb950' if copied['done'] else '#60a5fa')))

    # Close when clicking anywhere outside
    popup.bind('<FocusOut>', lambda e: popup.destroy())
    popup.focus_set()


# =============================================================================
# Main application
# =============================================================================

class CRISISGui:
    def __init__(self, root):
        self.root = root
        self.root.title('CRISIS Model')
        self.root.geometry('1020x820')
        self.root.resizable(True, True)
        self.root.minsize(800, 600)

        self._vars:    dict = {}
        self._frames:  dict = {}
        self._widgets: dict = {}
        self._running = False

        self._params_list, self._categories = _load_params()
        _apply_styles(root)
        self._build_menu()
        self._build_ui()
        self.root.protocol('WM_DELETE_WINDOW', self._on_close)

    # ------------------------------------------------------------------
    # Menu bar
    # ------------------------------------------------------------------

    def _build_menu(self):
        mb = tk.Menu(self.root, bg=C['surface'], fg=C['text'],
                     activebackground=C['accent_light'],
                     activeforeground=C['accent'],
                     relief='flat', borderwidth=0)
        fm = tk.Menu(mb, tearoff=0,
                     bg=C['surface'], fg=C['text'],
                     activebackground=C['accent_light'],
                     activeforeground=C['accent'])
        fm.add_command(label='Save Configuration…',  command=self._save_config)
        fm.add_command(label='Load Configuration…',  command=self._load_config)
        fm.add_separator()
        fm.add_command(label='Reload Parameters',    command=self._reload_params)
        fm.add_separator()
        fm.add_command(label='Exit',                 command=self.root.destroy)
        mb.add_cascade(label='File', menu=fm)
        self.root.config(menu=mb)

    # ------------------------------------------------------------------
    # Full UI layout
    # ------------------------------------------------------------------

    def _build_ui(self):
        self._vars.clear()
        self._frames.clear()
        self._widgets.clear()

        # ── Header banner ───────────────────────────────────────────────
        hdr = tk.Frame(self.root, bg=C['header_bg'])
        hdr.pack(fill=tk.X, side=tk.TOP)

        # Left block: title + subtitle
        title_block = tk.Frame(hdr, bg=C['header_bg'])
        title_block.pack(side=tk.LEFT, padx=(22, 0), pady=26)

        tk.Label(title_block,
                 text='CRISIS  Model',
                 bg=C['header_bg'], fg=C['text_on_dark'],
                 font=F['title']).pack(anchor='w')
        tk.Label(title_block,
                 text='Coupled Regional Rainfall-Induced & Seismic Slope Instability Simulations',
                 bg=C['header_bg'], fg=C['header_sub'],
                 font=F['subtitle']).pack(anchor='w')

        # Right block: logo (white bg) + gradient strip (navy→white) to its left
        logo_lbl = tk.Label(hdr, bg='white', borderwidth=0)
        logo_lbl.pack(side=tk.RIGHT, padx=0, pady=0)

        grad_lbl = tk.Label(hdr, bg=C['header_bg'], borderwidth=0)
        grad_lbl.pack(side=tk.RIGHT, padx=0, pady=0)

        def _resize_header(event=None):
            h = hdr.winfo_height()
            if h < 10:
                return
            # Leave a gap at the bottom so the accent line clears the logo text
            logo_h = max(10, h - 14)
            logo_img = _load_logo(height_px=logo_h)
            if logo_img is not None:
                logo_lbl.config(image=logo_img)
                logo_lbl.image = logo_img
            grad_img = _make_gradient(width=350, height=logo_h)
            if grad_img is not None:
                grad_lbl.config(image=grad_img)
                grad_lbl.image = grad_img

        hdr.bind('<Configure>', _resize_header)

        # ── Thin accent line ────────────────────────────────────────────
        tk.Frame(self.root, bg=C['accent'], height=3).pack(fill=tk.X)

        # ── Bottom toolbar (packed before notebook so BOTTOM gets space first) ──
        toolbar = tk.Frame(self.root, bg=C['bg'], padx=12, pady=4)
        toolbar.pack(fill=tk.X, side=tk.BOTTOM)

        tk.Frame(toolbar, bg=C['border'], height=1).pack(fill=tk.X, pady=(0, 4))

        btn_row = tk.Frame(toolbar, bg=C['bg'])
        btn_row.pack(fill=tk.X)

        self._run_btn = _flat_btn(
            btn_row, '▶   Run Model', self._run_model,
            bg=C['success'], fg=C['text_on_dark'], hover_bg=C['success_dark'],
            font=F['btn_run'], padx=22, pady=9
        )
        self._run_btn.pack(side=tk.LEFT, padx=(0, 10))

        _flat_btn(btn_row, '💾  Save Config', self._save_config,
                  bg=C['surface'], fg=C['text'], hover_bg=C['accent_light']
                  ).pack(side=tk.LEFT, padx=3)

        _flat_btn(btn_row, '📂  Load Config', self._load_config,
                  bg=C['surface'], fg=C['text'], hover_bg=C['accent_light']
                  ).pack(side=tk.LEFT, padx=3)

        _flat_btn(btn_row, '✕  Clear Config', self._clear_config,
                  bg=C['surface'], fg=C['danger'], hover_bg='#fde8e8'
                  ).pack(side=tk.LEFT, padx=3)

        _flat_btn(btn_row, 'ℹ  About', self._show_about,
                  bg=C['surface'], fg=C['text'], hover_bg=C['accent_light']
                  ).pack(side=tk.LEFT, padx=3)

        # ── Contact block at far right ───────────────────────────────────
        contact_block = tk.Frame(btn_row, bg=C['bg'])
        contact_block.pack(side=tk.RIGHT, padx=(0, 4))

        tk.Label(contact_block, text='CONTACT US',
                 bg=C['bg'], fg=C['text_muted'],
                 font=('Segoe UI', 8, 'bold')).pack(anchor='e', pady=(0, 3))

        for name, email in [('Mirna Kassem',     'kassem_mirna@berkeley.edu'),
                             ('Dimitrios Zekkos', 'zekkos@berkeley.edu')]:
            crow = tk.Frame(contact_block, bg=C['bg'])
            crow.pack(anchor='e', pady=1)

            tk.Label(crow, text=name,
                     bg=C['bg'], fg=C['text'],
                     font=('Segoe UI', 10)).pack(side=tk.LEFT, padx=(0, 6))

            mail_btn = tk.Label(crow, text='✉',
                                bg=C['bg'], fg=C['accent'],
                                font=('Segoe UI', 12),
                                cursor='hand2')
            mail_btn.pack(side=tk.LEFT)
            mail_btn.bind('<Button-1>',
                          lambda e, b=mail_btn, addr=email: _show_email_popup(self.root, b, addr))
            mail_btn.bind('<Enter>', lambda e, b=mail_btn: b.config(fg=C['accent_dark']))
            mail_btn.bind('<Leave>', lambda e, b=mail_btn: b.config(fg=C['accent']))

        # ── Notebook ────────────────────────────────────────────────────
        nb_frame = tk.Frame(self.root, bg=C['bg'], padx=12, pady=8)
        nb_frame.pack(fill=tk.BOTH, expand=True)

        self._notebook = ttk.Notebook(nb_frame)
        self._notebook.pack(fill=tk.BOTH, expand=True)

        self._tab_inners: dict = {}
        self._forward_tab_refs = []   # (frame, title) – shown only in forward mode
        self._back_tab_refs    = []   # (frame, title) – shown only in back mode

        _BA_CATEGORIES = {
            'BA Raster Inputs',
            'BA Hydrology',
            'BA Scalar Inputs',
            'BA Shear Strength',
            'Mapped Landslide Inventory',
        }
        _BA_DISPLAY = {
            'BA Raster Inputs':           '  Topography  ',
            'BA Hydrology':               '  Hydrological Formulations  ',
            'BA Scalar Inputs':           '  Scalar Inputs  ',
            'BA Shear Strength':          '  Shear Strength  ',
            'Mapped Landslide Inventory': '  Mapped Landslide Inventory  ',
        }

        for cat in self._categories:
            tab_outer = tk.Frame(self._notebook, bg=C['surface'])
            display   = _BA_DISPLAY.get(cat, f'  {cat}  ')
            self._notebook.add(tab_outer, text=display)
            _, inner = _scrollable_frame(tab_outer, bg=C['surface'])
            inner.columnconfigure(1, weight=1)
            self._tab_inners[cat] = inner
            if cat in _BA_CATEGORIES:
                self._back_tab_refs.append((tab_outer, display))
            elif cat != 'Project Details':
                self._forward_tab_refs.append((tab_outer, display))

        # ── Running Log tab ──────────────────────────────────────────────
        log_tab = tk.Frame(self._notebook, bg=C['log_bg'])
        self._notebook.add(log_tab, text='  Running Log  ')
        self._log_tab = log_tab
        # Running Log is shown in both forward and back modes, but hidden on startup

        log_container = tk.Frame(log_tab, bg=C['log_bg'])
        log_container.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        self._log_box = tk.Text(
            log_container,
            state=tk.DISABLED,
            font=F['log'],
            bg=C['log_bg'],
            fg=C['log_fg'],
            insertbackground=C['log_fg'],
            selectbackground=C['accent'],
            relief='flat',
            borderwidth=0,
            padx=14, pady=12,
            wrap=tk.WORD
        )
        log_sb = ttk.Scrollbar(log_container, command=self._log_box.yview)
        self._log_box.configure(yscrollcommand=log_sb.set)
        log_sb.pack(side=tk.RIGHT, fill=tk.Y)
        self._log_box.pack(fill=tk.BOTH, expand=True)

        self._log_box.tag_configure('ok',   foreground='#3fb950')
        self._log_box.tag_configure('err',  foreground='#f85149')
        self._log_box.tag_configure('info', foreground=C['log_accent'])
        self._log_box.tag_configure('dim',  foreground='#484f58')

        # ── Outputs tab ──────────────────────────────────────────────────
        outputs_tab = tk.Frame(self._notebook, bg=C['surface'])
        self._notebook.add(outputs_tab, text='  Outputs  ')
        self._forward_tab_refs.append((outputs_tab, '  Outputs  '))

        # Summary note at the top of the Outputs tab
        note_frame = tk.Frame(outputs_tab, bg=C['accent_light'],
                              highlightthickness=1, highlightbackground=C['border'])
        note_frame.pack(fill=tk.X, padx=10, pady=(10, 4))
        note_txt = tk.Text(note_frame, wrap='word', relief='flat', bd=0,
                           bg=C['accent_light'], fg=C['text'],
                           font=F['ui'], cursor='arrow', padx=12, pady=6, height=6)
        note_txt.pack(fill=tk.X)
        note_txt.insert('end',
            'All outputs are written to a sub-folder named "{Project Title}_Outputs" within the selected output directory. '
            'The following files are generated:\n'
            '  •  Shapefiles (.shp) — Spatial geometry of predicted landslides, '
            'compatible with ArcGIS Pro, QGIS, and other GIS platforms.\n'
            '  •  Individual_landslides_with_time.xlsx — Per-landslide statistics '
            '(triggering depth, area, volume) for both incremental and cumulative counts. '
            'Each column represents one time point, with the first column corresponding to t = 0.\n'
            '  •  Total_landslides_with_time.xlsx — Domain-wide cumulative totals '
            'aggregated over all landslides. Each row represents one time point, with a leading Time column '
            '(t = 0, 1, 2, ...) identifying it.')
        note_txt.config(state='disabled')

        self._outputs_canvas_frame = tk.Frame(outputs_tab, bg=C['surface'],
                                              highlightbackground=C['border'],
                                              highlightthickness=1)
        self._outputs_canvas_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(4, 10))
        self._placeholder_canvas(self._outputs_canvas_frame,
                                'Run the model to view DEM with predicted landslides.')

        # ── BA Outputs tab ───────────────────────────────────────────────
        ba_out_tab = tk.Frame(self._notebook, bg=C['surface'])
        self._notebook.add(ba_out_tab, text='  Outputs  ')
        self._ba_outputs_tab = ba_out_tab   # added explicitly after Running Log

        # Canvas area (histograms + MC plot) — fills the whole tab
        self._ba_outputs_canvas_frame = tk.Frame(
            ba_out_tab, bg=C['surface'],
            highlightbackground=C['border'], highlightthickness=1)
        self._ba_outputs_canvas_frame.pack(fill=tk.BOTH, expand=True,
                                           padx=10, pady=10)
        self._placeholder_canvas(self._ba_outputs_canvas_frame,
                                 'Plots appear after the back-analysis run completes.\n'
                                 'This may take a while for large rasters or many landslides.')

        self._create_widgets()
        self._build_tab_extras()

        # ── Cell Inspector tab — built once, added/removed dynamically ──────
        self._inspector_frame = tk.Frame(self._notebook, bg=C['surface'])
        self._build_cell_inspector_tab(self._inspector_frame)
        self._inspector_in_notebook = False

        self._sync_inspector()
        for key in ('time_points', 'Triggering_Indicator'):
            v = self._vars.get(key)
            if v:
                v[0].trace_add('write', self._sync_inspector)
        trig_widget = self._widgets.get('Triggering_Indicator')
        if trig_widget:
            trig_widget.bind('<<ComboboxSelected>>', lambda e: self._sync_inspector())

        # Watch analysis_type param: fire _update_analysis_mode on every change
        _at = self._vars.get('analysis_type')
        if _at:
            _at[0].trace_add('write', self._update_analysis_mode)
        _at_widget = self._widgets.get('analysis_type')
        if _at_widget:
            _at_widget.bind('<<ComboboxSelected>>',
                            lambda e: self._update_analysis_mode())

        # Also re-run _update_analysis_mode when Triggering_Indicator changes,
        # so the Earthquake Inputs tab appears/disappears immediately
        _trig = self._vars.get('Triggering_Indicator')
        if _trig:
            _trig[0].trace_add('write', self._update_analysis_mode)
        if trig_widget:
            trig_widget.bind('<<ComboboxSelected>>',
                             lambda e: self._update_analysis_mode())

        self._update_analysis_mode()


    # ------------------------------------------------------------------
    # Widget generation (fully data-driven — no param names here)
    # ------------------------------------------------------------------

    def _create_widgets(self):
        row_counters = {cat: 0 for cat in self._categories}

        for p in self._params_list:
            cat = p.get('category', self._categories[0])
            if cat not in self._tab_inners:
                continue

            parent = self._tab_inners[cat]
            row    = row_counters[cat]
            key    = p['key']
            ptype  = p['type']
            default = p.get('default', '')

            # ------------------------------------------------------------------
            # Note / info box — full-width, no label or input widget
            # ------------------------------------------------------------------
            if ptype == 'note':
                note_frame = tk.Frame(parent, bg=C['accent_light'],
                                      highlightthickness=1,
                                      highlightbackground=C['border'])
                note_frame.grid(row=row, column=0, columnspan=3,
                                sticky='ew', padx=10, pady=(10, 4))
                note_frame.columnconfigure(0, weight=1)

                note_height = p.get('height', 2)

                txt = tk.Text(note_frame, wrap='word', relief='flat', bd=0,
                              bg=C['accent_light'], fg=C['text'],
                              font=F['ui'], cursor='arrow',
                              padx=12, pady=5,
                              height=note_height)
                txt.pack(fill='x')

                _link_idx = [0]
                for seg, url in p.get('text', []):
                    if url:
                        tag = f'note_link_{_link_idx[0]}'
                        _link_idx[0] += 1
                        txt.tag_configure(tag, foreground=C['accent'],
                                          underline=True)
                        txt.tag_bind(tag, '<Button-1>',
                                     lambda e, u=url: webbrowser.open(u))
                        txt.tag_bind(tag, '<Enter>',
                                     lambda e, w=txt: w.config(cursor='hand2'))
                        txt.tag_bind(tag, '<Leave>',
                                     lambda e, w=txt: w.config(cursor='arrow'))
                        txt.insert('end', seg, (tag,))
                    else:
                        txt.insert('end', seg)

                txt.config(state='disabled')
                self._frames[key] = note_frame
                row_counters[cat] += 1
                continue

            # Alternate row background for readability
            row_bg = C['surface'] if row % 2 == 0 else C['surface_alt']

            # Row container
            pframe = tk.Frame(parent, bg=row_bg)
            pframe.grid(row=row, column=0, columnspan=3,
                        sticky='ew', padx=0, pady=0)
            pframe.columnconfigure(1, weight=1)
            parent.columnconfigure(0, weight=1)
            self._frames[key] = pframe

            # Left padding strip (accent colour for required fields)
            strip_col = C['accent'] if p.get('required') else row_bg
            tk.Frame(pframe, bg=strip_col, width=3).grid(
                row=0, column=0, sticky='ns', padx=0)

            # Label
            lbl = tk.Label(pframe,
                           text=p['label'],
                           bg=row_bg, fg=C['text'],
                           font=F['label'],
                           width=p.get('label_width', 36), anchor='w')
            lbl.grid(row=0, column=1, sticky='w', padx=(10, 8), pady=7)
            if p.get('description'):
                _Tooltip(lbl, p['description'])
                # Small info dot to hint tooltip exists
                tk.Label(pframe, text='ⓘ',
                         bg=row_bg, fg=C['text_muted'],
                         font=F['ui_sm']).grid(row=0, column=1,
                                               sticky='e', padx=(0, 4))

            # Column 2 width defaults to 360 px but can be overridden per param
            # with a 'widget_width' key in crisis_params.py.
            pframe.columnconfigure(2, minsize=p.get('widget_width', 360), weight=0)

            if ptype == 'choice':
                choices = p['choices']
                var = tk.StringVar()
                for raw, display in choices:
                    if raw == default:
                        var.set(display)
                        break
                else:
                    var.set(str(default))

                combo = ttk.Combobox(pframe,
                                     textvariable=var,
                                     values=[d for _, d in choices],
                                     state='readonly',
                                     font=F['ui'])
                combo.grid(row=0, column=2, sticky='ew', padx=(0, 14), pady=5)
                self._vars[key] = (var, 'choice', choices)
                self._widgets[key] = combo
                var.trace_add('write', lambda *_: self._update_visibility())
                if key == 'analysis_type':
                    combo.bind('<<ComboboxSelected>>',
                               lambda e: self._update_analysis_mode())
                if key == 'Hydrological_model_indicator':
                    var.trace_add('write', self._update_hydro_dependencies)
                if key == 'Strength_status':
                    var.trace_add('write', lambda *_: self._update_strength_plot_visibility())
                if key == 'Strength_variation_with_theta':
                    var.trace_add('write', self._schedule_hydro_plot)
                if key == 'ba_Strength_variation_with_theta':
                    var.trace_add('write', self._schedule_ba_hydro_plot)

            elif ptype == 'searchable_choice':
                choices = p['choices']
                all_displays = [f"{raw} — {label}" for raw, label in choices]

                var = tk.StringVar(value='')

                combo = ttk.Combobox(pframe,
                                     textvariable=var,
                                     values=all_displays,
                                     font=F['ui'])
                combo.grid(row=0, column=2, sticky='ew', padx=(0, 14), pady=(5, 1))

                hint = ttk.Label(pframe,
                                 text='🔍  Type to filter, then click ▼ to see matching options',
                                 style='Muted.TLabel')
                hint.grid(row=1, column=2, sticky='w', padx=(2, 14), pady=(0, 4))

                def _make_filter(cb, all_d, v):
                    def _on_key(event=None):
                        if event and event.keysym in ('Down', 'Up', 'Return',
                                                       'Escape', 'Tab', 'Left', 'Right'):
                            return
                        typed = v.get().lower()
                        cb['values'] = [d for d in all_d if typed in d.lower()] or all_d
                    def _on_open():
                        typed = v.get().lower()
                        cb['values'] = [d for d in all_d if typed in d.lower()] or all_d
                    return _on_key, _on_open

                _on_key, _on_open = _make_filter(combo, all_displays, var)
                combo.bind('<KeyRelease>', _on_key)
                combo.configure(postcommand=_on_open)

                self._vars[key] = (var, 'searchable_choice', choices)
                self._widgets[key] = combo

            elif ptype in ('float', 'int', 'str'):
                var = tk.StringVar(value=str(default))
                entry = ttk.Entry(pframe, textvariable=var, font=F['ui'])
                entry.grid(row=0, column=2, sticky='ew', padx=(0, 14), pady=5)
                if key in ('z_min', 'z_max', 'Depth_points'):
                    var.trace_add('write', self._refresh_depth_choices)
                if key in ('ba_z_min', 'ba_z_max', 'ba_Depth_points'):
                    var.trace_add('write', self._refresh_ba_depth_choices)
                pass  # all int/float/str entries are editable by default
                self._vars[key] = (var, ptype, None)
                self._widgets[key] = entry

            elif ptype in ('file', 'dir'):
                var = tk.StringVar(value=str(default))

                if ptype == 'file':
                    ffilter = p.get('file_filter', 'All files (*.*)')
                    browse_cmd = lambda v=var, ff=ffilter: self._browse_file(v, ff)
                else:
                    browse_cmd = lambda v=var: self._browse_dir(v)

                # Sub-frame fills the full 360 px column; entry expands, Browse stays fixed
                inp_frame = tk.Frame(pframe, bg=row_bg)
                inp_frame.grid(row=0, column=2, sticky='ew', padx=(0, 14), pady=5)
                inp_frame.columnconfigure(0, weight=1)

                entry = ttk.Entry(inp_frame, textvariable=var, font=F['ui'])
                entry.grid(row=0, column=0, sticky='ew')

                browse_btn = _flat_btn(inp_frame, 'Browse…', browse_cmd,
                                       bg=C['accent_light'],
                                       fg=C['accent'],
                                       hover_bg=C['accent'],
                                       font=F['ui_sm'],
                                       padx=8, pady=3)
                browse_btn.bind('<Enter>',
                                lambda e, b=browse_btn: b.config(
                                    bg=C['accent'], fg=C['text_on_dark']))
                browse_btn.bind('<Leave>',
                                lambda e, b=browse_btn: b.config(
                                    bg=C['accent_light'], fg=C['accent']))
                browse_btn.grid(row=0, column=1, padx=(4, 0))
                self._vars[key] = (var, ptype, None)
                self._widgets[key] = entry
                if key == 'hydro_files_dir':
                    var.trace_add('write', self._update_hydro_dependencies)
                    var.trace_add('write', self._schedule_hydro_plot)
                if key == 'theta_files_dir':
                    var.trace_add('write', self._schedule_hydro_plot)
                    var.trace_add('write', self._update_theta_auto_fields)
                # dem_file and Cell_size are the only Topography raster inputs —
                # Slope, Flow Direction, and Flow Accumulation are all auto-derived
                # from them, so both must refresh the preview
                _RASTER_KEYS = {'dem_file', 'Cell_size', 'pga_file', 'pgv_file'}
                if key in _RASTER_KEYS:
                    var.trace_add('write', self._schedule_raster_plot)
                if key == 'dem_file':
                    var.trace_add('write', self._invalidate_hydro_vrange)
                    var.trace_add('write', self._auto_fill_dem_geo)
                _STRENGTH_KEYS = {'cohesion_3d_file', 'phi_3d_file'}
                if key in _STRENGTH_KEYS:
                    var.trace_add('write', self._schedule_strength_plot)
                # ── Back-analysis raster / hydro auto-plot traces ────────────
                # ba_dem_file is now the only BA raster input — Slope, Flow
                # Direction, and Flow Accumulation are all auto-derived from
                # it and ba_Cell_size, so both must refresh the preview (and
                # therefore the cache _run_back_analysis reuses on Run)
                if key in ('ba_dem_file', 'ba_Cell_size'):
                    var.trace_add('write', self._schedule_ba_raster_plot)
                if key == 'ba_dem_file':
                    var.trace_add('write', self._invalidate_ba_hydro_vrange)
                    var.trace_add('write', self._auto_fill_ba_dem_geo)
                if key == 'ba_hydro_files_dir':
                    var.trace_add('write', self._schedule_ba_hydro_plot)
                    var.trace_add('write', self._invalidate_ba_hydro_vrange)
                    var.trace_add('write', self._update_ba_hydro_auto_fields)
                if key == 'ba_theta_files_dir':
                    var.trace_add('write', self._schedule_ba_hydro_plot)
                    var.trace_add('write', self._update_ba_theta_auto_fields)
                # ── Size validation on every raster/coord file ───────────────
                _ALL_FILE_KEYS = {
                    'pga_file', 'pgv_file',
                    'cohesion_3d_file', 'phi_3d_file',
                }
                if key in _ALL_FILE_KEYS:
                    _k = key  # capture for closure
                    var.trace_add('write', lambda *_, _key=_k: self._schedule_size_check(_key))

            row_counters[cat] += 1

        # Padding row at the bottom of each tab
        for cat, inner in self._tab_inners.items():
            tk.Frame(inner, bg=C['surface'], height=8).grid(
                row=row_counters[cat], column=0, columnspan=4, sticky='ew')
            row_counters[cat] += 1

        self._tab_row_counters = row_counters
        self._update_visibility()
        self._update_hydro_dependencies()

    # ------------------------------------------------------------------
    # Hydrology-driven auto-calculation of time/depth counts
    # ------------------------------------------------------------------

    def _scan_hydro_dir(self, directory):
        """Return (num_time_points, num_depth_layers) from Pressure_head_T{t}Z{d+1}.h5 filenames."""
        import re
        pattern = re.compile(r'^Pressure_head_T(\d+)Z(\d+)\.h5$')
        times, depths = set(), set()
        try:
            for fname in os.listdir(directory):
                m = pattern.match(fname)
                if m:
                    times.add(int(m.group(1)))
                    depths.add(int(m.group(2)))
        except Exception:
            return None, None
        if not times or not depths:
            return None, None
        return max(times) + 1, max(depths)

    def _validate_hydro_dir(self, directory, expected_nd, expected_nt):
        """
        Check that every T has exactly expected_nd Z-files and every Z has exactly expected_nt T-files.
        Returns a list of error strings (empty = OK).
        """
        import re
        pattern = re.compile(r'^Pressure_head_T(\d+)Z(\d+)\.h5$')
        from collections import defaultdict
        t_to_z = defaultdict(set)
        z_to_t = defaultdict(set)
        try:
            for fname in os.listdir(directory):
                m = pattern.match(fname)
                if m:
                    t, z = int(m.group(1)), int(m.group(2))
                    t_to_z[t].add(z)
                    z_to_t[z].add(t)
        except Exception:
            return []

        errors = []
        bad_t = sorted(t for t, zs in t_to_z.items() if len(zs) != expected_nd)
        if bad_t:
            sample = bad_t[:5]
            counts = [len(t_to_z[t]) for t in sample]
            detail = ', '.join(f'T{t}({c} Z-files)' for t, c in zip(sample, counts))
            if len(bad_t) > 5:
                detail += f' … (+{len(bad_t)-5} more)'
            errors.append(
                f'Expected {expected_nd} depth file(s) per time step, but found mismatches:\n  {detail}')

        bad_z = sorted(z for z, ts in z_to_t.items() if len(ts) != expected_nt)
        if bad_z:
            sample = bad_z[:5]
            counts = [len(z_to_t[z]) for z in sample]
            detail = ', '.join(f'Z{z}({c} T-files)' for z, c in zip(sample, counts))
            if len(bad_z) > 5:
                detail += f' … (+{len(bad_z)-5} more)'
            errors.append(
                f'Expected {expected_nt} time file(s) per depth layer, but found mismatches:\n  {detail}')

        return errors

    def _check_hydro_raster_sizes(self, directory, pattern, expected_shape, sample_limit=10):
        """
        Check that a sample of .h5 files matching `pattern` in `directory` have a
        raster size equal to `expected_shape` (rows, cols). Sampled rather than
        exhaustive for GUI responsiveness — back_analysis.run_model() performs
        the authoritative full-file check at run time regardless.
        Returns a list of mismatch description strings (empty = all sampled OK).
        """
        import h5py as _h5
        mismatches = []
        checked = 0
        try:
            fnames = sorted(f for f in os.listdir(directory) if pattern.match(f))
        except Exception:
            return []
        for fname in fnames:
            if checked >= sample_limit:
                break
            fpath = os.path.join(directory, fname)
            try:
                with _h5.File(fpath, 'r') as f:
                    shape = f['data'].shape[:2]
            except Exception:
                continue
            checked += 1
            if tuple(shape) != tuple(expected_shape):
                mismatches.append(f'{fname}: {shape[0]}×{shape[1]}')
        return mismatches

    def _update_hydro_dependencies(self, *_):
        hydro_ind = self._get_raw_value('Hydrological_model_indicator')

        def _set_var(key, value):
            t = self._vars.get(key)
            if t:
                t[0].set(str(value))

        def _set_state(key, state):
            w = self._widgets.get(key)
            if w:
                try:
                    w.config(state=state)
                except Exception:
                    pass

        if hydro_ind == 0:  # ── Constant Ru ─────────────────────────────
            _set_var('Earthquake_time_point', 0)
            _set_state('Earthquake_time_point', 'disabled')
            # Force Bishop's χ = 1 (fully saturated) when using constant Ru
            _set_var('Strength_variation_with_theta',
                     "No — Fully saturated assumption, Bishop's effective stress coefficient χ = 1")
            # Hide the pore pressure plot panel
            if hasattr(self, '_hydro_plot_container'):
                try:
                    self._hydro_plot_container.grid_remove()
                except Exception:
                    pass

            # A constant Ru has no spatial or temporal variability — every
            # time step would reproduce an identical result — so warn (but
            # don't auto-fix; the user decides) if time_points is already
            # set to something other than 1.
            time_pts_raw = self._get_raw_value('time_points')
            if time_pts_raw not in (None, ''):
                try:
                    time_pts_val = int(time_pts_raw)
                except (TypeError, ValueError):
                    time_pts_val = None
                if time_pts_val is not None and time_pts_val != 1:
                    messagebox.showwarning(
                        'Constant Ru with Multiple Time Points',
                        f'"Number of time points" is set to {time_pts_raw}, but a constant Ru '
                        'has no spatial or temporal variability — every time step would produce '
                        'an identical result.\n\nConsider setting "Number of time points" to 1.')

        else:               # ── Variable pore pressure files ────────────
            _set_state('Earthquake_time_point', 'normal')
            _set_state('time_points', 'normal')
            # Show the pore pressure plot panel
            if hasattr(self, '_hydro_plot_container'):
                try:
                    self._hydro_plot_container.grid()
                except Exception:
                    pass

            hydro_dir = self._get_raw_value('hydro_files_dir')
            if hydro_dir and os.path.isdir(hydro_dir):
                dem_path = self._get_raw_value('dem_file')
                user_nd_raw = self._get_raw_value('Depth_points')
                user_nt_raw = self._get_raw_value('time_points')

                # The directory scan, DEM read, and per-file size sampling
                # below are all disk I/O — done in a background thread so a
                # large DEM (or a slow/cloud-synced drive) can't freeze the
                # GUI on every directory change.
                def _work():
                    n_t, n_d = self._scan_hydro_dir(hydro_dir)
                    if n_t is None:
                        return ('no_files', None)

                    user_nd = int(user_nd_raw) if user_nd_raw else None
                    user_nt = int(user_nt_raw) if user_nt_raw else None
                    errs = []
                    # Only validate if the user has already entered values —
                    # never auto-fill; let the user decide.
                    if user_nd is not None and user_nd != n_d:
                        errs.append(
                            f'"Number of subsurface layers" is set to {user_nd} but the '
                            f'pore pressure directory contains {n_d} depth layer(s).\n'
                            f'Please update "Number of subsurface layers" to {n_d}.')
                    if user_nt is not None and user_nt != n_t:
                        errs.append(
                            f'"Number of time points" is set to {user_nt} but the '
                            f'pore pressure directory contains {n_t} time step(s).\n'
                            f'Please update "Number of time points" to {n_t}.')
                    if not errs and user_nd is not None and user_nt is not None:
                        errs = self._validate_hydro_dir(hydro_dir, user_nd, user_nt)

                    # Also check that the pore pressure rasters match the DEM's
                    # size — a mismatch used to be silently cropped in
                    # forward_analysis.py, which could misalign pore pressure
                    # data against elevation with no warning
                    if dem_path:
                        try:
                            import re as _re
                            dem_shape = self._load_dem_array(dem_path).shape[:2]
                            size_mismatches = self._check_hydro_raster_sizes(
                                hydro_dir, _re.compile(r'^Pressure_head_T\d+Z\d+\.h5$'), dem_shape)
                            if size_mismatches:
                                errs.append(
                                    f'The following pore pressure files do not match the DEM size '
                                    f'({dem_shape[0]}×{dem_shape[1]}):\n  ' + '\n  '.join(size_mismatches))
                        except Exception:
                            pass

                    return ('checked', errs)

                def _done(result):
                    kind, errs = result if result else (None, None)
                    if kind == 'no_files':
                        self._log(
                            'No Pressure_head_T{t}Z{d+1}.h5 files found in the specified directory.\n',
                            tag='err')
                        return
                    if kind != 'checked':
                        return
                    if errs:
                        messagebox.showerror(
                            'Pore Pressure Directory Mismatch',
                            'The pore pressure directory does not match the specified inputs:\n\n' +
                            '\n\n'.join(errs) +
                            '\n\nThe directory entry has been cleared. Please fix the inputs and re-select the directory.')
                        # Clear the directory so the user must re-select a consistent one
                        if 'hydro_files_dir' in self._vars:
                            self._vars['hydro_files_dir'][0].set('')
                        if hasattr(self, '_hydro_canvas_frame'):
                            self._clear_canvas_frame(self._hydro_canvas_frame)
                            self._placeholder_canvas(self._hydro_canvas_frame,
                                                     'Set pore pressure directory above to auto-plot.')

                self._run_in_bg(_work, _done, cancel_attr='_hydro_dep_check_gen')

    def _refresh_depth_choices(self, *_):
        try:
            z_min = float(self._get_raw_value('z_min') or 0.125)
            z_max = float(self._get_raw_value('z_max') or 1.875)
            n_d   = int(self._get_raw_value('Depth_points') or 8)
            if n_d < 1:
                n_d = 1
            depths = np.linspace(z_min, z_max, n_d)
            values = [f'{v:.4f}' for v in depths]
        except Exception:
            values = []
        for combo_attr, var_attr in (('_hydro_d_combo', '_hydro_d_var'),
                                     ('_strength_d_combo', '_strength_d_var')):
            combo = getattr(self, combo_attr, None)
            if combo is None:
                continue
            current = getattr(self, var_attr).get()
            combo['values'] = values
            if values and current not in values:
                getattr(self, var_attr).set(values[0])

    def _update_strength_plot_visibility(self):
        if not hasattr(self, '_strength_plot_container'):
            return
        if self._get_raw_value('Strength_status') == 1:
            self._strength_plot_container.grid()
        else:
            self._strength_plot_container.grid_remove()

    # ------------------------------------------------------------------
    # Auto-plot debounce helpers
    # ------------------------------------------------------------------

    def _schedule_raster_plot(self, *_):
        if hasattr(self, '_raster_plot_job'):
            self.root.after_cancel(self._raster_plot_job)
        self._raster_plot_job = self.root.after(400, self._plot_rasters)

    def _schedule_hydro_plot(self, *_):
        if hasattr(self, '_hydro_plot_job'):
            self.root.after_cancel(self._hydro_plot_job)
        self._hydro_plot_job = self.root.after(400, lambda: self._plot_hydro(silent=True))

    def _compute_hydro_vrange(self, hydro_dir):
        """Scan all Pressure_head_T{t}Z{d+1}.h5 files and return (vmin, vmax) across the full dataset."""
        import re
        pattern = re.compile(r'^Pressure_head_T\d+Z\d+\.h5$')

        # Physical cap: pressure head cannot exceed the maximum soil depth (z_max).
        # This is the same clip the model applies when loading pressure files, and it
        # prevents a handful of hydrological-simulation artifacts from blowing up the scale.
        try:
            p_cap = float(self._get_raw_value('z_max') or 0)
            if p_cap <= 0:
                p_cap = None
        except (TypeError, ValueError):
            p_cap = None

        # NoData mask from DEM (values < -900).  The DEM is trimmed by one row
        # inside run_model, so its shape may differ from the pressure grids; we
        # handle both possibilities below.
        dem_path = self._get_raw_value('dem_file')
        dem_mask     = None   # shape == pressure shape  (trimmed DEM)
        dem_mask_raw = None   # shape == raw DEM shape
        if dem_path:
            try:
                dem_arr = self._load_dem_array(dem_path)
                dem_mask_raw = np.isnan(dem_arr)
                # Also build the row-trimmed version that matches run_model loading
                dem_mask = np.isnan(dem_arr[:-1, :])
            except Exception:
                pass

        gmin, gmax = np.inf, -np.inf
        for fname in os.listdir(hydro_dir):
            if not pattern.match(fname):
                continue
            try:
                a = self._load_h5_array(os.path.join(hydro_dir, fname))
                # Apply NoData mask whichever shape matches
                if dem_mask is not None and dem_mask.shape == a.shape:
                    a[dem_mask] = np.nan
                elif dem_mask_raw is not None and dem_mask_raw.shape == a.shape:
                    a[dem_mask_raw] = np.nan
                # Clip to physical maximum (same as model's np.minimum(P2D, depth))
                if p_cap is not None:
                    a = np.minimum(a, p_cap)
                lo = np.nanmin(a)
                hi = np.nanmax(a)
                if lo < gmin: gmin = lo
                if hi > gmax: gmax = hi
            except Exception:
                continue

        if np.isinf(gmin) or np.isinf(gmax):
            return None, None
        # Round outward to nearest integer for a clean scale
        return float(np.floor(gmin)), float(np.ceil(gmax))

    def _invalidate_hydro_vrange(self, *_):
        self._hydro_vrange_dir = None

    def _compute_theta_vrange(self, theta_dir):
        """Scan all Theta_T{t}Z{d+1}.h5 files and return (vmin, vmax) across the full dataset."""
        import re
        pattern = re.compile(r'^Theta_T\d+Z\d+\.h5$')
        dem_path = self._get_raw_value('dem_file')
        dem_mask = None
        if dem_path:
            try:
                dem_arr = self._load_dem_array(dem_path)
                dem_mask = np.isnan(dem_arr)
            except Exception:
                pass

        gmin, gmax = np.inf, -np.inf
        for fname in os.listdir(theta_dir):
            if not pattern.match(fname):
                continue
            try:
                a = self._load_h5_array(os.path.join(theta_dir, fname))
                if dem_mask is not None and dem_mask.shape == a.shape:
                    a[dem_mask] = np.nan
                lo = np.nanmin(a)
                hi = np.nanmax(a)
                if lo < gmin: gmin = lo
                if hi > gmax: gmax = hi
            except Exception:
                continue

        if np.isinf(gmin) or np.isinf(gmax):
            return None, None
        return float(np.floor(gmin * 100) / 100), float(np.ceil(gmax * 100) / 100)

    def _schedule_strength_plot(self, *_):
        if hasattr(self, '_strength_plot_job'):
            self.root.after_cancel(self._strength_plot_job)
        self._strength_plot_job = self.root.after(400, lambda: self._plot_strength(silent=True))

    # ── Back-analysis debounce helpers ────────────────────────────────

    def _schedule_ba_raster_plot(self, *_):
        if hasattr(self, '_ba_raster_plot_job'):
            self.root.after_cancel(self._ba_raster_plot_job)
        self._ba_raster_plot_job = self.root.after(400, self._plot_ba_rasters)

    def _schedule_ba_hydro_plot(self, *_):
        if hasattr(self, '_ba_hydro_plot_job'):
            self.root.after_cancel(self._ba_hydro_plot_job)
        self._ba_hydro_plot_job = self.root.after(400, lambda: self._plot_ba_hydro(silent=True))

    def _invalidate_ba_hydro_vrange(self, *_):
        self._ba_hydro_vrange_dir = None

    def _update_ba_hydro_auto_fields(self, *_):
        """Validate ba_hydro_files_dir against user-entered ba_Depth_points and
        ba_Time_points. All disk I/O (directory scan, DEM read, per-file size
        sampling) runs in a background thread — a large DEM or a slow/cloud-
        synced drive could otherwise freeze the GUI on every directory change."""
        hydro_dir = self._get_raw_value('ba_hydro_files_dir')
        if not hydro_dir or not os.path.isdir(hydro_dir):
            return

        user_nd_raw = self._get_raw_value('ba_Depth_points')
        user_nt_raw = self._get_raw_value('ba_Time_points')
        dem_path = self._get_raw_value('ba_dem_file')

        def _work():
            import re
            n_t, n_d = self._scan_hydro_dir(hydro_dir)
            if n_t is None:
                return None

            user_nd = int(user_nd_raw) if user_nd_raw else None
            user_nt = int(user_nt_raw) if user_nt_raw else None

            errs = []
            if user_nd is not None and user_nd != n_d:
                errs.append(
                    f'"Number of subsurface layers" is set to {user_nd} but the '
                    f'pore pressure directory contains {n_d} depth layer(s).\n'
                    f'Please update "Number of subsurface layers" to {n_d}.')
            if user_nt is not None and user_nt != n_t:
                errs.append(
                    f'"Number of time points" is set to {user_nt} but the '
                    f'pore pressure directory contains {n_t} time step(s).\n'
                    f'Please update "Number of time points" to {n_t}.')
            if not errs and user_nd is not None and user_nt is not None:
                errs = self._validate_hydro_dir(hydro_dir, user_nd, user_nt)

            # Also check that the pore pressure rasters match the DEM's size —
            # a mismatch used to be silently cropped in back_analysis.py, which
            # could misalign pore pressure data against elevation with no warning
            if dem_path:
                try:
                    dem_shape = self._load_dem_array(dem_path).shape[:2]
                    size_mismatches = self._check_hydro_raster_sizes(
                        hydro_dir, re.compile(r'^Pressure_head_T\d+Z\d+\.h5$'), dem_shape)
                    if size_mismatches:
                        errs.append(
                            f'The following pore pressure files do not match the DEM size '
                            f'({dem_shape[0]}×{dem_shape[1]}):\n  ' + '\n  '.join(size_mismatches))
                except Exception:
                    pass

            return errs

        def _done(errs):
            if not errs:
                return
            messagebox.showerror(
                'Pore Pressure Directory Mismatch',
                'The pore pressure directory does not match the specified inputs:\n\n' +
                '\n\n'.join(errs) +
                '\n\nThe directory entry has been cleared. Please fix the inputs and re-select the directory.')
            if 'ba_hydro_files_dir' in self._vars:
                self._vars['ba_hydro_files_dir'][0].set('')
            if hasattr(self, '_ba_hydro_canvas_frame'):
                self._clear_canvas_frame(self._ba_hydro_canvas_frame)
                self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                         'Browse a pore pressure directory above to plot.')

        self._run_in_bg(_work, _done, cancel_attr='_ba_hydro_dep_check_gen')

    def _update_ba_theta_auto_fields(self, *_):
        """Check ba_theta_files_dir rasters against the DEM's size, same as
        _update_ba_hydro_auto_fields does for the pore pressure directory.
        Runs in a background thread — see _update_ba_hydro_auto_fields."""
        theta_dir = self._get_raw_value('ba_theta_files_dir')
        if not theta_dir or not os.path.isdir(theta_dir):
            return
        dem_path = self._get_raw_value('ba_dem_file')
        if not dem_path:
            return

        def _work():
            import re
            try:
                dem_shape = self._load_dem_array(dem_path).shape[:2]
                return self._check_hydro_raster_sizes(
                    theta_dir, re.compile(r'^Theta_T\d+Z\d+\.h5$'), dem_shape), dem_shape
            except Exception:
                return None

        def _done(result):
            if not result:
                return
            size_mismatches, dem_shape = result
            if size_mismatches:
                messagebox.showerror(
                    'Volumetric Water Content Directory Mismatch',
                    f'The following volumetric water content files do not match the DEM size '
                    f'({dem_shape[0]}×{dem_shape[1]}):\n\n  ' + '\n  '.join(size_mismatches) +
                    '\n\nThe directory entry has been cleared. Please fix the inputs and re-select the directory.')
                if 'ba_theta_files_dir' in self._vars:
                    self._vars['ba_theta_files_dir'][0].set('')
                if hasattr(self, '_ba_hydro_canvas_frame'):
                    self._clear_canvas_frame(self._ba_hydro_canvas_frame)
                    self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                             'Browse a pore pressure directory above to plot.')

        self._run_in_bg(_work, _done, cancel_attr='_ba_theta_dep_check_gen')

    def _update_theta_auto_fields(self, *_):
        """Check theta_files_dir rasters against the DEM's size (forward
        analysis), same as _update_ba_theta_auto_fields does for back-analysis.
        Runs in a background thread — see _update_ba_hydro_auto_fields."""
        theta_dir = self._get_raw_value('theta_files_dir')
        if not theta_dir or not os.path.isdir(theta_dir):
            return
        dem_path = self._get_raw_value('dem_file')
        if not dem_path:
            return

        def _work():
            import re
            try:
                dem_shape = self._load_dem_array(dem_path).shape[:2]
                return self._check_hydro_raster_sizes(
                    theta_dir, re.compile(r'^Theta_T\d+Z\d+\.h5$'), dem_shape), dem_shape
            except Exception:
                return None

        def _done(result):
            if not result:
                return
            size_mismatches, dem_shape = result
            if size_mismatches:
                messagebox.showerror(
                    'Volumetric Water Content Directory Mismatch',
                    f'The following volumetric water content files do not match the DEM size '
                    f'({dem_shape[0]}×{dem_shape[1]}):\n\n  ' + '\n  '.join(size_mismatches) +
                    '\n\nThe directory entry has been cleared. Please fix the inputs and re-select the directory.')
                if 'theta_files_dir' in self._vars:
                    self._vars['theta_files_dir'][0].set('')
                if hasattr(self, '_hydro_canvas_frame'):
                    self._clear_canvas_frame(self._hydro_canvas_frame)
                    self._placeholder_canvas(self._hydro_canvas_frame,
                                             'Set pore pressure directory above to auto-plot.')

        self._run_in_bg(_work, _done, cancel_attr='_theta_dep_check_gen')

    def _refresh_ba_depth_choices(self, *_):
        try:
            z_min = float(self._get_raw_value('ba_z_min') or 0.125)
            z_max = float(self._get_raw_value('ba_z_max') or 1.875)
            n_d   = int(self._get_raw_value('ba_Depth_points') or 8)
            if n_d < 1:
                n_d = 1
            depths = np.linspace(z_min, z_max, n_d)
            values = [f'{v:.4f}' for v in depths]
        except Exception:
            values = []
        for combo_attr, var_attr in (('_ba_hydro_d_combo', '_ba_hydro_d_var'),):
            combo = getattr(self, combo_attr, None)
            if combo is None:
                continue
            current = getattr(self, var_attr).get()
            combo['values'] = values
            if values and current not in values:
                getattr(self, var_attr).set(values[0])

    def _load_ba_dem_mask(self):
        """Return boolean NaN mask from BA DEM (True = invalid cell). None if unavailable."""
        path = self._get_raw_value('ba_dem_file')
        if not path:
            return None
        try:
            return np.isnan(self._load_dem_array(path))
        except Exception:
            return None

    def _compute_ba_hydro_vrange(self, hydro_dir):
        """Scan all Pressure_head_T{t}Z{d}.h5 in the BA hydro dir, return (vmin, vmax)."""
        import re
        pattern = re.compile(r'^Pressure_head_T\d+Z\d+\.h5$')
        try:
            p_cap = float(self._get_raw_value('ba_z_max') or 0)
            if p_cap <= 0:
                p_cap = None
        except (TypeError, ValueError):
            p_cap = None
        dem_path = self._get_raw_value('ba_dem_file')
        dem_mask = dem_mask_raw = None
        if dem_path:
            try:
                dem_arr = self._load_dem_array(dem_path)
                dem_mask_raw = np.isnan(dem_arr)
                dem_mask = np.isnan(dem_arr[:-1, :])
            except Exception:
                pass
        gmin, gmax = np.inf, -np.inf
        for fname in os.listdir(hydro_dir):
            if not pattern.match(fname):
                continue
            try:
                a = self._load_h5_array(os.path.join(hydro_dir, fname))
                if dem_mask is not None and dem_mask.shape == a.shape:
                    a[dem_mask] = np.nan
                elif dem_mask_raw is not None and dem_mask_raw.shape == a.shape:
                    a[dem_mask_raw] = np.nan
                if p_cap is not None:
                    a = np.minimum(a, p_cap)
                lo, hi = np.nanmin(a), np.nanmax(a)
                if lo < gmin: gmin = lo
                if hi > gmax: gmax = hi
            except Exception:
                continue
        if np.isinf(gmin) or np.isinf(gmax):
            return None, None
        return float(np.floor(gmin)), float(np.ceil(gmax))

    def _compute_ba_theta_vrange(self, theta_dir):
        """Scan all Theta_T{t}Z{d}.h5 in the BA theta dir, return (vmin, vmax)."""
        import re
        pattern = re.compile(r'^Theta_T\d+Z\d+\.h5$')
        dem_path = self._get_raw_value('ba_dem_file')
        dem_mask = None
        if dem_path:
            try:
                dem_arr = self._load_dem_array(dem_path)
                dem_mask = np.isnan(dem_arr)
            except Exception:
                pass
        gmin, gmax = np.inf, -np.inf
        for fname in os.listdir(theta_dir):
            if not pattern.match(fname):
                continue
            try:
                a = self._load_h5_array(os.path.join(theta_dir, fname))
                if dem_mask is not None and dem_mask.shape == a.shape:
                    a[dem_mask] = np.nan
                lo, hi = np.nanmin(a), np.nanmax(a)
                if lo < gmin: gmin = lo
                if hi > gmax: gmax = hi
            except Exception:
                continue
        if np.isinf(gmin) or np.isinf(gmax):
            return None, None
        return float(np.floor(gmin * 100) / 100), float(np.ceil(gmax * 100) / 100)

    # ------------------------------------------------------------------
    # Tab extras — plot panels added after param widgets
    # ------------------------------------------------------------------

    def _build_tab_extras(self):
        self._build_raster_plot_panel()
        self._build_hydro_plot_panel()
        self._build_strength_plot_panel()
        self._build_ba_raster_plot_panel()
        self._build_ba_hydro_plot_panel()

    def _plot_section_header(self, inner, row, title):
        tk.Frame(inner, bg=C['border'], height=1).grid(
            row=row, column=0, columnspan=4, sticky='ew', padx=0, pady=(6, 0))
        row += 1
        hdr = tk.Frame(inner, bg=C['surface'])
        hdr.grid(row=row, column=0, columnspan=4, sticky='ew', padx=10, pady=(6, 2))
        tk.Label(hdr, text=title, bg=C['surface'], fg=C['accent'],
                 font=F['ui_bold']).pack(side=tk.LEFT)
        return hdr, row + 1

    def _placeholder_canvas(self, parent, text):
        lbl = tk.Label(parent, text=text,
                       bg=C['surface'], fg=C['text_muted'],
                       font=F['ui_sm'], pady=18)
        lbl.pack()
        return lbl

    # ── Raster Inputs plot panel ─────────────────────────────────────

    def _build_raster_plot_panel(self):
        cat = 'Topography'
        inner = self._tab_inners.get(cat)
        if inner is None:
            return
        row = self._tab_row_counters.get(cat, 0)

        hdr_frame, row = self._plot_section_header(inner, row, 'Raster Preview')


        self._raster_canvas_frame = tk.Frame(inner, bg=C['surface'],
                                              highlightbackground=C['border'],
                                              highlightthickness=1)
        self._raster_canvas_frame.grid(row=row, column=0, columnspan=4,
                                        sticky='ew', padx=10, pady=(2, 10))
        self._placeholder_canvas(self._raster_canvas_frame,
                                  'Browse a raster file above to auto-plot.')
        self._tab_row_counters[cat] = row + 1

    # ── Hydrological Formulations plot panel ─────────────────────────

    def _build_hydro_plot_panel(self):
        cat = 'Hydrological Formulations'
        inner = self._tab_inners.get(cat)
        if inner is None:
            return
        row = self._tab_row_counters.get(cat, 0)

        # Single container — hide/show this one frame to toggle the whole panel
        self._hydro_plot_container = tk.Frame(inner, bg=C['surface'])
        self._hydro_plot_container.grid(row=row, column=0, columnspan=4,
                                         sticky='ew', padx=0, pady=0)

        # Separator + header inside the container
        tk.Frame(self._hydro_plot_container, bg=C['border'], height=1).pack(
            fill=tk.X, pady=(6, 0))
        hdr = tk.Frame(self._hydro_plot_container, bg=C['surface'])
        hdr.pack(fill=tk.X, padx=10, pady=(6, 2))
        tk.Label(hdr, text='Pore Pressure & Volumetric Water Content Preview',
                 bg=C['surface'], fg=C['accent'],
                 font=F['ui_bold']).pack(side=tk.LEFT)


        ctrl = tk.Frame(self._hydro_plot_container, bg=C['surface'])
        ctrl.pack(anchor='w', padx=10, pady=(2, 4))

        tk.Label(ctrl, text='Time step:', bg=C['surface'], fg=C['text'],
                 font=F['ui']).pack(side=tk.LEFT)
        self._hydro_t_var = tk.StringVar(value='0')
        ttk.Spinbox(ctrl, from_=0, to=999, textvariable=self._hydro_t_var,
                    width=6, font=F['ui']).pack(side=tk.LEFT, padx=(4, 16))
        self._hydro_t_var.trace_add('write', self._schedule_hydro_plot)

        tk.Label(ctrl, text='Depth (m):', bg=C['surface'], fg=C['text'],
                 font=F['ui']).pack(side=tk.LEFT)
        self._hydro_d_var = tk.StringVar()
        self._hydro_d_combo = ttk.Combobox(ctrl, textvariable=self._hydro_d_var,
                                            width=10, font=F['ui'], state='readonly')
        self._hydro_d_combo.pack(side=tk.LEFT, padx=(4, 16))
        self._hydro_d_var.trace_add('write', self._schedule_hydro_plot)

        self._hydro_canvas_frame = tk.Frame(self._hydro_plot_container, bg=C['surface'],
                                             highlightbackground=C['border'],
                                             highlightthickness=1)
        self._hydro_canvas_frame.pack(fill=tk.X, padx=10, pady=(2, 10))
        self._placeholder_canvas(self._hydro_canvas_frame,
                                  'Set pore pressure directory above to auto-plot.')
        self._tab_row_counters[cat] = row + 1
        self._refresh_depth_choices()

    # ── Shear Strength plot panel ─────────────────────────────────────

    def _build_strength_plot_panel(self):
        cat = 'Shear Strength'
        inner = self._tab_inners.get(cat)
        if inner is None:
            return
        row = self._tab_row_counters.get(cat, 0)

        # Single container — hidden when uniform strength is selected
        self._strength_plot_container = tk.Frame(inner, bg=C['surface'])
        self._strength_plot_container.grid(row=row, column=0, columnspan=4,
                                            sticky='ew', padx=0, pady=0)
        row += 1

        tk.Frame(self._strength_plot_container, bg=C['border'], height=1).pack(
            fill=tk.X, pady=(6, 0))
        hdr = tk.Frame(self._strength_plot_container, bg=C['surface'])
        hdr.pack(fill=tk.X, padx=10, pady=(6, 2))
        tk.Label(hdr, text='Strength Parameter Preview',
                 bg=C['surface'], fg=C['text'],
                 font=F['label']).pack(side=tk.LEFT)

        ctrl = tk.Frame(self._strength_plot_container, bg=C['surface'])
        ctrl.pack(anchor='w', padx=10, pady=(2, 4))

        tk.Label(ctrl, text='Depth (m):', bg=C['surface'], fg=C['text'],
                 font=F['ui']).pack(side=tk.LEFT)
        self._strength_d_var = tk.StringVar()
        self._strength_d_combo = ttk.Combobox(ctrl, textvariable=self._strength_d_var,
                                               width=10, font=F['ui'], state='readonly')
        self._strength_d_combo.pack(side=tk.LEFT, padx=(4, 16))
        self._strength_d_var.trace_add('write', self._schedule_strength_plot)

        self._strength_canvas_frame = tk.Frame(self._strength_plot_container, bg=C['surface'],
                                                highlightbackground=C['border'],
                                                highlightthickness=1)
        self._strength_canvas_frame.pack(fill=tk.X, padx=10, pady=(2, 10))
        self._placeholder_canvas(self._strength_canvas_frame,
                                  'Load 3D cohesion and friction angle files above to auto-plot.')

        self._tab_row_counters[cat] = row

        # Set initial visibility and depth choices
        self._refresh_depth_choices()
        self._update_strength_plot_visibility()

    # ── BA Raster Inputs plot panel ───────────────────────────────────

    def _build_ba_raster_plot_panel(self):
        cat = 'BA Raster Inputs'
        inner = self._tab_inners.get(cat)
        if inner is None:
            return
        row = self._tab_row_counters.get(cat, 0)

        _, row = self._plot_section_header(inner, row, 'Raster Preview')

        self._ba_raster_canvas_frame = tk.Frame(inner, bg=C['surface'],
                                                highlightbackground=C['border'],
                                                highlightthickness=1)
        self._ba_raster_canvas_frame.grid(row=row, column=0, columnspan=4,
                                          sticky='ew', padx=10, pady=(2, 10))
        self._placeholder_canvas(self._ba_raster_canvas_frame,
                                 'Browse a raster file above to auto-plot.')
        self._tab_row_counters[cat] = row + 1

    # ── BA Hydrological Formulations plot panel ───────────────────────

    def _build_ba_hydro_plot_panel(self):
        cat = 'BA Hydrology'
        inner = self._tab_inners.get(cat)
        if inner is None:
            return
        row = self._tab_row_counters.get(cat, 0)

        self._ba_hydro_plot_container = tk.Frame(inner, bg=C['surface'])
        self._ba_hydro_plot_container.grid(row=row, column=0, columnspan=4,
                                           sticky='ew', padx=0, pady=0)

        tk.Frame(self._ba_hydro_plot_container, bg=C['border'], height=1).pack(
            fill=tk.X, pady=(6, 0))
        hdr = tk.Frame(self._ba_hydro_plot_container, bg=C['surface'])
        hdr.pack(fill=tk.X, padx=10, pady=(6, 2))
        tk.Label(hdr, text='Pore Pressure & Volumetric Water Content Preview',
                 bg=C['surface'], fg=C['accent'],
                 font=F['ui_bold']).pack(side=tk.LEFT)

        ctrl = tk.Frame(self._ba_hydro_plot_container, bg=C['surface'])
        ctrl.pack(anchor='w', padx=10, pady=(2, 4))

        tk.Label(ctrl, text='Time step:', bg=C['surface'], fg=C['text'],
                 font=F['ui']).pack(side=tk.LEFT)
        self._ba_hydro_t_var = tk.StringVar(value='0')
        ttk.Spinbox(ctrl, from_=0, to=999, textvariable=self._ba_hydro_t_var,
                    width=6, font=F['ui']).pack(side=tk.LEFT, padx=(4, 16))
        self._ba_hydro_t_var.trace_add('write', self._schedule_ba_hydro_plot)

        tk.Label(ctrl, text='Depth (m):', bg=C['surface'], fg=C['text'],
                 font=F['ui']).pack(side=tk.LEFT)
        self._ba_hydro_d_var = tk.StringVar()
        self._ba_hydro_d_combo = ttk.Combobox(ctrl, textvariable=self._ba_hydro_d_var,
                                               width=10, font=F['ui'], state='readonly')
        self._ba_hydro_d_combo.pack(side=tk.LEFT, padx=(4, 16))
        self._ba_hydro_d_var.trace_add('write', self._schedule_ba_hydro_plot)

        self._ba_hydro_canvas_frame = tk.Frame(self._ba_hydro_plot_container, bg=C['surface'],
                                               highlightbackground=C['border'],
                                               highlightthickness=1)
        self._ba_hydro_canvas_frame.pack(fill=tk.X, padx=10, pady=(2, 10))
        self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                 'Set pore pressure directory above to auto-plot.')
        self._tab_row_counters[cat] = row + 1
        self._refresh_ba_depth_choices()

    # ------------------------------------------------------------------
    # Plot helpers
    # ------------------------------------------------------------------

    def _clear_canvas_frame(self, frame):
        for w in frame.winfo_children():
            w.destroy()

    def _embed_figure(self, fig, frame, natural_size=False):
        canvas = FigureCanvasTkAgg(fig, master=frame)
        toolbar = _NavTB2(canvas, frame, pack_toolbar=False)
        toolbar.update()
        toolbar.pack(side=tk.BOTTOM, fill=tk.X)
        canvas.draw()
        if natural_size:
            w = int(fig.get_figwidth()  * fig.get_dpi())
            h = int(fig.get_figheight() * fig.get_dpi())
            canvas.get_tk_widget().pack(fill=tk.NONE, expand=False)
            canvas.get_tk_widget().config(width=w, height=h)
        else:
            canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        return canvas

    def _run_in_bg(self, work_fn, done_fn, cancel_attr=None):
        """Run work_fn() in a background thread; schedule done_fn(result) on the main thread.

        If cancel_attr is a string, a generation counter is kept on self under that name.
        A done_fn call is silently dropped if a newer request superseded it while it was running.
        """
        if cancel_attr:
            gen = getattr(self, cancel_attr, 0) + 1
            setattr(self, cancel_attr, gen)
            my_gen = gen
            _orig_done = done_fn
            def done_fn(result):   # noqa: F811
                if getattr(self, cancel_attr, None) == my_gen:
                    _orig_done(result)

        def _thread():
            try:
                result = work_fn()
            except Exception as exc:
                import traceback as _tb
                self._log(f'Background plot error: {exc}\n{_tb.format_exc()}\n', tag='err')
                result = None
            self.root.after(0, lambda: done_fn(result))

        threading.Thread(target=_thread, daemon=True).start()

    def _load_h5_array(self, path):
        import h5py
        with h5py.File(path, 'r') as f:
            return f['data'][:].astype(np.float32)

    @staticmethod
    def _import_load_dem():
        """
        Import load_dem from whichever analysis module is present —
        forward_analysis.py and back_analysis.py each carry their own
        identical copy (same convention as compute_slope_gradient8 and the
        other shared helpers imported elsewhere in this file), so this works
        regardless of whether the deployment is forward-only, back-analysis-
        only, or has both.
        """
        try:
            from forward_analysis import load_dem
        except ImportError:
            from back_analysis import load_dem
        return load_dem

    def _load_dem_array(self, path):
        """
        Load a DEM as a float32 array, supporting either an HDF5 file (.h5,
        via _load_h5_array) or a georeferenced raster such as a GeoTIFF
        (.tif/.tiff, via load_dem/rasterio). No-data cells are set to NaN
        either way — via the '< -900' convention for HDF5 (matching the rest
        of the app), or from the GeoTIFF's own embedded no-data value. Only
        the DEM field accepts GeoTIFF; other rasters (PGA, PGV, cohesion/phi
        3D arrays, pore pressure, theta) remain HDF5-only and must keep using
        _load_h5_array directly.
        """
        if Path(path).suffix.lower() not in ('.tif', '.tiff'):
            arr = self._load_h5_array(path)
            arr[arr < -900] = np.nan
            return arr
        arr, _geo = self._import_load_dem()(path)
        return arr

    def _dem_geo_from_tif(self, path):
        """
        (Cell_size, xmin, ymin) auto-derived from a GeoTIFF DEM's embedded
        geotransform, or None if `path` is not a .tif/.tiff (e.g. an HDF5
        DEM, which carries no georeferencing). Raises ValueError for a
        rotated/sheared or non-square-pixel GeoTIFF, same as load_dem.
        """
        if Path(path).suffix.lower() not in ('.tif', '.tiff'):
            return None
        _arr, geo = self._import_load_dem()(path)
        return geo

    def _apply_dem_geo_autofill(self, dem_key, cell_size_key, xmin_key, ymin_key):
        """
        When dem_key points to a GeoTIFF, fill cell_size_key/xmin_key/ymin_key
        from its embedded geotransform (see _dem_geo_from_tif) — this is the
        whole point of accepting a GeoTIFF DEM: Cell_size/xmin/ymin no longer
        need to be typed in by hand. No-op for an HDF5 DEM, which carries no
        georeferencing to auto-fill from, or if the file can't be read yet
        (e.g. still mid-edit while typing a path).

        Runs the actual file read in a background thread — a large regional
        GeoTIFF (or a slow/cloud-synced drive) can take long enough to read
        that doing it synchronously on the Tk main thread would freeze the
        whole GUI on every keystroke/selection of the DEM field.
        """
        dem_path = self._get_raw_value(dem_key)
        if not dem_path:
            return

        def _work():
            try:
                return self._dem_geo_from_tif(dem_path), None
            except Exception as exc:
                return None, exc

        def _done(result):
            geo, err = result if result else (None, None)
            if err is not None:
                self._log(f'Could not read GeoTIFF georeferencing from {dem_path}: {err}\n', tag='err')
                return
            if geo is None:
                return
            cell_size, xmin, ymin = geo
            for key, val in ((cell_size_key, cell_size), (xmin_key, xmin), (ymin_key, ymin)):
                t = self._vars.get(key)
                if t:
                    t[0].set(str(val))
            self._log(
                f'Auto-filled from GeoTIFF DEM: {cell_size_key}={cell_size:.6g}, '
                f'{xmin_key}={xmin:.6g}, {ymin_key}={ymin:.6g}\n', tag='info')

        self._run_in_bg(_work, _done, cancel_attr=f'_dem_geo_autofill_gen_{dem_key}')

    def _auto_fill_dem_geo(self, *_):
        self._apply_dem_geo_autofill('dem_file', 'Cell_size', 'xmin', 'ymin')

    def _auto_fill_ba_dem_geo(self, *_):
        self._apply_dem_geo_autofill('ba_dem_file', 'ba_Cell_size', 'ba_xmin', 'ba_ymin')

    def _no_mpl_error(self):
        messagebox.showerror('Missing Package',
                             'matplotlib and numpy are required for plotting.\n'
                             'Run: run_crisis.bat --reinstall')

    # ── Raster size validation ────────────────────────────────────────

    def _validate_raster_sizes(self, is_ba=False):
        """
        Read the DEM shape, then check every configured input raster.
        Returns a list of human-readable error strings (empty = all OK).
        Non-h5 files or missing files are silently skipped.
        Tolerance: forward-analysis rasters (_chk) allow ±1 row or column
        (some arrays are trimmed by one row in the physics code). Back-analysis
        pore pressure rasters (_chk_hydro_dir) require an exact match — the
        back-analysis physics code no longer silently crops a size mismatch,
        it fails loudly instead, so this check must not tolerate one either.
        """
        import h5py as _h5
        from pathlib import Path as _P
        prefix = 'ba_' if is_ba else ''

        dem_path = self._get_raw_value(f'{prefix}dem_file')
        if not dem_path:
            return []
        try:
            R, C = self._load_dem_array(dem_path).shape[:2]
        except Exception as exc:
            return [f'Cannot read DEM to validate sizes: {exc}']

        errors = []

        def _chk(key, label, kind='2d'):
            path = self._get_raw_value(key)
            if not path:
                return
            try:
                with _h5.File(path, 'r') as _f:
                    sh = _f['data'].shape
            except Exception:
                return  # not an h5 file or missing — skip silently
            if kind == '2d' or kind == '3d':
                r2, c2 = sh[0], sh[1]
                if abs(r2 - R) > 1 or abs(c2 - C) > 1:
                    errors.append(
                        f'{label}: {r2}×{c2} is incompatible with DEM {R}×{C}')
            elif kind == 'x':
                # x-coordinate vector — length should match column count
                n = sh[-1] if len(sh) >= 2 else sh[0]
                if abs(n - C) > 1:
                    errors.append(
                        f'{label}: length {n} is incompatible with DEM columns {C}')
            elif kind == 'y':
                # y-coordinate vector — length should match row count
                n = sh[0]
                if abs(n - R) > 1:
                    errors.append(
                        f'{label}: length {n} is incompatible with DEM rows {R}')

        def _chk_hydro_dir(dir_key, label, file_prefix='Pressure_head'):
            hydro_dir = self._get_raw_value(dir_key)
            if not hydro_dir:
                return
            sample = _P(hydro_dir) / f'{file_prefix}_T0Z1.h5'
            if not sample.exists():
                # try first available file
                candidates = list(_P(hydro_dir).glob(f'{file_prefix}_T*Z1.h5'))
                if not candidates:
                    return
                sample = candidates[0]
            try:
                with _h5.File(sample, 'r') as _f:
                    sh = _f['data'].shape
                r2, c2 = sh[0], sh[1]
                # Exact match required — no tolerance (see docstring above)
                if r2 != R or c2 != C:
                    errors.append(
                        f'{label}: {r2}×{c2} is incompatible with DEM {R}×{C}')
            except Exception:
                pass

        if is_ba:
            _chk_hydro_dir(f'{prefix}hydro_files_dir', 'Pore Pressure files')
            _chk_hydro_dir(f'{prefix}theta_files_dir', 'Volumetric Water Content files',
                          file_prefix='Theta')
        else:
            _chk('pga_file',              'PGA')
            _chk('pgv_file',              'PGV')
            _chk('cohesion_3d_file',      'Cohesion 3D', kind='3d')
            _chk('phi_3d_file',           'Friction Angle 3D', kind='3d')
            _chk_hydro_dir('hydro_files_dir', 'Pore Pressure files')
            _chk_hydro_dir('theta_files_dir', 'Volumetric Water Content files',
                          file_prefix='Theta')

        return errors

    # ── Per-file size check (fires on Browse) ────────────────────────

    def _schedule_size_check(self, key):
        """Debounced: wait 300 ms after the last write before checking."""
        attr = '_size_check_after'
        if hasattr(self, attr) and getattr(self, attr):
            self.root.after_cancel(getattr(self, attr))
        setattr(self, attr, self.root.after(300, lambda: self._check_file_vs_dem(key)))

    def _check_file_vs_dem(self, key):
        setattr(self, '_size_check_after', None)
        path = self._get_raw_value(key)
        if not path:
            return

        # Decide which DEM to compare against
        is_ba = key.startswith('ba_')
        dem_key = 'ba_dem_file' if is_ba else 'dem_file'
        dem_path = self._get_raw_value(dem_key)
        if not dem_path:
            return

        # Both reads are disk I/O — done in a background thread so a large
        # DEM (or a slow/cloud-synced drive) can't freeze the GUI while
        # typing/browsing a raster file.
        def _work():
            import h5py as _h5
            try:
                R, C = self._load_dem_array(dem_path).shape[:2]
            except Exception:
                return None
            try:
                with _h5.File(path, 'r') as _f:
                    sh = _f['data'].shape
            except Exception:
                return None  # not an h5 file — skip silently
            return (R, C, sh)

        def _done(result):
            if result is None:
                return
            R, C, sh = result
            label = key.replace('ba_', '').replace('_file', '').replace('_', ' ').title()

            def _reject(msg):
                messagebox.showerror('Raster Size Mismatch', msg)
                if key in self._vars:
                    self._vars[key][0].set('')
                # Immediately refresh the raster plot so the rejected raster disappears
                _FWD_RASTER_KEYS = {'dem_file', 'pga_file', 'pgv_file'}
                if key in _FWD_RASTER_KEYS:
                    self.root.after(10, self._schedule_raster_plot)

            # x-coordinate: compare to column count
            if 'x_coord' in key:
                n = sh[-1] if len(sh) >= 2 else sh[0]
                if abs(n - C) > 1:
                    _reject(f'{label}: length {n} is incompatible with DEM columns {C}.\n\n'
                            f'Expected length ≈ {C}.')
                return

            # y-coordinate: compare to row count
            if 'y_coord' in key:
                n = sh[0]
                if abs(n - R) > 1:
                    _reject(f'{label}: length {n} is incompatible with DEM rows {R}.\n\n'
                            f'Expected length ≈ {R}.')
                return

            # 2D / 3D raster: compare first two dims
            if len(sh) >= 2:
                r2, c2 = sh[0], sh[1]
                if abs(r2 - R) > 1 or abs(c2 - C) > 1:
                    _reject(f'{label}: {r2}×{c2} is incompatible with DEM {R}×{C}.\n\n'
                            f'Expected size ≈ {R}×{C}.')

        self._run_in_bg(_work, _done, cancel_attr=f'_size_check_gen_{key}')

    # ── Display downsampling helper ───────────────────────────────────

    @staticmethod
    def _ds(arr, max_px=800):
        """Subsample a 2D array for display only, using the same stride on both axes
        to preserve the true aspect ratio of the data."""
        if arr is None:
            return arr
        r, c = arr.shape[:2]
        s = max(1, max(r, c) // max_px)
        return arr[::s, ::s]

    # ── DEM mask helper ───────────────────────────────────────────────

    def _load_dem_mask(self):
        """Return boolean NaN mask from DEM (True = invalid cell). None if DEM unavailable."""
        path = self._get_raw_value('dem_file')
        if not path:
            return None
        try:
            return np.isnan(self._load_dem_array(path))
        except Exception:
            return None

    def _apply_dem_mask(self, arr, mask):
        """Set arr[mask] = NaN where mask is True. Works for 2D arrays."""
        if mask is None or arr is None:
            return arr
        arr = arr.copy()
        if arr.shape == mask.shape:
            arr[mask] = np.nan
        return arr

    # ── Plot rasters ─────────────────────────────────────────────────

    def _plot_rasters(self):
        if not _HAS_MPL:
            self._no_mpl_error()
            return

        self._clear_canvas_frame(self._raster_canvas_frame)
        self._placeholder_canvas(self._raster_canvas_frame, 'Loading rasters…')

        # Snapshot GUI state needed by the background thread
        cell_size     = self._get_raw_value('Cell_size') or 1.0
        trig_ind      = self._get_raw_value('Triggering_Indicator')
        keys = {
            'dem':   self._get_raw_value('dem_file'),
            'pga':   self._get_raw_value('pga_file') if trig_ind == 2 else None,
            'pgv':   self._get_raw_value('pgv_file') if trig_ind == 2 else None,
        }

        def _work():
            def _load(path, label):
                if not path:
                    return None
                try:
                    if label == 'DEM':
                        return self._load_dem_array(path)
                    arr = self._load_h5_array(path)
                    arr[arr < -900] = np.nan
                    return arr
                except Exception as exc:
                    self._log(f'Could not load {label}: {exc}\n', tag='err')
                    return None

            dem = _load(keys['dem'], 'DEM')
            pga = _load(keys['pga'], 'PGA') if keys['pga'] else None
            pgv = _load(keys['pgv'], 'PGV') if keys['pgv'] else None

            # Slope, Flow Direction, and Flow Accumulation are all auto-derived
            # from the DEM — no slope_file / flow_direction_file /
            # flow_accumulation_file. Flow direction/accumulation take
            # roughly a minute on a large DEM (they run in this background
            # thread, so the GUI stays responsive).
            slope = fd = area_counts = None
            if dem is not None:
                try:
                    # Imported from forward_analysis (not back_analysis) so the
                    # forward-mode preview has no dependency on back_analysis.py
                    # being present — forward_analysis.py carries its own copy
                    # of these functions for exactly this reason.
                    from forward_analysis import (
                        compute_slope_gradient8,
                        compute_flow_direction_topotoolbox,
                        compute_flow_accumulation_topotoolbox,
                    )
                    slope = compute_slope_gradient8(dem, float(cell_size), unit='degree')
                    self._log('Computing flow direction for preview (this can take ~1 minute)…\n')
                    fd = compute_flow_direction_topotoolbox(dem, float(cell_size))
                    area_counts = compute_flow_accumulation_topotoolbox(fd)
                    # Cache so _run_model can reuse these instead of
                    # recomputing Slope/Flow Direction/Flow Accumulation from
                    # scratch when the DEM file and Cell_size haven't changed
                    self._fwd_raster_cache = {
                        'dem_file': keys['dem'],
                        'Cell_size': float(cell_size),
                        'Elev_Mat': dem,
                        'Slope_Mat': slope,
                        'FD_Mat': fd,
                        'Area_Mat_counts': area_counts,
                    }
                except Exception as exc:
                    self._log(f'Could not compute Slope/Flow Direction/Flow Accumulation: {exc}\n', tag='err')
            area = area_counts

            # Size compatibility check — flag any raster that doesn't match DEM
            if dem is not None:
                _dem_r, _dem_c = dem.shape[:2]
                _size_errs = []
                for _arr, _lbl in [(slope, 'Slope'), (fd, 'Flow Direction'),
                                   (area, 'Flow Accumulation'),
                                   (pga, 'PGA'), (pgv, 'PGV')]:
                    if _arr is not None:
                        _r, _c = _arr.shape[:2]
                        if abs(_r - _dem_r) > 1 or abs(_c - _dem_c) > 1:
                            _size_errs.append(
                                f'{_lbl}: {_r}×{_c} is incompatible with DEM {_dem_r}×{_dem_c}')
                if _size_errs:
                    return ('size_error',
                            'Raster size mismatch — cannot plot:\n\n' +
                            '\n'.join(f'  •  {e}' for e in _size_errs))

            if area is not None:
                area = area * float(cell_size) ** 2
                area[area < -900] = np.nan

            loaded = [x for x in [dem, slope, fd, area, pga, pgv] if x is not None]
            if not loaded:
                return 'no_data'

            mask  = np.isnan(dem) if dem is not None else None
            slope = self._apply_dem_mask(slope, mask)
            fd    = self._apply_dem_mask(fd,    mask)
            area  = self._apply_dem_mask(area,  mask)
            pga   = self._apply_dem_mask(pga,   mask)
            pgv   = self._apply_dem_mask(pgv,   mask)

            panels = []
            if dem   is not None: panels.append((dem,   'Elevation (m)',               'terrain',  False))
            if slope is not None: panels.append((slope, 'Slope Inclination (degrees)', 'viridis',  False))
            if fd    is not None: panels.append((fd,    'Flow Direction (D8)',          'twilight', False))
            if area  is not None: panels.append((area,  'Flow Accumulation (m²)',       'plasma',   True))
            if pga   is not None: panels.append((pga,   'Peak Ground Acceleration (g)', 'hot_r',   False))
            if pgv   is not None: panels.append((pgv,   'Peak Ground Velocity (cm/s)', 'YlOrRd',   False))

            n    = len(panels)
            cols = min(n, 3)
            rows = (n + cols - 1) // cols
            fig  = Figure(figsize=(cols * 4.2, rows * 3.6), dpi=96, facecolor=C['surface'])
            for i, (data, title, cmap_name, do_log) in enumerate(panels):
                ax = fig.add_subplot(rows, cols, i + 1)
                plot_data  = np.log10(np.where(data <= 0, np.nan, data)) if do_log else data
                plot_data  = CRISISGui._ds(plot_data)
                cbar_label = (title + ' (log₁₀)') if do_log else title
                cmap = _mpl.colormaps[cmap_name].copy()
                cmap.set_bad(color='white')
                im = ax.imshow(np.ma.masked_invalid(plot_data), cmap=cmap,
                               aspect='equal', interpolation='nearest')
                cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
                cbar.set_label(cbar_label, rotation=270, labelpad=14, fontsize=8)
                ax.set_title(title, fontsize=9, pad=5, fontweight='bold')
                ax.axis('off')
            fig.tight_layout(pad=1.8)
            return fig

        def _done(result):
            self._clear_canvas_frame(self._raster_canvas_frame)
            if isinstance(result, tuple) and result[0] == 'size_error':
                self._placeholder_canvas(self._raster_canvas_frame,
                                         'Fix raster sizes above to plot.')
            elif result == 'no_data':
                self._placeholder_canvas(self._raster_canvas_frame,
                                         'Browse a raster file above to auto-plot.')
            elif result is not None:
                self._embed_figure(result, self._raster_canvas_frame)

        self._run_in_bg(_work, _done, cancel_attr='_raster_plot_gen')

    # ── BA: Plot rasters ──────────────────────────────────────────────

    def _plot_ba_rasters(self):
        if not _HAS_MPL:
            self._no_mpl_error()
            return
        if not hasattr(self, '_ba_raster_canvas_frame'):
            return

        self._clear_canvas_frame(self._ba_raster_canvas_frame)
        self._placeholder_canvas(self._ba_raster_canvas_frame, 'Loading rasters…')

        cell_size = self._get_raw_value('ba_Cell_size') or 1.0
        keys = {
            'dem': self._get_raw_value('ba_dem_file'),
        }

        def _work():
            def _load(path, label):
                if not path:
                    return None
                try:
                    return self._load_dem_array(path)
                except Exception as exc:
                    self._log(f'Could not load {label}: {exc}\n', tag='err')
                    return None

            dem = _load(keys['dem'], 'BA DEM')

            # Slope, Flow Direction, and Flow Accumulation are all auto-derived
            # from the DEM — no ba_slope_file / ba_flow_direction_file /
            # ba_flow_accumulation_file. Flow direction/accumulation take
            # roughly a minute on a large DEM (they run in this background
            # thread, so the GUI stays responsive).
            slope = fd = area_counts = None
            if dem is not None:
                try:
                    from back_analysis import (
                        compute_slope_gradient8,
                        compute_flow_direction_topotoolbox,
                        compute_flow_accumulation_topotoolbox,
                    )
                    slope = compute_slope_gradient8(dem, float(cell_size), unit='degree')
                    self._log('Computing flow direction for preview (this can take ~1 minute)…\n')
                    fd = compute_flow_direction_topotoolbox(dem, float(cell_size))
                    area_counts = compute_flow_accumulation_topotoolbox(fd)
                    # Cache so _run_back_analysis can reuse these instead of
                    # recomputing Slope/Flow Direction/Flow Accumulation from
                    # scratch when the DEM file and Cell_size haven't changed
                    self._ba_raster_cache = {
                        'dem_file': keys['dem'],
                        'Cell_size': float(cell_size),
                        'Elev_Mat': dem,
                        'Slope_Mat': slope,
                        'FD_Mat': fd,
                        'Area_Mat_counts': area_counts,
                    }
                except Exception as exc:
                    self._log(f'Could not compute Slope/Flow Direction/Flow Accumulation: {exc}\n', tag='err')
            area = area_counts

            # Size compatibility check
            if dem is not None:
                _dem_r, _dem_c = dem.shape[:2]
                _size_errs = []
                for _arr, _lbl in [(slope, 'Slope'), (fd, 'Flow Direction'),
                                   (area, 'Flow Accumulation')]:
                    if _arr is not None:
                        _r, _c = _arr.shape[:2]
                        if abs(_r - _dem_r) > 1 or abs(_c - _dem_c) > 1:
                            _size_errs.append(
                                f'{_lbl}: {_r}×{_c} is incompatible with DEM {_dem_r}×{_dem_c}')
                if _size_errs:
                    return ('size_error',
                            'Raster size mismatch — cannot plot:\n\n' +
                            '\n'.join(f'  •  {e}' for e in _size_errs))

            if area is not None:
                area = area * float(cell_size) ** 2
                area[area < -900] = np.nan

            loaded = [x for x in [dem, slope, fd, area] if x is not None]
            if not loaded:
                return None

            mask  = np.isnan(dem) if dem is not None else None
            slope = self._apply_dem_mask(slope, mask)
            fd    = self._apply_dem_mask(fd,    mask)
            area  = self._apply_dem_mask(area,  mask)

            panels = []
            if dem   is not None: panels.append((dem,   'Elevation (m)',               'terrain',  False))
            if slope is not None: panels.append((slope, 'Slope Inclination (degrees)', 'viridis',  False))
            if fd    is not None: panels.append((fd,    'Flow Direction (D8)',          'twilight', False))
            if area  is not None: panels.append((area,  'Flow Accumulation (m²)',       'plasma',   True))

            n    = len(panels)
            cols = min(n, 3)
            rows = (n + cols - 1) // cols
            fig  = Figure(figsize=(cols * 4.2, rows * 3.6), dpi=96, facecolor=C['surface'])
            for i, (data, title, cmap_name, do_log) in enumerate(panels):
                ax = fig.add_subplot(rows, cols, i + 1)
                plot_data  = np.log10(np.where(data <= 0, np.nan, data)) if do_log else data
                plot_data  = CRISISGui._ds(plot_data)
                cbar_label = (title + ' (log₁₀)') if do_log else title
                cmap = _mpl.colormaps[cmap_name].copy()
                cmap.set_bad(color='white')
                im = ax.imshow(np.ma.masked_invalid(plot_data), cmap=cmap,
                               aspect='equal', interpolation='nearest')
                cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.04)
                cbar.set_label(cbar_label, rotation=270, labelpad=14, fontsize=8)
                ax.set_title(title, fontsize=9, pad=5, fontweight='bold')
                ax.axis('off')
            fig.tight_layout(pad=1.8)
            return fig

        def _done(result):
            self._clear_canvas_frame(self._ba_raster_canvas_frame)
            if isinstance(result, tuple) and result[0] == 'size_error':
                self._placeholder_canvas(self._ba_raster_canvas_frame,
                                         'Fix raster sizes above to plot.')
            elif result is not None:
                self._embed_figure(result, self._ba_raster_canvas_frame)
            else:
                self._placeholder_canvas(self._ba_raster_canvas_frame,
                                         'Browse a raster file above to auto-plot.')

        self._run_in_bg(_work, _done, cancel_attr='_ba_raster_plot_gen')

    # ── Plot pore pressure (and optionally volumetric water content) ──

    def _plot_hydro(self, silent=False):
        if not _HAS_MPL:
            if not silent:
                self._no_mpl_error()
            return

        hydro_ind = self._get_raw_value('Hydrological_model_indicator')
        if hydro_ind != 1:
            if not silent:
                messagebox.showinfo('Constant Ru',
                                    'Pore pressure is constant (Ru model).\n'
                                    'Nothing to plot spatially.')
            return

        hydro_dir = self._get_raw_value('hydro_files_dir')
        if not hydro_dir:
            if not silent:
                messagebox.showwarning('Missing Input',
                                       'Specify the Pore Pressure Files Directory first.')
            return

        # Validate inputs on main thread before launching background work
        try:
            t = int(self._hydro_t_var.get())
        except ValueError:
            messagebox.showerror('Invalid Input', 'Time step must be an integer.')
            return

        try:
            z_min_v  = float(self._get_raw_value('z_min') or 0.125)
            z_max_v  = float(self._get_raw_value('z_max') or 1.875)
            n_d_v    = int(self._get_raw_value('Depth_points') or 8)
            depths_v = np.linspace(z_min_v, z_max_v, n_d_v)
            sel_depth = float(self._hydro_d_var.get())
            d         = int(np.argmin(np.abs(depths_v - sel_depth))) + 1
            depth_str = f'{sel_depth:.4f} m'
        except Exception:
            d         = 1
            sel_depth = 0.125
            depth_str = 'Layer 1'

        fname_p = os.path.join(hydro_dir, f'Pressure_head_T{t}Z{d}.h5')
        if not os.path.exists(fname_p):
            if not silent:
                messagebox.showerror('File Not Found',
                                     f'Could not find:\n{fname_p}\n\n'
                                     'Check the directory and the time/depth values.')
            return

        # Show loading placeholder immediately
        self._clear_canvas_frame(self._hydro_canvas_frame)
        self._placeholder_canvas(self._hydro_canvas_frame, 'Loading pore pressure data…')

        # Snapshot all GUI state needed by the background thread
        dem_path        = self._get_raw_value('dem_file')
        use_theta       = self._get_raw_value('Strength_variation_with_theta') == 1
        theta_dir       = self._get_raw_value('theta_files_dir') if use_theta else None
        cached_vrange   = getattr(self, '_hydro_vrange_dir', None)
        cached_vmin     = getattr(self, '_hydro_vmin', None)
        cached_vmax     = getattr(self, '_hydro_vmax', None)
        cached_th_dir   = getattr(self, '_theta_vrange_dir', None)
        cached_th_vmin  = getattr(self, '_theta_vmin', None)
        cached_th_vmax  = getattr(self, '_theta_vmax', None)

        def _make_dem_mask(target_shape, d_path):
            if not d_path:
                return None
            try:
                dem_arr = self._load_dem_array(d_path)
                tr, tc = target_shape
                if dem_arr.shape[0] == tr + 1:
                    dem_arr = dem_arr[:-1, :]
                mask = np.isnan(dem_arr)
                mr, mc = mask.shape
                if mr != tr:
                    mask = mask[:tr, :] if mr > tr else np.pad(mask, ((0, tr - mr), (0, 0)), constant_values=True)
                if mc != tc:
                    mask = mask[:, :tc] if mc > tc else np.pad(mask, ((0, 0), (0, tc - mc)), constant_values=True)
                return mask
            except Exception:
                return None

        def _work():
            # Load pressure file
            try:
                arr_p = self._load_h5_array(fname_p)
            except Exception as exc:
                return ('load_error', str(exc))

            # Size check against DEM
            if dem_path:
                try:
                    import h5py as _h5c
                    with _h5c.File(dem_path, 'r') as _f:
                        _dr, _dc = _f['data'].shape[:2]
                    _pr, _pc = arr_p.shape[:2]
                    if abs(_pr - _dr) > 1 or abs(_pc - _dc) > 1:
                        return ('size_error',
                                f'Pore Pressure file: {_pr}×{_pc} is incompatible '
                                f'with DEM {_dr}×{_dc}')
                except Exception:
                    pass

            try:
                arr_p = np.minimum(arr_p, float(sel_depth))
            except Exception:
                pass
            dem_mask_p = _make_dem_mask(arr_p.shape, dem_path)
            if dem_mask_p is not None:
                arr_p[dem_mask_p] = np.nan

            # Optionally load theta
            arr_th = None
            if use_theta and theta_dir:
                fname_th = os.path.join(theta_dir, f'Theta_T{t}Z{d}.h5')
                if os.path.exists(fname_th):
                    try:
                        arr_th = self._load_h5_array(fname_th)
                        dem_mask_th = _make_dem_mask(arr_th.shape, dem_path)
                        if dem_mask_th is not None:
                            arr_th[dem_mask_th] = np.nan
                    except Exception:
                        arr_th = None

            # Global colour scale — compute if not cached for this directory
            nonlocal cached_vrange, cached_vmin, cached_vmax
            nonlocal cached_th_dir, cached_th_vmin, cached_th_vmax
            if cached_vrange != hydro_dir:
                cached_vmin, cached_vmax = self._compute_hydro_vrange(hydro_dir)
                cached_vrange = hydro_dir
                # Store back so next call hits the cache
                self._hydro_vrange_dir = hydro_dir
                self._hydro_vmin       = cached_vmin
                self._hydro_vmax       = cached_vmax

            vmin_th = vmax_th = None
            if arr_th is not None and theta_dir:
                if cached_th_dir != theta_dir:
                    cached_th_vmin, cached_th_vmax = self._compute_theta_vrange(theta_dir)
                    cached_th_dir = theta_dir
                    self._theta_vrange_dir = theta_dir
                    self._theta_vmin       = cached_th_vmin
                    self._theta_vmax       = cached_th_vmax
                vmin_th, vmax_th = cached_th_vmin, cached_th_vmax

            # Build matplotlib figure — figure sized to match raster aspect ratio
            blues = _mpl.colormaps['Blues'].copy()
            blues.set_bad(color='white')
            ncols  = 2 if arr_th is not None else 1
            _ar_p  = arr_p.shape[0] / max(arr_p.shape[1], 1)
            _pw_p  = 3.2
            _ph_p  = _pw_p * _ar_p + 0.6   # +0.6 for title and colorbar label
            fig    = Figure(figsize=(ncols * _pw_p, _ph_p), dpi=96, facecolor=C['surface'])

            ax1 = fig.add_subplot(1, ncols, 1)
            im1 = ax1.imshow(np.ma.masked_invalid(CRISISGui._ds(arr_p)), cmap=blues,
                             aspect='equal', interpolation='nearest',
                             vmin=cached_vmin, vmax=cached_vmax)
            fig.colorbar(im1, ax=ax1, fraction=0.046, pad=0.04, label='Pressure Head (m)')
            ax1.set_title(f'Pore Pressure Head — Time {t},  Depth = {depth_str}',
                          fontsize=10, pad=6, fontweight='bold')
            ax1.axis('off')

            if arr_th is not None:
                ax2 = fig.add_subplot(1, 2, 2)
                im2 = ax2.imshow(np.ma.masked_invalid(CRISISGui._ds(arr_th)), cmap=blues,
                                 aspect='equal', interpolation='nearest',
                                 vmin=vmin_th, vmax=vmax_th)
                fig.colorbar(im2, ax=ax2, fraction=0.046, pad=0.04,
                             label='Volumetric Water Content θ (m³/m³)')
                ax2.set_title(f'Volumetric Water Content θ — Time {t},  Depth = {depth_str}',
                              fontsize=10, pad=6, fontweight='bold')
                ax2.axis('off')

            fig.tight_layout(pad=1.5)
            return ('ok', fig)

        def _done(result):
            self._clear_canvas_frame(self._hydro_canvas_frame)
            if result is None:
                self._placeholder_canvas(self._hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
                return
            kind, payload = result
            if kind == 'size_error':
                self._placeholder_canvas(self._hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
            elif kind == 'load_error':
                messagebox.showerror('Load Error', payload)
                self._placeholder_canvas(self._hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
            else:
                self._embed_figure(payload, self._hydro_canvas_frame)

        self._run_in_bg(_work, _done, cancel_attr='_hydro_plot_gen')

    # ── BA: Plot pore pressure (and optionally volumetric water content) ──

    def _plot_ba_hydro(self, silent=False):
        if not _HAS_MPL:
            if not silent:
                self._no_mpl_error()
            return
        if not hasattr(self, '_ba_hydro_canvas_frame'):
            return

        hydro_dir = self._get_raw_value('ba_hydro_files_dir')
        if not hydro_dir:
            if not silent:
                messagebox.showwarning('Missing Input',
                                       'Specify the BA Pore Pressure Files Directory first.')
            return

        try:
            t = int(self._ba_hydro_t_var.get())
        except ValueError:
            messagebox.showerror('Invalid Input', 'Time step must be an integer.')
            return

        try:
            z_min_v  = float(self._get_raw_value('ba_z_min') or 0.125)
            z_max_v  = float(self._get_raw_value('ba_z_max') or 1.875)
            n_d_v    = int(self._get_raw_value('ba_Depth_points') or 8)
            depths_v = np.linspace(z_min_v, z_max_v, n_d_v)
            sel_depth = float(self._ba_hydro_d_var.get())
            d         = int(np.argmin(np.abs(depths_v - sel_depth))) + 1
            depth_str = f'{sel_depth:.4f} m'
        except Exception:
            d         = 1
            sel_depth = 0.125
            depth_str = 'Layer 1'

        fname_p = os.path.join(hydro_dir, f'Pressure_head_T{t}Z{d}.h5')
        if not os.path.exists(fname_p):
            if not silent:
                messagebox.showerror('File Not Found',
                                     f'Could not find:\n{fname_p}\n\n'
                                     'Check the directory and the time/depth values.')
            return

        self._clear_canvas_frame(self._ba_hydro_canvas_frame)
        self._placeholder_canvas(self._ba_hydro_canvas_frame, 'Loading pore pressure data…')

        dem_path       = self._get_raw_value('ba_dem_file')
        use_theta      = self._get_raw_value('ba_Strength_variation_with_theta') == 1
        theta_dir      = self._get_raw_value('ba_theta_files_dir') if use_theta else None
        cached_vrange  = getattr(self, '_ba_hydro_vrange_dir', None)
        cached_vmin    = getattr(self, '_ba_hydro_vmin', None)
        cached_vmax    = getattr(self, '_ba_hydro_vmax', None)
        cached_th_dir  = getattr(self, '_ba_theta_vrange_dir', None)
        cached_th_vmin = getattr(self, '_ba_theta_vmin', None)
        cached_th_vmax = getattr(self, '_ba_theta_vmax', None)

        def _make_ba_dem_mask(target_shape, d_path):
            if not d_path:
                return None
            try:
                dem_arr = self._load_dem_array(d_path)
                tr, tc = target_shape
                if dem_arr.shape[0] == tr + 1:
                    dem_arr = dem_arr[:-1, :]
                mask = np.isnan(dem_arr)
                mr, mc = mask.shape
                if mr != tr:
                    mask = mask[:tr, :] if mr > tr else np.pad(mask, ((0, tr - mr), (0, 0)), constant_values=True)
                if mc != tc:
                    mask = mask[:, :tc] if mc > tc else np.pad(mask, ((0, 0), (0, tc - mc)), constant_values=True)
                return mask
            except Exception:
                return None

        def _work():
            try:
                arr_p = self._load_h5_array(fname_p)
            except Exception as exc:
                return ('load_error', str(exc))

            # Size check against BA DEM
            if dem_path:
                try:
                    import h5py as _h5c
                    with _h5c.File(dem_path, 'r') as _f:
                        _dr, _dc = _f['data'].shape[:2]
                    _pr, _pc = arr_p.shape[:2]
                    if abs(_pr - _dr) > 1 or abs(_pc - _dc) > 1:
                        return ('size_error',
                                f'Pore Pressure file: {_pr}×{_pc} is incompatible '
                                f'with DEM {_dr}×{_dc}')
                except Exception:
                    pass

            try:
                arr_p = np.minimum(arr_p, float(sel_depth))
            except Exception:
                pass
            dem_mask_p = _make_ba_dem_mask(arr_p.shape, dem_path)
            if dem_mask_p is not None:
                arr_p[dem_mask_p] = np.nan

            arr_th = None
            if use_theta and theta_dir:
                fname_th = os.path.join(theta_dir, f'Theta_T{t}Z{d}.h5')
                if os.path.exists(fname_th):
                    try:
                        arr_th = self._load_h5_array(fname_th)
                        dem_mask_th = _make_ba_dem_mask(arr_th.shape, dem_path)
                        if dem_mask_th is not None:
                            arr_th[dem_mask_th] = np.nan
                    except Exception:
                        arr_th = None

            nonlocal cached_vrange, cached_vmin, cached_vmax
            nonlocal cached_th_dir, cached_th_vmin, cached_th_vmax
            if cached_vrange != hydro_dir:
                cached_vmin, cached_vmax = self._compute_ba_hydro_vrange(hydro_dir)
                cached_vrange = hydro_dir
                self._ba_hydro_vrange_dir = hydro_dir
                self._ba_hydro_vmin       = cached_vmin
                self._ba_hydro_vmax       = cached_vmax

            vmin_th = vmax_th = None
            if arr_th is not None and theta_dir:
                if cached_th_dir != theta_dir:
                    cached_th_vmin, cached_th_vmax = self._compute_ba_theta_vrange(theta_dir)
                    cached_th_dir = theta_dir
                    self._ba_theta_vrange_dir = theta_dir
                    self._ba_theta_vmin       = cached_th_vmin
                    self._ba_theta_vmax       = cached_th_vmax
                vmin_th, vmax_th = cached_th_vmin, cached_th_vmax

            blues = _mpl.colormaps['Blues'].copy()
            blues.set_bad(color='white')
            ncols = 2 if arr_th is not None else 1
            fig   = Figure(figsize=(6 * ncols, 4.5), dpi=96, facecolor=C['surface'])

            ax1 = fig.add_subplot(1, ncols, 1)
            im1 = ax1.imshow(np.ma.masked_invalid(CRISISGui._ds(arr_p)), cmap=blues,
                             aspect='equal', interpolation='nearest',
                             vmin=cached_vmin, vmax=cached_vmax)
            fig.colorbar(im1, ax=ax1, fraction=0.046, pad=0.04, label='Pressure Head (m)')
            ax1.set_title(f'Pore Pressure Head — Time {t},  Depth = {depth_str}',
                          fontsize=10, pad=6, fontweight='bold')
            ax1.axis('off')

            if arr_th is not None:
                ax2 = fig.add_subplot(1, 2, 2)
                im2 = ax2.imshow(np.ma.masked_invalid(CRISISGui._ds(arr_th)), cmap=blues,
                                 aspect='equal', interpolation='nearest',
                                 vmin=vmin_th, vmax=vmax_th)
                fig.colorbar(im2, ax=ax2, fraction=0.046, pad=0.04,
                             label='Volumetric Water Content θ (m³/m³)')
                ax2.set_title(f'Volumetric Water Content θ — Time {t},  Depth = {depth_str}',
                              fontsize=10, pad=6, fontweight='bold')
                ax2.axis('off')

            fig.tight_layout(pad=1.5)
            return ('ok', fig)

        def _done(result):
            self._clear_canvas_frame(self._ba_hydro_canvas_frame)
            if result is None:
                self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
                return
            kind, payload = result
            if kind == 'size_error':
                self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
            elif kind == 'load_error':
                messagebox.showerror('Load Error', payload)
                self._placeholder_canvas(self._ba_hydro_canvas_frame,
                                         'Set pore pressure directory above to auto-plot.')
            else:
                self._embed_figure(payload, self._ba_hydro_canvas_frame)

        self._run_in_bg(_work, _done, cancel_attr='_ba_hydro_plot_gen')

    # ── BA Outputs: summary generation + plots ───────────────────────

    def _run_ba_outputs(self):
        if not _HAS_MPL:
            return
        base  = self._get_raw_value('output_dir') or ''
        title = (self._get_raw_value('project_title') or 'CRISIS').strip() or 'CRISIS'
        out_dir = str(Path(base) / f'{title}_Outputs') if base else ''
        if not out_dir or not os.path.isdir(out_dir):
            return   # folder not ready yet — silently skip

        # Snapshot GUI state needed by the background thread (reading Tk
        # variables must happen on the main thread — see _plot_rasters above).
        ba_ctx = dict(
            dem_path        = self._get_raw_value('ba_dem_file'),
            mapped_shp_path = self._get_raw_value('ba_mapped_ls_shapefile') or '',
            cell_size       = self._get_raw_value('ba_Cell_size'),
            xmin_dem        = self._get_raw_value('ba_xmin'),
            ymin_dem        = self._get_raw_value('ba_ymin'),
        )

        def _work():
            try:
                summary_df, excluded = self._generate_ba_summary(out_dir)
            except Exception as exc:
                return ('summary_error', str(exc))
            if summary_df is None or len(summary_df) == 0:
                return ('no_data', None)
            try:
                fig = self._plot_ba_outputs(summary_df, out_dir, ba_ctx)
            except Exception as exc:
                return ('plot_error', str(exc))
            return ('ok', (fig, excluded, out_dir))

        def _done(result):
            kind, payload = result
            if kind == 'summary_error':
                self._log(f'Outputs summary error: {payload}\n', tag='err')
            elif kind == 'no_data':
                pass   # no LD files yet — nothing to summarize
            elif kind == 'plot_error':
                self._log(f'Outputs plot error: {payload}\n', tag='err')
            else:
                fig, excluded, out_dir = payload
                if excluded:
                    self._log(f'Excluded {len(excluded)} landslide(s) (non-positive effective normal stress).\n',
                              tag='dim')
                self._log(f'Outputs saved to: {out_dir}\n', tag='ok')
                self._log(f'Summary saved → {out_dir}\\Back_Calculated_Landslides_Summary.xlsx\n',
                          tag='ok')
                self._clear_canvas_frame(self._ba_outputs_canvas_frame)
                self._embed_figure(fig, self._ba_outputs_canvas_frame)
                try:
                    import matplotlib.pyplot as plt
                    plt.close(fig)
                except Exception:
                    pass
            # Logged last, after the Excel summary/plot are actually done —
            # otherwise the user sees "completed successfully" while the
            # summary file is still being written in the background.
            self._log('\n✓  Back-analysis completed successfully.\n', tag='ok')

        self._run_in_bg(_work, _done)

    @staticmethod
    def _generate_ba_summary(out_dir):
        """Generate Back_Calculated_Landslides_Summary.xlsx from all LD xlsx files.

        back_analysis.run_model() now writes this same workbook itself right
        after a run (so the CLI/HPC path gets a summary too, not just the
        GUI). If it's already there, read it back instead of re-scanning and
        recomputing from every acceptable_combinations_for_LD_*.xlsx file —
        this function still does the full recompute as a fallback, e.g. when
        viewing Outputs from an older run that predates that change."""
        import re
        import pandas as pd
        out_path = Path(out_dir)
        summary_path = out_path / 'Back_Calculated_Landslides_Summary.xlsx'
        if summary_path.exists():
            try:
                summary_df = pd.read_excel(summary_path, sheet_name='Sheet1')
                try:
                    excluded_df = pd.read_excel(summary_path, sheet_name='Excluded_Landslides')
                    excluded = excluded_df['Excluded_Landslide_Number'].tolist()
                except Exception:
                    excluded = []
                if len(summary_df) > 0:
                    return summary_df, excluded
            except Exception:
                pass   # fall through to full recompute below

        pattern = re.compile(r'^acceptable_combinations_for_LD_(\d+)\.xlsx$', re.IGNORECASE)
        ld_files = sorted(
            [(int(m.group(1)), f) for f in out_path.iterdir()
             if (m := pattern.match(f.name))],
            key=lambda x: x[0]
        )
        rows, excluded = [], []
        for num, fpath in ld_files:
            try:
                df_ld = pd.read_excel(fpath, header=0)
            except Exception:
                continue
            M = df_ld.values
            if M.shape[1] < 21:
                continue
            Total_error  = M[:, 0].astype(float)
            CD_error     = M[:, 1].astype(float)
            Area_error   = M[:, 2].astype(float)
            Volume_error = M[:, 3].astype(float)
            Phi          = M[:, 4].astype(float)
            cohesion     = M[:, 5].astype(float)
            Area_modeled = M[:, 7].astype(float)
            Vol_modeled  = M[:, 8].astype(float)
            Trig_row     = M[:, 14].astype(float)
            Trig_col     = M[:, 15].astype(float)
            Trig_Slope   = M[:, 16].astype(float)
            Trig_Depths  = M[:, 17].astype(float)
            ENS          = M[:, -2].astype(float)  # Effective_Normal_Stress
            SS           = M[:, -1].astype(float)  # Shear_Strength
            pos = ENS > 0
            if not np.any(pos):
                excluded.append(num)
                continue
            ENS = ENS[pos]; SS = SS[pos]; cohesion = cohesion[pos]
            Phi = Phi[pos]; Trig_Depths = Trig_Depths[pos]
            Trig_Slope = Trig_Slope[pos]; Trig_row = Trig_row[pos]
            Trig_col = Trig_col[pos]; Total_error = Total_error[pos]
            CD_error = CD_error[pos]; Area_error = Area_error[pos]
            Volume_error = Volume_error[pos]
            Area_modeled = Area_modeled[pos]; Vol_modeled = Vol_modeled[pos]
            mapped_area = Area_modeled / (1 + Area_error)
            mapped_vol  = Vol_modeled  / (1 + Volume_error)
            avg_depth   = mapped_vol / mapped_area
            rows.append([
                num, mapped_area[0], mapped_vol[0],
                np.mean(ENS), np.std(ENS, ddof=0),
                np.mean(SS),  np.std(SS,  ddof=0),
                np.mean(SS) / np.mean(ENS),
                np.mean(Trig_Slope),
                np.mean(cohesion), np.std(cohesion, ddof=0),
                np.mean(Phi),      np.std(Phi,      ddof=0),
                np.mean(Trig_Depths), np.std(Trig_Depths, ddof=0),
                Total_error[0], CD_error[0], Area_error[0], Volume_error[0],
                avg_depth[0], Trig_row[0], Trig_col[0],
            ])
        if not rows:
            return None, excluded
        cols = [
            'Landslide Number', 'Mapped Area (m^2)', 'Mapped Volume (m^3)',
            'Average Effective Normal Stress (kPa)', 'Effective Normal Stress SD (kPa)',
            'Average Shear Strength (kPa)', 'Shear Strength SD (kPa)',
            'Normalized Shear Strength (unitless)', 'Slope of Triggering Cell (deg)',
            'Average cohesion (kPa)', 'cohesion SD (kPa)', 'Average Phi (deg)', 'Phi SD (deg)',
            'Average Triggering depth (m)', 'Triggering depth SD (m)',
            'Total error (unitless)', 'CD error (unitless)', 'Area error (unitless)', 'Volume error (unitless)',
            'Average Mapped depth (m)',
            'Row of triggering cell in main DEM (unitless)', 'Col of triggering cell in main DEM (unitless)',
        ]
        summary_df = pd.DataFrame(rows, columns=cols)
        summary_path = out_path / 'Back_Calculated_Landslides_Summary.xlsx'
        with pd.ExcelWriter(str(summary_path), engine='openpyxl') as writer:
            summary_df.to_excel(writer, sheet_name='Sheet1', index=False)
            if excluded:
                pd.DataFrame({'Excluded_Landslide_Number': excluded}).to_excel(
                    writer, sheet_name='Excluded_Landslides', index=False)
            # Auto-fit column widths to header text on every sheet
            for sheet in writer.sheets.values():
                for col_cells in sheet.iter_cols():
                    header = col_cells[0].value  # first row is the header
                    if header is not None:
                        col_cells[0].column_letter  # trigger dimension caching
                        sheet.column_dimensions[col_cells[0].column_letter].width = \
                            max(len(str(header)) + 2, 10)
        return summary_df, excluded

    def _plot_ba_outputs(self, df, out_dir='', ba_ctx=None):
        """Build the 4-histogram + Mohr-Coulomb figure and return it (does not
        touch any Tk widgets — safe to call from a background thread;
        the caller embeds/closes the returned figure on the main thread)."""
        ba_ctx = ba_ctx or {}
        import matplotlib.pyplot as plt
        import matplotlib.gridspec as gridspec

        # ── Extract columns ──────────────────────────────────────────────
        total_err  = df['Total error (unitless)'].dropna().values
        cd_err     = df['CD error (unitless)'].dropna().values
        area_err   = df['Area error (unitless)'].dropna().values
        vol_err    = df['Volume error (unitless)'].dropna().values
        ens        = df['Average Effective Normal Stress (kPa)'].dropna().values
        ss         = df['Average Shear Strength (kPa)'].dropna().values
        ss_sd      = df['Shear Strength SD (kPa)'].dropna().values
        # Align ens/ss/ss_sd to common valid index
        mc_mask = np.isfinite(ens) & np.isfinite(ss) & np.isfinite(ss_sd)
        ens, ss, ss_sd = ens[mc_mask], ss[mc_mask], ss_sd[mc_mask]

        # ── Figure layout ────────────────────────────────────────────────
        # Histograms (left) | DEM (middle-left) | spacer | Mohr-Coulomb (right)
        fig = plt.figure(figsize=(22, 9), facecolor='white')
        # Col 0: histogram block, col 1: DEM, col 2: spacer, col 3: MC
        gs_outer = gridspec.GridSpec(1, 4, figure=fig,
                                     left=0.02, right=0.98,
                                     top=0.92, bottom=0.20,
                                     wspace=0.08,
                                     width_ratios=[1.7, 1, 0.30, 1])
        gs_hist = gridspec.GridSpecFromSubplotSpec(2, 2, subplot_spec=gs_outer[0],
                                                   wspace=0.08, hspace=0.38)
        ax1  = fig.add_subplot(gs_hist[0, 0])   # Total error
        ax2  = fig.add_subplot(gs_hist[0, 1])   # Area error
        ax3  = fig.add_subplot(gs_hist[1, 0])   # Volume error
        ax4  = fig.add_subplot(gs_hist[1, 1])   # Location (CD) error
        ax_d = fig.add_subplot(gs_outer[1])      # DEM (middle)
        ax5  = fig.add_subplot(gs_outer[3])      # Mohr-Coulomb (right)

        bar_color = (0.3010, 0.7450, 0.9330)
        font_name = 'DejaVu Serif'   # closest to Times New Roman available cross-platform

        def _style(ax):
            ax.set_box_aspect(1)
            for spine in ax.spines.values():
                spine.set_linewidth(1.5)
            ax.tick_params(length=5, width=1.2, labelsize=10)

        def _hist(ax, data, bins, xlabel, label, zero_line=False):
            counts, edges, patches = ax.hist(
                data, bins=bins, color=bar_color, alpha=0.6,
                edgecolor='black', linewidth=0.8)
            _style(ax)
            ax.set_xlabel(xlabel, fontsize=9, fontweight='bold', fontfamily=font_name)
            ax.set_ylabel('Frequency', fontsize=9, fontweight='bold', fontfamily=font_name)
            ax.set_xlim(edges[0], edges[-1])
            if counts.max() > 0:
                ax.set_ylim(0, counts.max() + max(counts.max() * 0.12, 2))
            if zero_line and counts.max() > 0:
                ax.axvline(0, color='red', linewidth=2)
            ax.text(0.95, 0.92, label, transform=ax.transAxes,
                    ha='right', va='top', fontsize=11,
                    fontweight='bold', fontfamily=font_name)

        _hist(ax1, total_err,
              bins=np.arange(0, 1.75, 0.25),
              xlabel='Total Error', label='(a)')
        _hist(ax2, area_err,
              bins=np.arange(-1, 1.25, 0.25),
              xlabel='Relative Area Error', label='(b)', zero_line=True)
        _hist(ax3, vol_err,
              bins=np.arange(-1, 1.25, 0.25),
              xlabel='Relative Volume Error', label='(c)', zero_line=True)
        _hist(ax4, cd_err,
              bins=np.arange(0, 1.75, 0.25),
              xlabel='Location Error', label='(d)')

        # ── DEM plot with landslide overlays (middle) ─────────────────────
        import geopandas as gpd

        dem_path        = ba_ctx.get('dem_path')
        mapped_shp_path = ba_ctx.get('mapped_shp_path') or ''
        try:
            cell_size = float(ba_ctx.get('cell_size') or 1.0)
            xmin_dem  = float(ba_ctx.get('xmin_dem') or 0.0)
            ymin_dem  = float(ba_ctx.get('ymin_dem') or 0.0)
        except Exception:
            cell_size, xmin_dem, ymin_dem = 1.0, 0.0, 0.0

        if dem_path:
            try:
                dem = self._load_dem_array(dem_path).astype(float)
                if dem.shape[0] > 1:
                    dem = dem[:-1, :]
                n_rows, n_cols = dem.shape
                xmax_dem = xmin_dem + n_cols * cell_size
                ymax_dem = ymin_dem + n_rows * cell_size
                extent   = [xmin_dem, xmax_dem, ymin_dem, ymax_dem]
                im = ax_d.imshow(self._ds(dem), cmap='terrain', aspect='auto',
                                 origin='upper', extent=extent,
                                 interpolation='nearest')
                cbar = fig.colorbar(im, ax=ax_d, fraction=0.046, pad=0.04)
                cbar.set_label('Elevation (m)', fontsize=9,
                               fontweight='bold', fontfamily=font_name)
                cbar.ax.tick_params(labelsize=8)
            except Exception:
                ax_d.text(0.5, 0.5, 'DEM unavailable', transform=ax_d.transAxes,
                          ha='center', va='center', fontsize=10,
                          fontfamily=font_name, color='grey')
        else:
            ax_d.text(0.5, 0.5, 'No DEM file\nspecified', transform=ax_d.transAxes,
                      ha='center', va='center', fontsize=10,
                      fontfamily=font_name, color='grey')

        import matplotlib.patches as mpatches
        _legend_patches = []
        gdf_modeled = None
        gdf_mapped  = None

        # Overlay modeled landslides first — drawn manually for reliable zorder
        if out_dir and os.path.isdir(out_dir):
            modeled_shps = list(Path(out_dir).glob('Predicted_landslides_time_*.shp'))
            if modeled_shps:
                try:
                    gdfs = [gpd.read_file(str(s)) for s in modeled_shps]
                    import pandas as _pd
                    gdf_modeled = gpd.GeoDataFrame(
                        _pd.concat(gdfs, ignore_index=True),
                        crs=gdfs[0].crs) if len(gdfs) > 1 else gdfs[0]
                    # Batched draw (one collection instead of one ax.fill()
                    # call per polygon) — several times faster for hundreds
                    # of landslides.
                    gdf_modeled.plot(ax=ax_d, facecolor='red', edgecolor='red',
                                     linewidth=0.6, alpha=0.7, zorder=2)
                    _legend_patches.append(
                        mpatches.Patch(facecolor='red', edgecolor='red',
                                       linewidth=0.8, label='Modeled landslides'))
                except Exception as _e:
                    self._log(f'Modeled landslides overlay error: {_e}\n', tag='err')

        # Overlay mapped landslides on top — black outline only, drawn last
        if mapped_shp_path and os.path.isfile(mapped_shp_path):
            try:
                gdf_mapped = gpd.read_file(mapped_shp_path)
                # Batched draw — see comment above.
                gdf_mapped.boundary.plot(ax=ax_d, color='black', linewidth=0.8, zorder=20)
                _legend_patches.append(
                    mpatches.Patch(facecolor='none', edgecolor='black',
                                   linewidth=0.8, label='Mapped landslides'))
            except Exception as _e:
                self._log(f'Mapped landslides overlay error: {_e}\n', tag='err')

        # Legend below DEM
        if _legend_patches:
            ax_d.legend(handles=_legend_patches,
                        loc='upper center', bbox_to_anchor=(0.5, -0.09),
                        ncol=len(_legend_patches), frameon=False,
                        fontsize=7, prop={'family': font_name})

        # North arrow — above the colorbar (right side of DEM axes), clip_on=False
        if dem_path:
            # x=1.07 places it over the colorbar strip; y slightly below previous
            ax_d.annotate('', xy=(1.07, 1.04), xytext=(1.07, 0.96),
                          xycoords='axes fraction', textcoords='axes fraction',
                          annotation_clip=False,
                          arrowprops=dict(arrowstyle='->', color='black', lw=1.5))
            ax_d.text(1.07, 1.05, 'N', transform=ax_d.transAxes,
                      ha='center', va='bottom', fontsize=9,
                      fontweight='bold', fontfamily=font_name, clip_on=False)

            # Scale bar — below the DEM axes, above the legend
            # Compute real-world length from data coordinates
            _xl, _xr = ax_d.get_xlim()
            _bar_m   = (_xr - _xl) * 0.20          # metres for 20 % of x-span
            _bar_label = f'{_bar_m:.0f} m'
            # Draw in axes-fraction space so it sits outside the DEM image
            _bx0, _bx1 = 0.05, 0.25                # 5 %–25 % of axes width
            _by        = -0.02                       # just below the axes bottom
            _th        = 0.012                      # tick half-height in fraction
            ax_d.plot([_bx0, _bx1], [_by, _by],
                      transform=ax_d.transAxes, color='black', lw=1.5,
                      clip_on=False, zorder=10)
            for _bx in (_bx0, _bx1):
                ax_d.plot([_bx, _bx], [_by - _th, _by + _th],
                          transform=ax_d.transAxes, color='black', lw=1.2,
                          clip_on=False, zorder=10)
            ax_d.text((_bx0 + _bx1) / 2, _by - _th * 2,
                      _bar_label, transform=ax_d.transAxes,
                      ha='center', va='top', fontsize=6,
                      fontfamily=font_name, clip_on=False, zorder=10)

        ax_d.set_xticks([])
        ax_d.set_yticks([])
        for spine in ax_d.spines.values():
            spine.set_visible(False)

        # ── Click-to-zoom on DEM ─────────────────────────────────────────
        # Capture variables needed by the handler. Reuse the gdf_modeled /
        # gdf_mapped GeoDataFrames already parsed above instead of re-reading
        # hundreds of shapefiles from disk on every click (that re-read was
        # the main source of the multi-second lag on each zoom click).
        # The DEM is downsampled once here too, rather than per click.
        _dem_data      = CRISISGui._ds(dem) if dem_path else None
        _dem_extent    = extent if dem_path else None
        _mods_gdf      = gdf_modeled
        _mapped_gdf    = gdf_mapped
        _zoom_radius   = 500   # half-width in map units (metres) for initial zoom
        # Holds the persistent zoom figure/axes across clicks (built once,
        # then just re-centered) — see _on_dem_click below.
        _zoom_state    = {'fig': None, 'ax': None}

        def _on_dem_click(event, _dem=_dem_data, _ext=_dem_extent,
                          _mgdf=_mapped_gdf, _modgdf=_mods_gdf,
                          _font=font_name):
            if event.inaxes is not ax_d or event.button != 1:
                return
            cx, cy = event.xdata, event.ydata
            if cx is None or cy is None:
                return

            import matplotlib.pyplot as _plt
            r = _zoom_radius

            zfig, zax = _zoom_state['fig'], _zoom_state['ax']
            zoom_window_alive = zfig is not None and _plt.fignum_exists(zfig.number)

            if not zoom_window_alive:
                # (Re)build the zoom window and its static content once — the
                # DEM image and landslide overlays never change between
                # clicks, only which area is in view, so this only needs to
                # run on the very first click (or after the user closes the
                # zoom window and clicks again).
                zfig, zax = _plt.subplots(figsize=(8, 8), facecolor='white')
                zfig.canvas.manager.set_window_title('DEM Zoom')

                if _dem is not None and _ext is not None:
                    zax.imshow(_dem, cmap='terrain', aspect='equal',
                               origin='upper', extent=_ext, interpolation='nearest')

                # Modeled landslides first (behind) — reuse already-parsed
                # geometries, drawn as one batched collection instead of one
                # ax.fill() call per polygon (several times faster for
                # hundreds of landslides — this used to be most of the delay
                # on the first zoom click).
                if _modgdf is not None and len(_modgdf) > 0:
                    try:
                        _modgdf.plot(ax=zax, facecolor='red', edgecolor='red',
                                     linewidth=0.8, alpha=0.7, zorder=2)
                    except Exception:
                        pass

                # Mapped landslides on top (drawn last, highest zorder) — reused too
                if _mgdf is not None and len(_mgdf) > 0:
                    try:
                        _mgdf.boundary.plot(ax=zax, color='black', linewidth=1.0, zorder=20)
                    except Exception:
                        pass

                zax.set_xticks([])
                zax.set_yticks([])
                for _sp in zax.spines.values():
                    _sp.set_visible(False)

                # Scroll to zoom
                def _on_scroll(ev, _a=zax, _f=zfig):
                    factor = 0.7 if ev.button == 'up' else 1.4
                    xl, xr = _a.get_xlim(); yb, yt = _a.get_ylim()
                    mx, my = ev.xdata or (xl+xr)/2, ev.ydata or (yb+yt)/2
                    _a.set_xlim(mx + (xl-mx)*factor, mx + (xr-mx)*factor)
                    _a.set_ylim(my + (yb-my)*factor, my + (yt-my)*factor)
                    _f.canvas.draw_idle()

                zfig.canvas.mpl_connect('scroll_event', _on_scroll)
                _plt.tight_layout(pad=0.5)
                _plt.show(block=False)
                try:
                    zfig.canvas.manager.window.wm_geometry('+200+30')
                except Exception:
                    pass
                _zoom_state['fig'], _zoom_state['ax'] = zfig, zax

            # Re-center the existing window on the newly clicked point — cheap
            # (no image/polygon redraw needed) so repeated clicks feel instant.
            zax.set_xlim(cx - r, cx + r)
            zax.set_ylim(cy - r, cy + r)
            zfig.canvas.draw_idle()
            try:
                zfig.canvas.manager.window.lift()
            except Exception:
                pass

        fig.canvas.mpl_connect('button_press_event', _on_dem_click)
        ax_d.text(0.5, -0.18, 'Click DEM to zoom  |  scroll to zoom in/out',
                  transform=ax_d.transAxes, ha='center', va='top',
                  fontsize=9, fontstyle='italic', color='grey',
                  fontfamily=font_name)

        # ── Mohr-Coulomb plot ────────────────────────────────────────────
        ax5.set_title('Mohr Coulomb Failure Envelope',
                      fontsize=12, fontweight='bold', fontfamily=font_name, pad=32)
        ax5.text(0.5, 1.02, 'Each dot represents a mapped landslide',
                 transform=ax5.transAxes,
                 ha='center', va='bottom',
                 fontsize=8, fontstyle='italic', fontfamily=font_name)
        legend_handles = []
        if len(ens) >= 2:
            # Replicate Statistics_of_linear_regression.m exactly
            n     = len(ens)
            X_avg = ens.mean();  Y_avg = ss.mean()
            X_SD  = np.sqrt(np.sum((ens - X_avg)**2) / n)
            Y_SD  = np.sqrt(np.sum((ss  - Y_avg)**2) / n)
            rxy   = np.sum(((ens - X_avg)/X_SD) * ((ss - Y_avg)/Y_SD)) / n
            a_m   = rxy * Y_SD / X_SD
            b_m   = Y_avg - a_m * X_avg

            # Standard error of the model (uses residuals from mean Y, matching .m)
            et = ss - Y_avg
            s  = np.sqrt(np.sum(et**2) / max(n - 2, 1))
            stdev_a = (s / np.sqrt(n)) * (1.0 / X_SD)
            stdev_b = (s / np.sqrt(n)) * np.sqrt(1 + X_avg**2 / X_SD**2)

            a_u = a_m + stdev_a;  b_u = b_m + stdev_b
            a_l = a_m - stdev_a;  b_l = b_m - stdev_b
            if b_l < 0:
                b_l = 0.0

            phi_m = np.degrees(np.arctan(a_m))
            phi_u = np.degrees(np.arctan(a_u))
            phi_l = np.degrees(np.arctan(a_l))

            x_range = np.linspace(0, max(ens.max() * 1.1, 5), 200)

            ln_m, = ax5.plot(x_range, a_m * x_range + b_m, 'k--', linewidth=1.8)
            ln_u, = ax5.plot(x_range, a_u * x_range + b_u, 'k:',  linewidth=1.2)
            ln_l, = ax5.plot(x_range, a_l * x_range + b_l, 'k:',  linewidth=1.2)
            eb    = ax5.errorbar(ens, ss, yerr=ss_sd, fmt='o', color='blue',
                                 markerfacecolor='blue', markersize=5,
                                 linewidth=0.8, capsize=3)

            legend_handles = [
                (ln_m, f"Mean: c' = {b_m:.2f} kPa,  φ' = {phi_m:.2f}°"),
                (ln_u, f"Upper Bound: c' = {b_u:.2f} kPa,  φ' = {phi_u:.2f}°"),
                (ln_l, f"Lower Bound: c' = {b_l:.2f} kPa,  φ' = {phi_l:.2f}°"),
                (eb,   'Back-calculated data'),
            ]

            # Axis limits: cover data + error bars + envelope lines, with 10% padding
            x_lim = max(ens.max() * 1.15, 5)
            y_top = max((ss + ss_sd).max(), a_u * x_lim + b_u)
            y_lim = max(y_top * 1.15, 5)
            ax5.set_xlim(0, x_lim)
            ax5.set_ylim(0, y_lim)

        ax5.set_xlabel("Effective Normal Stress, σ' (kPa)",
                       fontsize=9, fontweight='bold', fontfamily=font_name)
        ax5.set_ylabel("Shear Strength, τ (kPa)",
                       fontsize=9, fontweight='bold', fontfamily=font_name)
        _style(ax5)
        ax5.set_box_aspect(1)
        if legend_handles:
            handles, labels = zip(*legend_handles)
            ax5.legend(handles, labels,
                       loc='upper center', bbox_to_anchor=(0.5, -0.12),
                       ncol=1, frameon=False,
                       fontsize=8,
                       prop={'family': font_name, 'weight': 'bold'})

        return fig

    # ── Plot shear strength ──────────────────────────────────────────

    def _plot_strength(self, silent=False):
        if not _HAS_MPL:
            if not silent:
                self._no_mpl_error()
            return

        if self._get_raw_value('Strength_status') != 1:
            if not silent:
                messagebox.showinfo('Uniform Strength',
                                    'Strength is spatially uniform.\n'
                                    'Nothing to plot — c and φ are constant everywhere.')
            return

        c_file   = self._get_raw_value('cohesion_3d_file')
        phi_file = self._get_raw_value('phi_3d_file')
        if not c_file or not phi_file:
            if not silent:
                messagebox.showwarning('Missing Files',
                                       'Specify both 3D cohesion and friction angle files first.')
            return

        try:
            z_min_s  = float(self._get_raw_value('z_min') or 0.125)
            z_max_s  = float(self._get_raw_value('z_max') or 1.875)
            n_d_s    = int(self._get_raw_value('Depth_points') or 8)
            depths_s = np.linspace(z_min_s, z_max_s, n_d_s)
            sel_depth_s = float(self._strength_d_var.get())
            d         = int(np.argmin(np.abs(depths_s - sel_depth_s)))
            depth_str = f'{sel_depth_s:.4f} m'
        except Exception:
            d         = 0
            depth_str = 'Layer 1'

        self._clear_canvas_frame(self._strength_canvas_frame)
        self._placeholder_canvas(self._strength_canvas_frame, 'Loading shear strength arrays…')

        dem_path      = self._get_raw_value('dem_file')
        cache_key     = (c_file, phi_file)
        cached_key    = getattr(self, '_strength_vrange_key', None)
        cached_c_vmin = getattr(self, '_strength_c_vmin', None)
        cached_c_vmax = getattr(self, '_strength_c_vmax', None)
        cached_p_vmin = getattr(self, '_strength_phi_vmin', None)
        cached_p_vmax = getattr(self, '_strength_phi_vmax', None)

        def _work():
            try:
                c_arr   = self._load_h5_array(c_file)
                phi_arr = self._load_h5_array(phi_file)
            except Exception as exc:
                return ('load_error', str(exc))

            if c_arr.ndim != 3 or phi_arr.ndim != 3:
                return ('shape_error', f'Expected 3D arrays (rows × cols × depths).')
            if d < 0 or d >= c_arr.shape[2]:
                return ('depth_error', f'Array has {c_arr.shape[2]} depth layers (0-based).')

            # Size check against DEM
            if dem_path:
                try:
                    import h5py as _h5c
                    with _h5c.File(dem_path, 'r') as _f:
                        _dr, _dc = _f['data'].shape[:2]
                    _size_errs = []
                    for _arr, _lbl in [(c_arr, 'Cohesion 3D'), (phi_arr, 'Friction Angle 3D')]:
                        _r, _c = _arr.shape[0], _arr.shape[1]
                        if abs(_r - _dr) > 1 or abs(_c - _dc) > 1:
                            _size_errs.append(
                                f'{_lbl}: {_r}×{_c} is incompatible with DEM {_dr}×{_dc}')
                    if _size_errs:
                        return ('size_error',
                                'Raster size mismatch — cannot plot:\n\n' +
                                '\n'.join(f'  •  {e}' for e in _size_errs))
                except Exception:
                    pass

            # DEM mask
            try:
                dem_arr = self._load_dem_array(dem_path) if dem_path else None
                mask = np.isnan(dem_arr) if dem_arr is not None else None
            except Exception:
                mask = None

            c_2d   = self._apply_dem_mask(c_arr[:, :, d],   mask)
            phi_2d = self._apply_dem_mask(phi_arr[:, :, d], mask)

            nonlocal cached_key, cached_c_vmin, cached_c_vmax, cached_p_vmin, cached_p_vmax
            if cached_key != cache_key:
                c_vals   = c_arr[np.isfinite(c_arr)]
                phi_vals = phi_arr[np.isfinite(phi_arr)]
                cached_c_vmin = float(np.floor(c_vals.min()))   if len(c_vals)   else None
                cached_c_vmax = float(np.ceil(c_vals.max()))    if len(c_vals)   else None
                cached_p_vmin = float(np.floor(phi_vals.min())) if len(phi_vals) else None
                cached_p_vmax = float(np.ceil(phi_vals.max()))  if len(phi_vals) else None
                cached_key = cache_key
                self._strength_vrange_key = cache_key
                self._strength_c_vmin     = cached_c_vmin
                self._strength_c_vmax     = cached_c_vmax
                self._strength_phi_vmin   = cached_p_vmin
                self._strength_phi_vmax   = cached_p_vmax

            cmap_c   = _mpl.colormaps['YlOrBr'].copy(); cmap_c.set_bad(color='white')
            cmap_phi = _mpl.colormaps['RdYlGn'].copy(); cmap_phi.set_bad(color='white')

            fig = Figure(figsize=(9, 4), dpi=96, facecolor=C['surface'])
            ax1 = fig.add_subplot(1, 2, 1)
            im1 = ax1.imshow(np.ma.masked_invalid(CRISISGui._ds(c_2d)), cmap=cmap_c,
                             aspect='auto', interpolation='nearest',
                             vmin=cached_c_vmin, vmax=cached_c_vmax)
            fig.colorbar(im1, ax=ax1, fraction=0.046, pad=0.04, label='c (kPa)')
            ax1.set_title(f'Cohesion c — Depth = {depth_str}', fontsize=10, pad=5, fontweight='bold')
            ax1.axis('off')

            ax2 = fig.add_subplot(1, 2, 2)
            im2 = ax2.imshow(np.ma.masked_invalid(CRISISGui._ds(phi_2d)), cmap=cmap_phi,
                             aspect='auto', interpolation='nearest',
                             vmin=cached_p_vmin, vmax=cached_p_vmax)
            fig.colorbar(im2, ax=ax2, fraction=0.046, pad=0.04, label='φ (degrees)')
            ax2.set_title(f'Friction Angle φ — Depth = {depth_str}', fontsize=10, pad=5, fontweight='bold')
            ax2.axis('off')

            fig.tight_layout(pad=1.5)
            return ('ok', fig)

        def _done(result):
            self._clear_canvas_frame(self._strength_canvas_frame)
            if result is None:
                self._placeholder_canvas(self._strength_canvas_frame,
                                         'Load 3D cohesion and friction angle files above to auto-plot.')
                return
            kind, payload = result
            if kind == 'ok':
                self._embed_figure(payload, self._strength_canvas_frame)
            else:
                if kind != 'size_error':
                    messagebox.showerror({
                        'load_error':  'Load Error',
                        'shape_error': 'Array Shape Error',
                        'depth_error': 'Depth Out of Range',
                    }.get(kind, 'Error'), payload)
                self._placeholder_canvas(self._strength_canvas_frame,
                                         'Load 3D cohesion and friction angle files above to auto-plot.')

        self._run_in_bg(_work, _done, cancel_attr='_strength_plot_gen')

    # ── Plot outputs (DEM + landslides) ──────────────────────────────────

    def _plot_outputs(self):
        if not _HAS_MPL:
            return

        _title     = (self._get_raw_value('project_title') or 'CRISIS').strip() or 'CRISIS'
        output_dir = Path(self._get_raw_value('output_dir')) / f'{_title}_Outputs'
        if not output_dir.exists():
            return

        dem_path = self._get_raw_value('dem_file')
        if not dem_path:
            return

        fid_file = output_dir / 'FID_all_times.h5'
        if not fid_file.exists():
            return

        self._clear_canvas_frame(self._outputs_canvas_frame)
        self._placeholder_canvas(self._outputs_canvas_frame, 'Building outputs plot…')

        # Snapshot all GUI state needed by the background thread
        hydro_ind       = self._get_raw_value('Hydrological_model_indicator')
        trig_ind        = int(self._get_raw_value('Triggering_Indicator') or 1)
        eq_time_pt      = self._get_raw_value('Earthquake_time_point')
        cell_size_val   = self._get_raw_value('Cell_size') or 1.0
        xmin_val        = self._get_raw_value('xmin') or 0.0
        ymin_val        = self._get_raw_value('ymin') or 0.0
        # Flow Accumulation for the stream overlay below is only used if it was
        # already computed for this exact DEM + Cell_size (e.g. by the
        # Topography preview) — recomputing it here would mean running the
        # ~1-minute Flow Direction step a second time just for a plot
        raster_cache    = getattr(self, '_fwd_raster_cache', None)

        def _work():
            # ── Load DEM ────────────────────────────────────────────────
            try:
                dem = self._load_dem_array(dem_path)
            except Exception:
                return None

            # ── Load FID raster ─────────────────────────────────────────
            try:
                import h5py
                with h5py.File(str(fid_file), 'r') as f:
                    fid = f['data'][:].astype(np.float32)
            except Exception:
                return None

            # ── Check for landslides ─────────────────────────────────────
            if not np.any(~np.isnan(fid) & (fid > 0)):
                return 'no_landslides'

            try:
                time_pts = int(self._get_raw_value('time_points') or 1)
            except (TypeError, ValueError):
                time_pts = 1
            show_lad_series = (hydro_ind == 1 and time_pts > 1)

            # ── Build figure ─────────────────────────────────────────────
            import matplotlib.gridspec as gridspec
            fig = Figure(figsize=(14, 12), dpi=96, facecolor=C['surface'])
            gs  = gridspec.GridSpec(3, 3, figure=fig, width_ratios=[1.6, 1.8, 1.0],
                                    hspace=0.55, wspace=0.35, height_ratios=[1, 1, 1])
            fig.subplots_adjust(left=0.02, right=0.97, top=0.91, bottom=0.16)
            ax_dem = fig.add_subplot(gs[:, 0])

            # Geographic extent — x/y coordinates are cheap to compute fresh
            # from xmin/ymin/Cell_size (no need to reload/cache anything)
            _fwd_extent = None
            try:
                from forward_analysis import compute_coordinate_vectors
                _rows_dem, _cols_dem = dem.shape[:2]
                _xc, _yc = compute_coordinate_vectors(
                    _rows_dem, _cols_dem, float(xmin_val), float(ymin_val), float(cell_size_val))
                _fwd_extent = [float(_xc.min()), float(_xc.max()),
                               float(_yc.min()), float(_yc.max())]
            except Exception:
                pass

            dem_cmap = _mpl.colormaps['terrain'].copy()
            dem_cmap.set_bad(color='white')
            _imshow_kw = dict(cmap=dem_cmap, aspect='equal', interpolation='nearest')
            if _fwd_extent:
                _imshow_kw.update(extent=_fwd_extent, origin='upper')
            im = ax_dem.imshow(np.ma.masked_invalid(CRISISGui._ds(dem)), **_imshow_kw)
            cbar = fig.colorbar(im, ax=ax_dem, fraction=0.046, pad=0.04, shrink=0.7)
            cbar.set_label('Elevation (m)', rotation=270, labelpad=14, fontsize=10)

            # Stream overlay — only drawn if Flow Accumulation was already
            # cached for this exact DEM + Cell_size; skipped otherwise rather
            # than recomputing it here (see raster_cache note above)
            _cache = raster_cache or {}
            _cache_matches = (
                _cache.get('dem_file') == dem_path
                and _cache.get('Area_Mat_counts') is not None
                and abs(float(_cache.get('Cell_size', float('nan'))) - float(cell_size_val)) < 1e-9
            )
            if _cache_matches:
                try:
                    flow_area = _cache['Area_Mat_counts'] * float(cell_size_val) ** 2
                    threshold = np.nanpercentile(flow_area[flow_area > 0], 90)
                    stream_mask_data = (flow_area > threshold) & ~np.isnan(flow_area)
                    overlay_st = np.full(flow_area.shape + (4,), [1, 1, 1, 0], dtype=np.float32)
                    overlay_st[stream_mask_data] = [0, 0, 0, 0.6]
                    _st_kw = dict(aspect='equal', interpolation='nearest')
                    if _fwd_extent:
                        _st_kw.update(extent=_fwd_extent, origin='upper')
                    ax_dem.imshow(CRISISGui._ds(overlay_st), **_st_kw)
                except Exception:
                    pass

            # Shapefiles
            import geopandas as gpd
            import pandas as _pd
            from matplotlib.patches import Patch
            from matplotlib.lines import Line2D

            _fwd_shps = sorted(output_dir.glob('Predicted_landslides_time_*.shp'))
            _preloaded_polys = []

            def _extract_polys(gdf):
                coords = []
                for _fg in gdf.geometry:
                    if _fg is None:
                        continue
                    _pl = (list(_fg.geoms) if _fg.geom_type == 'MultiPolygon'
                           else [_fg] if _fg.geom_type == 'Polygon' else [])
                    for _p in _pl:
                        coords.append(_p.exterior.xy)
                return coords

            if _fwd_shps:
                try:
                    _gdfs = [gpd.read_file(str(_s)) for _s in _fwd_shps]
                    _gfwd = (gpd.GeoDataFrame(_pd.concat(_gdfs, ignore_index=True),
                                              crs=_gdfs[0].crs)
                             if len(_gdfs) > 1 else _gdfs[0])
                    _preloaded_polys = _extract_polys(_gfwd)
                except Exception:
                    pass

            if _preloaded_polys:
                for _fxs, _fys in _preloaded_polys:
                    ax_dem.fill(_fxs, _fys, facecolor='red', edgecolor='red',
                                linewidth=0.5, alpha=0.75, zorder=3)
            else:
                landslide_mask = ~np.isnan(fid) & (fid > 0)
                overlay_ls = np.full(fid.shape + (4,), [1, 1, 1, 0], dtype=np.float32)
                overlay_ls[landslide_mask] = [1, 0, 0, 0.75]
                _ls_kw = dict(aspect='equal', interpolation='nearest')
                if _fwd_extent:
                    _ls_kw.update(extent=_fwd_extent, origin='upper')
                ax_dem.imshow(CRISISGui._ds(overlay_ls), **_ls_kw)

            # North arrow — inside upper-right corner
            ax_dem.annotate('', xy=(0.94, 0.97), xytext=(0.94, 0.88),
                            xycoords='axes fraction', textcoords='axes fraction',
                            annotation_clip=False,
                            arrowprops=dict(arrowstyle='->', color='black', lw=1.5))
            ax_dem.text(0.94, 0.98, 'N', transform=ax_dem.transAxes,
                        ha='center', va='bottom', fontsize=9, fontweight='bold')

            # Scale bar — inside lower-left corner
            _xl_f, _xr_f = ax_dem.get_xlim()
            _bar_m_f     = (_xr_f - _xl_f) * 0.20
            _bar_label_f = f'{_bar_m_f:.0f} m' if _fwd_extent else '(px)'
            _bx0_f, _bx1_f = 0.05, 0.25
            _by_f, _th_f    = 0.04, 0.012
            ax_dem.plot([_bx0_f, _bx1_f], [_by_f, _by_f],
                        transform=ax_dem.transAxes, color='black', lw=1.5, zorder=10)
            for _bxf in (_bx0_f, _bx1_f):
                ax_dem.plot([_bxf, _bxf], [_by_f - _th_f, _by_f + _th_f],
                            transform=ax_dem.transAxes, color='black', lw=1.2, zorder=10)
            ax_dem.text((_bx0_f + _bx1_f) / 2, _by_f - _th_f * 2.5,
                        _bar_label_f, transform=ax_dem.transAxes,
                        ha='center', va='top', fontsize=7, zorder=10)

            legend_elements = [
                Patch(facecolor='red', edgecolor='red', label='Predicted Landslides'),
                Line2D([0], [0], color='black', lw=2, label='Streams'),
            ]
            ax_dem.legend(handles=legend_elements, loc='upper center',
                         bbox_to_anchor=(0.5, -0.04), ncol=2,
                         fontsize=9, framealpha=0.95, edgecolor='black', fancybox=True)
            ax_dem.set_title('Final Landslide Distribution (Cumulative)',
                            fontsize=11, pad=8, fontweight='bold')
            ax_dem.axis('off')
            ax_dem.text(0.5, -0.17, 'Click DEM to zoom  |  scroll to zoom in/out',
                        transform=ax_dem.transAxes, ha='center', va='top',
                        fontsize=8, fontstyle='italic', color='grey')

            # Excel histograms
            excel_file = output_dir / 'Individual_landslides_with_time.xlsx'
            hist_col   = 2

            ax_depths = fig.add_subplot(gs[0, hist_col])
            try:
                import pandas as pd
                df_depths = pd.read_excel(excel_file, sheet_name='Triggering_Depths_Cumulative_m')
                data_depths = df_depths.iloc[:, -1].dropna().values
                if len(data_depths) > 0:
                    ax_depths.hist(data_depths, bins=20, color='#4285f4', edgecolor='black', alpha=0.7)
                    ax_depths.set_xlabel('Landslide Triggering Depth (m)', fontsize=9, fontweight='bold')
                    ax_depths.set_ylabel('Frequency', fontsize=9, fontweight='bold')
                    ax_depths.grid(axis='y', alpha=0.3, linestyle='--')
                    ax_depths.tick_params(labelsize=8)
            except Exception:
                ax_depths.text(0.5, 0.5, 'No data', ha='center', va='center',
                               transform=ax_depths.transAxes, fontsize=9)

            ax_areas = fig.add_subplot(gs[1, hist_col])
            try:
                df_areas = pd.read_excel(excel_file, sheet_name='Areas_Cumulative_m2')
                data_areas = df_areas.iloc[:, -1].dropna().values
                if len(data_areas) > 0:
                    ax_areas.hist(data_areas, bins=20, color='#34a853', edgecolor='black', alpha=0.7)
                    ax_areas.set_xlabel('Landslide Area (m$^2$)', fontsize=9, fontweight='bold')
                    ax_areas.set_ylabel('Frequency', fontsize=9, fontweight='bold')
                    ax_areas.grid(axis='y', alpha=0.3, linestyle='--')
                    ax_areas.tick_params(labelsize=8)
            except Exception:
                ax_areas.text(0.5, 0.5, 'No data', ha='center', va='center',
                              transform=ax_areas.transAxes, fontsize=9)

            ax_volumes = fig.add_subplot(gs[2, hist_col])
            try:
                df_volumes = pd.read_excel(excel_file, sheet_name='Volumes_Cumulative_m3')
                data_volumes = df_volumes.iloc[:, -1].dropna().values
                if len(data_volumes) > 0:
                    ax_volumes.hist(data_volumes, bins=20, color='#ea4335', edgecolor='black', alpha=0.7)
                    ax_volumes.set_xlabel('Landslide Volume (m$^3$)', fontsize=9, fontweight='bold')
                    ax_volumes.set_ylabel('Frequency', fontsize=9, fontweight='bold')
                    ax_volumes.grid(axis='y', alpha=0.3, linestyle='--')
                    ax_volumes.tick_params(labelsize=8)
            except Exception:
                ax_volumes.text(0.5, 0.5, 'No data', ha='center', va='center',
                                transform=ax_volumes.transAxes, fontsize=9)

            # Middle column: LAD
            ax_density = fig.add_subplot(gs[:, 1])
            if show_lad_series:
                try:
                    summary_file = output_dir / 'Total_landslides_with_time.xlsx'
                    if summary_file.exists():
                        df_density = pd.read_excel(summary_file)
                        total_ls_area = df_density.iloc[:, 5].values.astype(float)
                        cell_size_flt = float(cell_size_val)
                        study_area    = float(np.sum(~np.isnan(dem))) * cell_size_flt ** 2
                        density       = (total_ls_area / study_area) * 100.0
                        time_points_val = len(density)
                        time_axis       = np.arange(time_points_val)

                        _src_plotted = False
                        if trig_ind == 1:
                            ax_density.plot(time_axis, density, color='#c0392b', lw=2)
                            ax_density.fill_between(time_axis, density, alpha=0.18, color='#c0392b')
                            _src_plotted = True

                        if not _src_plotted:
                            try:
                                import glob as _iglob
                                ind_files = _iglob.glob(str(output_dir / 'Individual_landslides_with_time*.xlsx'))
                                if ind_files and trig_ind != 1:
                                    df_areas2 = pd.read_excel(ind_files[0], sheet_name='Areas_Cumulative_m2', header=0)
                                    df_src    = pd.read_excel(ind_files[0], sheet_name='Triggering_Sources_Cumulative', header=0)
                                    n_cols = min(df_areas2.shape[1], df_src.shape[1], time_points_val)
                                    d1 = np.zeros(n_cols); d2 = np.zeros(n_cols); d3 = np.zeros(n_cols)
                                    for col_i in range(n_cols):
                                        areas_col = pd.to_numeric(df_areas2.iloc[:, col_i], errors='coerce').fillna(0).values
                                        src_col   = pd.to_numeric(df_src.iloc[:, col_i],    errors='coerce').fillna(0).values
                                        d1[col_i] = np.sum(areas_col[src_col == 1]) / study_area * 100.0
                                        d2[col_i] = np.sum(areas_col[src_col == 2]) / study_area * 100.0
                                        d3[col_i] = np.sum(areas_col[src_col == 3]) / study_area * 100.0
                                    t_src = np.arange(n_cols)
                                    ax_density.plot(time_axis, density, color='black', lw=2, label='Total', zorder=5)
                                    for d_arr, col, lbl in [
                                        (d1, '#2980b9', 'Failures triggered by rainfall'),
                                        (d2, '#8e44ad', 'Failures triggered by both'),
                                        (d3, '#e67e22', 'Failures triggered by earthquake'),
                                    ]:
                                        if np.any(d_arr > 0):
                                            ax_density.plot(t_src, d_arr, color=col, lw=2, label=lbl)
                                    ax_density.legend(fontsize=7, loc='upper left', framealpha=0.7)
                                    _src_plotted = True
                            except Exception:
                                pass

                        if not _src_plotted:
                            ax_density.plot(time_axis, density, color='#c0392b', lw=2)
                            ax_density.fill_between(time_axis, density, alpha=0.18, color='#c0392b')

                        ax_density.set_xlabel('Time Point', fontsize=9, fontweight='bold')
                        ax_density.set_ylabel('Landslide Area Density (%)', fontsize=9, fontweight='bold')
                        ax_density.set_title('Landslide Area Density Over Time', fontsize=10, pad=6, fontweight='bold')
                        ax_density.set_xlim(0, time_points_val - 1)
                        ax_density.set_ylim(bottom=0)
                        ax_density.grid(alpha=0.3, linestyle='--')
                        ax_density.tick_params(labelsize=8)
                        ax_density.spines['left'].set_position('zero')
                        ax_density.spines['bottom'].set_position('zero')
                        ax_density.spines['right'].set_visible(False)
                        ax_density.spines['top'].set_visible(False)

                        try:
                            if trig_ind == 2 and eq_time_pt is not None:
                                eq_t  = int(eq_time_pt)
                                y_top = ax_density.get_ylim()[1]
                                label_y = y_top * 0.38
                                ax_density.plot([eq_t, eq_t], [label_y * 0.85, 0.0],
                                                color='red', lw=1.6, linestyle='--', zorder=6, clip_on=False)
                                ax_density.annotate('', xy=(eq_t, 0), xytext=(eq_t, 0.001 * y_top),
                                                    arrowprops=dict(arrowstyle='->', color='red', lw=1.6), zorder=7)
                                ax_density.text(eq_t, label_y, 'Earthquake', fontsize=8, fontweight='bold',
                                                color='red', ha='center', va='bottom', zorder=8)
                        except Exception:
                            pass
                    else:
                        ax_density.text(0.5, 0.5, 'No Total_landslides_with_time file found',
                                        ha='center', va='center', transform=ax_density.transAxes, fontsize=9)
                except Exception:
                    ax_density.text(0.5, 0.5, 'No data', ha='center', va='center',
                                    transform=ax_density.transAxes, fontsize=9)
            else:
                try:
                    cell_size_flt = float(cell_size_val)
                    ls_area    = float(np.sum(~np.isnan(fid) & (fid > 0))) * cell_size_flt ** 2
                    study_area = float(np.sum(~np.isnan(dem))) * cell_size_flt ** 2
                    lad_text   = f'{(ls_area / study_area) * 100.0:.2f} %'
                except Exception:
                    lad_text = 'N/A'

                ax_density.axis('off')
                ax_density.set_facecolor(C['surface'])
                ax_density.text(0.5, 0.72, 'Landslide Area Density',
                                ha='center', va='center', transform=ax_density.transAxes,
                                fontsize=12, fontweight='bold', color=C['text'])
                ax_density.text(0.5, 0.50, lad_text,
                                ha='center', va='center', transform=ax_density.transAxes,
                                fontsize=28, fontweight='bold', color='#c0392b')
                ax_density.text(0.5, 0.32, '(total landslide area / study area)',
                                ha='center', va='center', transform=ax_density.transAxes,
                                fontsize=9, color=C['text_muted'], style='italic')
                from matplotlib.patches import FancyBboxPatch
                box = FancyBboxPatch((0.15, 0.28), 0.70, 0.52,
                                     boxstyle='round,pad=0.02', linewidth=1.5,
                                     edgecolor='#c0392b', facecolor='#fdf2f2',
                                     transform=ax_density.transAxes, clip_on=False, zorder=0)
                ax_density.add_patch(box)

            _zoom_r = ((_fwd_extent[1] - _fwd_extent[0]) * 0.05
                       if _fwd_extent else dem.shape[1] * 0.05)

            # Pre-build the zoom figure here (background thread) with a
            # downsampled copy of the DEM so the first click is instant.
            _MAX_ZPX = 800
            _fy = max(1, dem.shape[0] // _MAX_ZPX)
            _fx = max(1, dem.shape[1] // _MAX_ZPX)
            _zoom_dem = dem[::_fy, ::_fx]

            zoom_fig = Figure(figsize=(7, 7), dpi=96, facecolor='white')
            zoom_fig.subplots_adjust(left=0, right=1, top=1, bottom=0)
            zoom_ax = zoom_fig.add_subplot(111)
            zoom_ax.set_xticks([]); zoom_ax.set_yticks([])
            for _sp in zoom_ax.spines.values():
                _sp.set_visible(False)
            _zkw = dict(cmap=dem_cmap, aspect='equal', interpolation='bilinear')
            if _fwd_extent:
                _zkw.update(extent=_fwd_extent, origin='upper')
            zoom_ax.imshow(np.ma.masked_invalid(_zoom_dem), **_zkw)
            for _fxs2, _fys2 in _preloaded_polys:
                zoom_ax.fill(_fxs2, _fys2, facecolor='red', edgecolor='red',
                             linewidth=0.8, alpha=0.75, zorder=3)

            return {
                'fig':      fig,
                'ax_dem':   ax_dem,
                'zoom_fig': zoom_fig,
                'zoom_ax':  zoom_ax,
                'zoom_r':   _zoom_r,
            }

        def _done(result):
            self._clear_canvas_frame(self._outputs_canvas_frame)
            if result is None:
                self._placeholder_canvas(self._outputs_canvas_frame,
                                         'Run the model to view DEM with predicted landslides.')
                return
            if result == 'no_landslides':
                msg_frame = tk.Frame(self._outputs_canvas_frame, bg=C['surface'])
                msg_frame.place(relx=0.5, rely=0.5, anchor='center')
                tk.Label(msg_frame, text='No landslides are triggered.',
                         bg=C['surface'], fg=C['text_muted'],
                         font=('Segoe UI', 14)).pack()
                return

            fig      = result['fig']
            ax_dem   = result['ax_dem']
            zoom_fig = result['zoom_fig']
            zoom_ax  = result['zoom_ax']
            _zoom_r  = result['zoom_r']

            # Create the zoom Toplevel now (main thread), but keep it hidden.
            # The Figure was already built in the background thread so no
            # heavy work happens here or on click.
            from matplotlib.backends.backend_tkagg import (
                FigureCanvasTkAgg as _FCTA, NavigationToolbar2Tk as _NavTB)
            zoom_top = tk.Toplevel(self.root)
            zoom_top.title('DEM Zoom')
            zoom_top.geometry('720x760+200+30')
            zoom_top.configure(bg='white')
            zoom_top.withdraw()   # hidden until first click
            zoom_top.protocol('WM_DELETE_WINDOW', zoom_top.withdraw)

            zcanvas = _FCTA(zoom_fig, master=zoom_top)
            _ztoolbar = _NavTB(zcanvas, zoom_top)
            _ztoolbar.update()
            _ztoolbar.pack(side='bottom', fill='x')
            zcanvas.get_tk_widget().pack(fill='both', expand=True)

            def _on_scroll_zoom(ev):
                factor = 0.7 if ev.button == 'up' else 1.4
                xl, xr = zoom_ax.get_xlim(); yb, yt = zoom_ax.get_ylim()
                mx = ev.xdata if ev.xdata is not None else (xl + xr) / 2
                my = ev.ydata if ev.ydata is not None else (yb + yt) / 2
                zoom_ax.set_xlim(mx + (xl - mx) * factor, mx + (xr - mx) * factor)
                zoom_ax.set_ylim(my + (yb - my) * factor, my + (yt - my) * factor)
                zcanvas.draw_idle()

            zoom_fig.canvas.mpl_connect('scroll_event', _on_scroll_zoom)

            # Trigger the very first draw of the zoom canvas 400 ms after the
            # main figure appears — runs on the Tk idle loop but the DEM is
            # downsampled so it completes in milliseconds.
            _zoom_drawn = [False]
            def _pre_draw_zoom():
                zcanvas.draw()
                _zoom_drawn[0] = True
            self.root.after(400, _pre_draw_zoom)

            def _on_fwd_dem_click(event, _r=_zoom_r):
                if event.inaxes is not ax_dem or event.button != 1:
                    return
                cx, cy = event.xdata, event.ydata
                if cx is None or cy is None:
                    return
                zoom_ax.set_xlim(cx - _r, cx + _r)
                zoom_ax.set_ylim(cy - _r, cy + _r)
                # If pre-draw hasn't fired yet, do it now (still fast)
                if not _zoom_drawn[0]:
                    zcanvas.draw()
                    _zoom_drawn[0] = True
                else:
                    zcanvas.draw_idle()
                zoom_top.deiconify()
                zoom_top.lift()

            canvas_widget = self._embed_figure(fig, self._outputs_canvas_frame)
            fig.canvas.mpl_connect('button_press_event', _on_fwd_dem_click)

        self._run_in_bg(_work, _done, cancel_attr='_outputs_plot_gen')


    # ------------------------------------------------------------------
    # Value helpers
    # ------------------------------------------------------------------

    def _get_raw_value(self, key):
        if key not in self._vars:
            return None
        var, vtype, choices = self._vars[key]
        raw = var.get()
        if vtype == 'choice':
            for rv, display in choices:
                if display == raw:
                    return rv
            return None
        elif vtype == 'searchable_choice':
            # Entry shows "EPSG:XXXX — Name"; extract the code before " — "
            if ' — ' in raw:
                return raw.split(' — ')[0].strip()
            # User may have typed a raw code directly; return as-is
            return raw if raw else None
        elif vtype == 'float':
            try:   return float(raw)
            except ValueError: return None
        elif vtype == 'int':
            try:   return int(raw)
            except ValueError: return None
        else:
            return raw

    def _collect_params(self):
        _BA_CATS = {'BA Raster Inputs', 'BA Hydrology', 'BA Scalar Inputs',
                    'BA Shear Strength', 'Mapped Landslide Inventory'}
        _FWD_CATS = {'Topography', 'Earthquake Inputs', 'Hydrological Formulations',
                     'Scalar Inputs', 'Shear Strength', 'Triggering Event'}
        mode = self._get_raw_value('analysis_type')  # 'forward', 'back', or None

        params, errors = {}, []
        for p in self._params_list:
            key = p['key']
            val = self._get_raw_value(key)
            if p.get('required', False):
                cat = p.get('category', '')
                # Skip required-field validation for the inactive mode.
                # Use 'not back' / 'not forward' so that an unset analysis_type
                # still skips the inactive side rather than validating both.
                if mode != 'back' and cat in _BA_CATS:
                    pass
                elif mode != 'forward' and cat in _FWD_CATS:
                    pass
                elif val is None or (isinstance(val, str) and not val.strip()):
                    errors.append(f'Required field is empty:  {p["label"]}')
            params[key] = val

        # xmin/ymin/Cell_size are always meters (they drive slope, area, and
        # volume calculations directly), so the output shapefiles are only
        # georeferenced correctly if the tagged CRS is itself projected in
        # meters. A geographic (lat/lon, degree-based) CRS silently produces
        # unusable shapefiles — no error at write time, but out-of-range
        # coordinates that blow up on reprojection/display in GIS software.
        crs_val = params.get('Coordinate_Reference_System')
        if crs_val:
            try:
                from pyproj import CRS as _PyprojCRS
            except ImportError:
                _PyprojCRS = None
            if _PyprojCRS is not None:
                try:
                    crs_obj = _PyprojCRS(str(crs_val))
                except Exception:
                    errors.append(
                        f'Coordinate Reference System "{crs_val}" is not a recognized '
                        'EPSG code or CRS string. Pick an entry from the dropdown, or '
                        'type a valid EPSG code (e.g. "EPSG:32619").'
                    )
                else:
                    if crs_obj.is_geographic:
                        errors.append(
                            'Coordinate Reference System is a geographic (lat/lon) CRS: '
                            f'{crs_val}. xmin, ymin, and Cell_size are always in meters, so '
                            'the CRS must be a projected system whose units are meters '
                            '(e.g. a UTM zone or State Plane in meters) — otherwise the '
                            'output shapefiles will not be usable in GIS software. '
                            'Pick a projected entry from the dropdown instead.'
                        )
        return params, errors

    # ------------------------------------------------------------------
    # Conditional visibility
    # ------------------------------------------------------------------

    def _update_visibility(self):
        for p in self._params_list:
            if 'show_if' not in p:
                continue
            frame = self._frames.get(p['key'])
            if frame is None:
                continue
            cond = p['show_if']
            if isinstance(cond, list):
                should_show = any(
                    self._get_raw_value(c['param']) == c['value'] for c in cond
                )
            else:
                should_show = (self._get_raw_value(cond['param']) == cond['value'])
            if should_show:
                frame.grid()
            else:
                frame.grid_remove()

    # ------------------------------------------------------------------
    # Dialogs
    # ------------------------------------------------------------------

    def _browse_file(self, var, file_filter):
        try:
            desc = file_filter.split('(')[0].strip()
            ext  = file_filter.split('(')[1].rstrip(')')
        except IndexError:
            desc, ext = 'All files', '*.*'
        path = filedialog.askopenfilename(
            filetypes=[(desc, ext), ('All files', '*.*')])
        if path:
            var.set(path)

    def _browse_dir(self, var):
        path = filedialog.askdirectory()
        if path:
            var.set(path)

    def _save_config(self):
        path = filedialog.asksaveasfilename(
            defaultextension='.json',
            filetypes=[('JSON config', '*.json'), ('All files', '*.*')],
            title='Save Configuration')
        if not path:
            return
        params, _ = self._collect_params()
        with open(path, 'w', encoding='utf-8') as fh:
            json.dump(params, fh, indent=2, default=str)
        messagebox.showinfo('Saved', f'Configuration saved to:\n{path}')

    def _load_config(self):
        path = filedialog.askopenfilename(
            filetypes=[('JSON config', '*.json'), ('All files', '*.*')],
            title='Load Configuration')
        if not path:
            return
        with open(path, 'r', encoding='utf-8') as fh:
            config = json.load(fh)
        for key, val in config.items():
            if key not in self._vars:
                continue
            var, vtype, choices = self._vars[key]
            if vtype in ('choice', 'searchable_choice') and choices:
                for rv, display in choices:
                    if str(rv) == str(val):
                        if vtype == 'searchable_choice':
                            var.set(f"{rv} — {display}")
                        else:
                            var.set(display)
                        break
                else:
                    var.set(str(val) if val is not None else '')
            else:
                var.set(str(val) if val is not None else '')
        self._update_visibility()

    def _on_close(self):
        if self._running:
            if not messagebox.askyesno(
                    'Model Running',
                    'The model is still running.\n\nExit anyway?',
                    icon='warning'):
                return

        answer = messagebox.askyesnocancel(
            'Save Project',
            'Would you like to save your configuration before exiting?\n\n'
            'Your inputs can be reloaded next time via "Load Config".')

        if answer is None:       # Cancel — stay in app
            return
        if answer:               # Yes — open save dialog, then exit
            path = filedialog.asksaveasfilename(
                defaultextension='.json',
                filetypes=[('CRISIS config', '*.json'), ('All files', '*.*')],
                title='Save Configuration')
            if path:
                params, _ = self._collect_params()
                with open(path, 'w', encoding='utf-8') as fh:
                    json.dump(params, fh, indent=2, default=str)
            # If they dismissed the save dialog without saving, still exit
        self.root.destroy()

    def _clear_config(self):
        if not messagebox.askyesno('Clear Configuration',
                                   'Reset all inputs to their default values?\n\n'
                                   'This cannot be undone.'):
            return
        self.root.after(0, self._reload_params)

    # ------------------------------------------------------------------
    # About dialog
    # ------------------------------------------------------------------

    def _show_about(self):
        about_text = (
            "CRISIS is an open-source, physics-based model designed to simulate 3D "
            "rainfall-induced and earthquake-triggered landslides at regional scales. "
            "It evaluates slope stability across a raster domain, generating realistic "
            "3D landslide geometries — each with a defined location, area, and volume — "
            "using a pseudo-3D upslope search algorithm.\n\n"
            "CRISIS operates in two modes: a forward predictive analysis that "
            "forecasts landslide occurrence over time, and a back-analysis that "
            "estimates shear strength parameters from a mapped landslide inventory "
            "across a large spatial scale.\n\n"
            "The back-analysis mode supports multi-core parallel processing, "
            "distributing the search over mapped landslides across available CPU "
            "cores to efficiently handle large inventories with hundreds of "
            "landslides and dense shear strength search grids.\n\n"
            "The model is designed to be highly flexible, accommodating a wide range "
            "of user needs — from simple simulations using uniform shear strength and "
            "a constant pore pressure coefficient, to advanced configurations with "
            "spatially variable 3D strength fields, transient pore pressure files, and "
            "coupled seismic displacement models. This makes CRISIS accessible to a "
            "broad user base, from practitioners seeking quick regional assessments to "
            "researchers conducting detailed process-based studies.\n\n"
            "CRISIS accepts spatially and temporally varying pore pressure heads "
            "estimated from external transient groundwater flow models, enabling "
            "high-resolution, time-evolving stability assessments throughout a storm "
            "sequence, or a constant pore pressure ratio (Ru) as a simple "
            "alternative. For coupled rainfall-earthquake scenarios, seismic "
            "displacement is estimated using established empirical models.\n\n"
            "Whether you are studying landslide hazards in mountainous catchments or "
            "assessing regional susceptibility to coupled rainfall-earthquake events, "
            "CRISIS offers a robust, high-resolution solution grounded in physical "
            "process understanding, capable of running on a standard laptop or a "
            "high-performance computing cluster."
        )

        citation_text = (
            "Kassem, M., & Zekkos, D. (2026). Assessing rainfall-induced landslide "
            "failure mechanisms in regional hydrological and hillslope stability "
            "simulations. Landslides, 1–20. https://doi.org/10.1007/s10346-025-02690-w"
        )

        dlg = tk.Toplevel(self.root)
        dlg.title('About CRISIS')
        dlg.configure(bg=C['surface'])
        dlg.geometry('660x600')
        dlg.minsize(480, 400)
        dlg.transient(self.root)
        dlg.grab_set()

        dlg.update_idletasks()
        px = self.root.winfo_rootx() + (self.root.winfo_width() - 660) // 2
        py = self.root.winfo_rooty() + (self.root.winfo_height() - 600) // 2
        dlg.geometry(f'+{max(px, 0)}+{max(py, 0)}')

        header = tk.Frame(dlg, bg=C['header_bg'])
        header.pack(fill=tk.X)
        tk.Label(header, text='About CRISIS',
                 bg=C['header_bg'], fg=C['text_on_dark'],
                 font=('Segoe UI', 15, 'bold')
                 ).pack(anchor='w', padx=18, pady=12)

        body = tk.Frame(dlg, bg=C['surface'], padx=18, pady=14)
        body.pack(fill=tk.BOTH, expand=True)

        text_frame = tk.Frame(body, bg=C['surface'])
        text_frame.pack(fill=tk.BOTH, expand=True)

        scrollbar = tk.Scrollbar(text_frame)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        txt = tk.Text(text_frame, wrap=tk.WORD, bg=C['surface'], fg=C['text'],
                       font=F['ui'], relief='flat', padx=4, pady=4,
                       yscrollcommand=scrollbar.set, cursor='arrow')
        txt.pack(fill=tk.BOTH, expand=True)
        scrollbar.config(command=txt.yview)

        txt.tag_configure('h', font=F['ui_bold'], foreground=C['accent'],
                           spacing3=8)
        txt.tag_configure('body', font=F['ui'], foreground=C['text'],
                           spacing3=10)
        txt.tag_configure('cite', font=('Consolas', 9), foreground=C['text_muted'],
                           spacing1=2)

        txt.insert(tk.END, 'About\n', 'h')
        txt.insert(tk.END, about_text + '\n\n', 'body')
        txt.insert(tk.END, 'Citation\n', 'h')
        txt.insert(tk.END, 'If you use CRISIS in published research, please cite:\n\n',
                   'body')
        txt.insert(tk.END, citation_text, 'cite')
        txt.config(state=tk.DISABLED)

        footer = tk.Frame(dlg, bg=C['surface'], padx=18, pady=(0, 14))
        footer.pack(fill=tk.X)

        def _copy_citation():
            self.root.clipboard_clear()
            self.root.clipboard_append(citation_text)
            self.root.update()
            copy_btn.config(text='✓  Copied')
            dlg.after(1200, lambda: copy_btn.config(text='⎘  Copy Citation'))

        copy_btn = _flat_btn(footer, '⎘  Copy Citation', _copy_citation,
                              bg=C['surface'], fg=C['accent'],
                              hover_bg=C['accent_light'])
        copy_btn.pack(side=tk.LEFT)

        _flat_btn(footer, 'Close', dlg.destroy,
                  bg=C['accent'], fg=C['text_on_dark'], hover_bg=C['accent_dark']
                  ).pack(side=tk.RIGHT)

        dlg.bind('<Escape>', lambda e: dlg.destroy())

    # ------------------------------------------------------------------
    # Cell Inspector tab
    # ------------------------------------------------------------------

    def _build_cell_inspector_tab(self, parent):
        """Build the Cell Inspector tab with profile plots + animation."""
        self._ci_anim_id   = None
        self._ci_load_after = None
        self._ci_data      = None
        self._ci_t         = 0

        # ── Control bar ─────────────────────────────────────────────
        ctrl = tk.Frame(parent, bg=C['surface'], pady=6)
        ctrl.pack(fill=tk.X, padx=12)

        tk.Label(ctrl, text='Triggering Cell #:', bg=C['surface'],
                 font=F['ui_bold']).pack(side=tk.LEFT)

        self._ci_rank_var = tk.StringVar(value='1')
        self._ci_spinbox  = ttk.Spinbox(ctrl, from_=1, to=9999, width=7,
                                        textvariable=self._ci_rank_var)
        self._ci_spinbox.pack(side=tk.LEFT, padx=(4, 8))

        # Auto-load on spinbox change (both keyboard confirm and arrow buttons)
        self._ci_rank_var.trace_add('write', lambda *_: self._ci_load())

        self._ci_info_lbl = tk.Label(ctrl,
                                     text='Run the forward simulation first to enable the Cell Inspector.',
                                     bg=C['surface'], fg=C['text_muted'],
                                     font=F['ui_sm'])
        self._ci_info_lbl.pack(side=tk.LEFT, padx=14)

        # Animation controls (right side)
        anim_f = tk.Frame(ctrl, bg=C['surface'])
        anim_f.pack(side=tk.RIGHT)

        self._ci_time_lbl = tk.Label(anim_f, text='Time: —',
                                     bg=C['surface'], font=F['ui_sm'], width=14)
        self._ci_time_lbl.pack(side=tk.RIGHT, padx=(10, 0))

        speed_f = tk.Frame(anim_f, bg=C['surface'])
        speed_f.pack(side=tk.RIGHT, padx=(8, 4))
        tk.Label(speed_f, text='← slower   Speed   faster →', bg=C['surface'],
                 fg=C['text_muted'], font=F['ui_sm']).pack(side=tk.TOP)
        self._ci_speed_var = tk.IntVar(value=1500)
        ttk.Scale(speed_f, from_=2000, to=100, variable=self._ci_speed_var,
                  orient='horizontal', length=110).pack(side=tk.TOP)

        _flat_btn(anim_f, '■  Stop', self._ci_stop,
                  bg=C['surface'], fg=C['danger'],
                  hover_bg='#fde8e8').pack(side=tk.RIGHT, padx=3)
        _flat_btn(anim_f, '▶  Play', self._ci_play,
                  bg=C['success'], fg=C['text_on_dark'],
                  hover_bg=C['success_dark']).pack(side=tk.RIGHT, padx=(0, 3))

        # ── Matplotlib figure ────────────────────────────────────────
        if not _HAS_MPL:
            tk.Label(parent, text='matplotlib not available',
                     bg=C['surface']).pack(expand=True)
            return

        fig = Figure(figsize=(13, 5.5), facecolor=C['surface'])
        fig.subplots_adjust(left=0.06, right=0.97,
                            bottom=0.11, top=0.90, wspace=0.42)
        ax_dem = fig.add_subplot(1, 3, 1)
        ax_p   = fig.add_subplot(1, 3, 2)
        ax_fs  = fig.add_subplot(1, 3, 3)
        self._ci_fig    = fig
        self._ci_ax_dem = ax_dem
        self._ci_ax_p   = ax_p
        self._ci_ax_fs  = ax_fs
        self._ci_cell_marker = None

        ax_dem.set_facecolor(C['surface'])
        ax_dem.axis('off')
        ax_dem.set_title('Triggering Cell Location', fontsize=10, fontweight='bold', pad=6)
        ax_dem.text(0.5, 0.5, 'Load a cell to\nview location',
                    ha='center', va='center', transform=ax_dem.transAxes,
                    fontsize=9, color=C['text_muted'])

        self._ci_line_p,  = ax_p.plot( [], [], color='#2980b9', lw=2.0)
        self._ci_line_fs, = ax_fs.plot([], [], color='#c0392b', lw=2.0)

        # FS = 1 reference line
        ax_fs.axvline(1.0, color='#999999', lw=1.2, linestyle='--', zorder=1)
        ax_fs.text(1.04, 0.93, 'FS = 1',
                   transform=ax_fs.get_xaxis_transform(),
                   fontsize=11, color='#666666', fontweight='bold')

        for ax in (ax_p, ax_fs):
            ax.set_ylabel('Depth (m)', fontsize=11, fontweight='bold')
            ax.tick_params(labelsize=10)
            ax.grid(alpha=0.3, linestyle='--')
            ax.set_facecolor(C['surface'])
            ax.text(0.5, 0.5, 'Load a cell to view profiles',
                    ha='center', va='center', transform=ax.transAxes,
                    fontsize=9, color=C['text_muted'])

        ax_p.set_xlabel('Pore Pressure Head (m)', fontsize=11, fontweight='bold')
        ax_fs.set_xlabel('Factor of Safety',       fontsize=11, fontweight='bold')
        ax_p.set_title( 'Pressure Head Profile',   fontsize=11, fontweight='bold', pad=6)
        ax_fs.set_title('Factor of Safety Profile',fontsize=11, fontweight='bold', pad=6)

        fig_frame = tk.Frame(parent, bg=C['surface'],
                             highlightbackground=C['border'], highlightthickness=1)
        fig_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(4, 10))
        canvas = FigureCanvasTkAgg(fig, master=fig_frame)
        canvas.draw()
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        self._ci_canvas = canvas

    def _ci_load(self):
        """Debounced entry point — waits 400 ms after last change before loading."""
        if hasattr(self, '_ci_load_after') and self._ci_load_after is not None:
            self.root.after_cancel(self._ci_load_after)
        self._ci_load_after = self.root.after(400, self._ci_load_now)

    def _ci_load_now(self):
        """Load pressure + FS data for the chosen cell (background thread)."""
        self._ci_load_after = None
        self._ci_stop()
        try:
            rank = int(self._ci_rank_var.get())
        except ValueError:
            return

        # Serve from cache if available (avoids re-opening all h5 files)
        if not hasattr(self, '_ci_cell_cache'):
            self._ci_cell_cache = {}
        if rank in self._ci_cell_cache:
            self._ci_data = self._ci_cell_cache[rank]
            self._ci_t = 0
            self._ci_on_load_done()
            return

        self._ci_info_lbl.config(
            text=f'Loading cell {rank} — this may take a moment on first load…',
            fg=C['accent'])
        self.root.update_idletasks()

        def _worker():
            try:
                import h5py
                import pandas as pd
                from forward_analysis import Compute_Fs_vector as _fs_fn

                # ── Individual file ──────────────────────────────────
                _title  = (self._get_raw_value('project_title') or 'CRISIS').strip() or 'CRISIS'
                out_dir = Path(self._get_raw_value('output_dir') or '.') / f'{_title}_Outputs'
                ind = sorted(out_dir.glob('Individual_landslides_with_time*.xlsx'))
                if not ind:
                    raise FileNotFoundError(
                        f'Individual_landslides_with_time*.xlsx not found in:\n{out_dir}\n\n'
                        'Run the forward simulation first, or check that the Project Title '
                        'and Output Directory match the folder where results were saved.'
                    )
                from openpyxl import load_workbook as _load_wb
                _wb = _load_wb(ind[0], read_only=True, data_only=True)
                _ws = _wb['Triggering_Cells_Cumulative']
                _raw = [row[-1].value for row in _ws.iter_rows(min_row=2)]
                _wb.close()
                last_col = np.array([int(v) for v in _raw if v is not None
                                     and str(v).strip() not in ('', 'None')],
                                    dtype=np.int64)
                n_cells = len(last_col)
                if rank < 1 or rank > n_cells:
                    raise ValueError(f'Cell # must be 1–{n_cells}.')
                lin_idx = int(last_col[rank - 1])

                # ── Grid shape from DEM ──────────────────────────────
                dem_f = self._get_raw_value('dem_file')
                if not dem_f:
                    raise FileNotFoundError('DEM file not configured.')
                rows_g, clms_g = self._load_dem_array(dem_f).shape
                ro, co = (int(v) for v in np.unravel_index(lin_idx, (rows_g, clms_g)))

                # ── Depths ───────────────────────────────────────────
                z_min  = float(self._get_raw_value('z_min')        or 0.125)
                z_max  = float(self._get_raw_value('z_max')        or 1.875)
                n_dep  = int(self._get_raw_value('Depth_points')   or 8)
                depths = np.linspace(z_min, z_max, n_dep, dtype=np.float64)

                # ── Slope — auto-derived from the DEM at just this cell ──
                slope_val = 30.0
                try:
                    from forward_analysis import compute_slope_gradient8
                    _dem_arr = self._load_dem_array(dem_f)
                    _cell_size_ci = float(self._get_raw_value('Cell_size') or 1.0)
                    slope_val = float(compute_slope_gradient8(
                        _dem_arr, _cell_size_ci, unit='degree')[ro, co])
                except Exception:
                    pass

                # ── Geotechnical params ──────────────────────────────
                Gamma     = float(self._get_raw_value('Gamma')   or 18.0)
                Gamma_W   = 9.81
                c_val     = float(self._get_raw_value('c')       or 0.0)
                Phi_val   = float(self._get_raw_value('Phi')     or 30.0)
                Theta_s   = float(self._get_raw_value('Theta_s') or 0.40)
                Theta_r   = float(self._get_raw_value('Theta_r') or 0.05)
                St_status = int(self._get_raw_value('Strength_status')               or 0)
                St_theta  = int(self._get_raw_value('Strength_variation_with_theta') or 0)

                C3D = Ph3D = None
                if St_status == 1:
                    c3f = self._get_raw_value('cohesion_3d_file')
                    p3f = self._get_raw_value('phi_3d_file')
                    if c3f:
                        with h5py.File(c3f, 'r') as f: C3D  = f['data'][:]
                    if p3f:
                        with h5py.File(p3f, 'r') as f: Ph3D = f['data'][:]

                # ── Pressure file directory ──────────────────────────
                hydro_dir = Path(self._get_raw_value('hydro_files_dir') or '.')
                theta_dir = Path(self._get_raw_value('theta_files_dir') or str(hydro_dir))

                t_files = sorted(
                    hydro_dir.glob('Pressure_head_T*Z1.h5'),
                    key=lambda ff: int(
                        ff.stem.replace('Pressure_head_T', '').replace('Z1', ''))
                )
                n_time = len(t_files)
                if n_time == 0:
                    raise FileNotFoundError(
                        f'No Pressure_head_T*Z1.h5 files found in:\n{hydro_dir}')

                # ── Pre-index existing files (avoids n_time×n_dep stat calls) ─
                _p_names = {f.name for f in hydro_dir.glob('Pressure_head_T*Z*.h5')}
                _t_names = ({f.name for f in theta_dir.glob('Theta_T*Z*.h5')}
                            if St_theta == 1 else set())

                # ── Load pressures + compute FS ──────────────────────
                P_all  = np.zeros((n_time, n_dep), dtype=np.float64)
                FS_all = np.zeros((n_time, n_dep), dtype=np.float64)

                for t in range(n_time):
                    P_vec  = np.zeros(n_dep,       dtype=np.float64)
                    Th_vec = np.full(n_dep, Theta_s, dtype=np.float64)
                    for d in range(n_dep):
                        pname = f'Pressure_head_T{t}Z{d+1}.h5'
                        if pname in _p_names:
                            with h5py.File(hydro_dir / pname, 'r') as f:
                                P_vec[d] = min(float(f['data'][ro, co]), depths[d])
                        if St_theta == 1:
                            tname = f'Theta_T{t}Z{d+1}.h5'
                            if tname in _t_names:
                                with h5py.File(theta_dir / tname, 'r') as f:
                                    Th_vec[d] = float(f['data'][ro, co])
                    P_all[t]  = P_vec
                    FS_all[t] = _fs_fn(St_status, c_val, Phi_val,
                                       Theta_s, Theta_r, C3D, Ph3D,
                                       Gamma_W, Gamma, P_vec, Th_vec,
                                       slope_val, depths, ro, co, St_theta)

                # ── DEM for location panel (cached across cells) ─────
                dem_arr = getattr(self, '_ci_dem_cache', None)
                dem_x   = getattr(self, '_ci_dem_x_cache', None)
                dem_y   = getattr(self, '_ci_dem_y_cache', None)
                if dem_arr is None:
                    try:
                        from forward_analysis import compute_coordinate_vectors
                        dem_arr = self._load_dem_array(dem_f)
                        self._ci_dem_cache = dem_arr
                        # x/y coordinate vectors, auto-derived from xmin/ymin/Cell_size
                        _xmin_ci = float(self._get_raw_value('xmin') or 0.0)
                        _ymin_ci = float(self._get_raw_value('ymin') or 0.0)
                        _cs_ci   = float(self._get_raw_value('Cell_size') or 1.0)
                        _rows_ci, _cols_ci = dem_arr.shape
                        dem_x, dem_y = compute_coordinate_vectors(
                            _rows_ci, _cols_ci, _xmin_ci, _ymin_ci, _cs_ci)
                        self._ci_dem_x_cache = dem_x
                        self._ci_dem_y_cache = dem_y
                    except Exception:
                        dem_arr = dem_x = dem_y = None

                self._ci_data = {
                    'ro': ro, 'co': co, 'lin_idx': lin_idx, 'rank': rank,
                    'depths': depths, 'P_all': P_all, 'FS_all': FS_all,
                    'n_time': n_time, 'slope': slope_val, 'n_cells': n_cells,
                    'dem_arr': dem_arr, 'dem_x': dem_x, 'dem_y': dem_y,
                }
                # Cache so revisiting a cell is instant
                if not hasattr(self, '_ci_cell_cache'):
                    self._ci_cell_cache = {}
                self._ci_cell_cache[rank] = self._ci_data
                if len(self._ci_cell_cache) > 30:
                    self._ci_cell_cache.pop(next(iter(self._ci_cell_cache)))
                self._ci_t = 0
                self.root.after(0, self._ci_on_load_done)

            except Exception as exc:
                err = str(exc)
                self.root.after(0, lambda: self._ci_on_load_error(err))

        threading.Thread(target=_worker, daemon=True).start()

    def _ci_on_load_done(self):
        d = self._ci_data
        self._ci_info_lbl.config(
            text=(f'Cell {d["rank"]} of {d["n_cells"]}  |  '
                  f'Row={d["ro"]}, Col={d["co"]}  |  '
                  f'Slope={d["slope"]:.1f}°  |  '
                  f'{d["n_time"]} time steps'),
            fg=C['text'])
        self._ci_spinbox.config(to=d['n_cells'])
        self._ci_setup_axes()
        self._ci_update_dem_panel(d)
        self._ci_update_frame(0)

    def _ci_on_load_error(self, msg):
        # If the output file simply doesn't exist yet (no run completed),
        # show a quiet hint in the label rather than a disruptive popup.
        if 'not found' in msg.lower() or 'no such file' in msg.lower():
            self._ci_info_lbl.config(
                text='Run the forward simulation first to enable the Cell Inspector.',
                fg=C['text_muted'])
            return
        self._ci_info_lbl.config(text=f'Error: {msg}', fg=C['danger'])
        messagebox.showerror('Load error', msg)

    def _ci_update_dem_panel(self, d):
        """Draw (or update) the DEM location panel with the triggering cell marked."""
        ax   = self._ci_ax_dem
        dem  = d.get('dem_arr')
        dem_x = d.get('dem_x')
        dem_y = d.get('dem_y')
        ro, co = d['ro'], d['co']

        if dem is None:
            return

        nr, nc = dem.shape

        # Build extent from coordinates if available
        if dem_x is not None and dem_y is not None and len(dem_x) >= nc and len(dem_y) >= nr:
            extent = [float(dem_x[:nc].min()), float(dem_x[:nc].max()),
                      float(dem_y[:nr].min()), float(dem_y[:nr].max())]
            mx = float(dem_x[min(co, len(dem_x) - 1)])
            my = float(dem_y[min(ro, len(dem_y) - 1)])
        else:
            extent = None
            mx, my = co, ro

        # Draw DEM image only once (skip if already drawn)
        if not getattr(self, '_ci_dem_drawn', False):
            ax.cla()
            dem_cmap = _mpl.colormaps['terrain'].copy()
            dem_cmap.set_bad(color='white')
            kw = dict(cmap=dem_cmap, aspect='equal', interpolation='nearest')
            if extent:
                kw.update(extent=extent, origin='upper')
            ax.imshow(np.ma.masked_invalid(CRISISGui._ds(dem)), **kw)
            ax.axis('off')
            ax.set_title('Triggering Cell Location', fontsize=10, fontweight='bold', pad=6)
            self._ci_dem_drawn = True

        # Remove previous marker and draw new one
        if self._ci_cell_marker is not None:
            try:
                self._ci_cell_marker.remove()
            except Exception:
                pass
            self._ci_cell_marker = None

        marker_obj, = ax.plot(mx, my, '*', color='red', markersize=14,
                              markeredgecolor='white', markeredgewidth=0.7,
                              zorder=10, clip_on=False)
        self._ci_cell_marker = marker_obj
        self._ci_canvas.draw_idle()

    def _ci_setup_axes(self):
        """Set fixed axis limits and static decorations."""
        d = self._ci_data
        if d is None:
            return
        depths = d['depths']
        z_min, z_max = float(depths[0]), float(depths[-1])

        ax_p, ax_fs = self._ci_ax_p, self._ci_ax_fs

        # Fixed x-limits
        ax_p.set_xlim(-3, 3)
        ax_fs.set_xlim(0, 5)

        # y-axis: z_min at top, z_max at bottom (depth increases downward)
        ax_p.set_ylim(z_max, z_min)
        ax_fs.set_ylim(z_max, z_min)

        # Vertical dashed line at pressure head = 0
        if not hasattr(self, '_ci_vline_p') or self._ci_vline_p not in ax_p.lines:
            self._ci_vline_p = ax_p.axvline(0, color='gray', lw=1.0,
                                             linestyle='--', zorder=1)

        # Vertical line at FS = 1
        if not hasattr(self, '_ci_vline_fs') or self._ci_vline_fs not in ax_fs.lines:
            self._ci_vline_fs = ax_fs.axvline(1, color='gray', lw=1.0,
                                               linestyle='--', zorder=1)

        # Failure text (hidden until triggered)
        if not hasattr(self, '_ci_fail_txt') or self._ci_fail_txt not in ax_fs.texts:
            self._ci_fail_txt = ax_fs.text(
                2.5, z_min + (z_max - z_min) * 0.08,
                'Failure!', color='red', fontsize=13, fontweight='bold',
                ha='center', va='bottom', zorder=10, visible=False)

        # Remove placeholder text
        for ax in (ax_p, ax_fs):
            for txt in list(ax.texts):
                if 'Load a cell' in txt.get_text():
                    txt.remove()

    def _ci_update_frame(self, t):
        """Redraw both profiles at time step t. Returns True if FS<=1 (failure)."""
        d = self._ci_data
        if d is None:
            return False
        t = max(0, min(t, d['n_time'] - 1))
        self._ci_t = t
        depths = d['depths']
        self._ci_line_p.set_data( d['P_all'][t],  depths)
        self._ci_line_fs.set_data(d['FS_all'][t], depths)
        self._ci_ax_p.set_title(
            f'Pressure Head Profile  (t = {t})',
            fontsize=10, fontweight='bold', pad=6)
        self._ci_ax_fs.set_title(
            f'Factor of Safety Profile  (t = {t})',
            fontsize=10, fontweight='bold', pad=6)
        self._ci_time_lbl.config(text=f'Time: {t} / {d["n_time"] - 1}')

        # Show "Failure!" if any FS layer <= 1
        failed = bool(np.any(d['FS_all'][t] <= 1.0))
        if hasattr(self, '_ci_fail_txt'):
            self._ci_fail_txt.set_visible(failed)

        self._ci_canvas.draw_idle()
        return failed

    def _ci_play(self):
        if self._ci_data is None or self._ci_anim_id is not None:
            return
        self._ci_tick()

    def _ci_tick(self):
        t = self._ci_t + 1
        if t >= self._ci_data['n_time']:
            self._ci_anim_id = None
            return
        failed = self._ci_update_frame(t)
        if failed:
            self._ci_anim_id = None
            return
        delay = max(50, int(self._ci_speed_var.get()))
        self._ci_anim_id = self.root.after(delay, self._ci_tick)

    def _ci_stop(self):
        if self._ci_anim_id is not None:
            self.root.after_cancel(self._ci_anim_id)
            self._ci_anim_id = None

    # ------------------------------------------------------------------
    # Reload parameters
    # ------------------------------------------------------------------

    def _reload_params(self):
        for widget in self.root.winfo_children():
            widget.destroy()
        self._vars.clear()
        self._frames.clear()
        self._params_list, self._categories = _load_params()
        _apply_styles(self.root)
        self._build_menu()
        self._build_ui()
        self._log('Parameters reloaded from crisis_params.py.\n', tag='info')

    # ------------------------------------------------------------------
    # Run model
    # ------------------------------------------------------------------

    def _run_model(self):
        if self._get_raw_value('analysis_type') == 'back':
            self._run_back_analysis()
            return
        if self._running:
            return
        params, errors = self._collect_params()
        if errors:
            messagebox.showerror('Validation Error', '\n'.join(errors))
            return

        # _validate_raster_sizes reads the DEM (and every other configured
        # raster) from disk — for a large regional DEM, or a slow/cloud-
        # synced drive, doing that synchronously right on the button click
        # would freeze the whole GUI before the actual run even starts. Run
        # it in a background thread and only proceed once it comes back.
        self._running = True
        self._run_btn.config(state=tk.DISABLED,
                             bg='#6b7280', activebackground='#6b7280')
        self._log('Checking input rasters against the DEM…\n', tag='dim')

        def _check_work():
            return self._validate_raster_sizes(is_ba=False)

        def _check_done(size_errors):
            if size_errors:
                self._running = False
                self._run_btn.config(state=tk.NORMAL,
                                     bg=C['success'], activebackground=C['success_dark'])
                messagebox.showerror(
                    'Raster Size Mismatch',
                    'The following inputs are incompatible with the DEM size '
                    'and cannot be used:\n\n' +
                    '\n'.join(f'  •  {e}' for e in size_errors))
                return
            self._start_forward_run(params)

        self._run_in_bg(_check_work, _check_done)

    def _start_forward_run(self, params):
        """Continuation of _run_model() once the pre-flight raster-size
        check has passed — actually starts the forward-analysis run."""
        self._ci_cell_cache = {}   # invalidate inspector cache for new run
        self._ci_dem_drawn  = False
        self._ci_dem_cache  = None
        self._ci_dem_x_cache = None
        self._ci_dem_y_cache = None

        # Switch to Running Log tab
        self._notebook.select(self._log_tab)

        self._log('─' * 64 + '\n', tag='dim')
        self._log('Starting CRISIS model run…\n', tag='info')

        # Reuse the Topography preview's Slope/Flow Direction/Flow Accumulation
        # if they were computed for this exact DEM + Cell_size — Flow
        # Direction alone can take about a minute, so this avoids paying that
        # cost twice for a single run
        raster_cache = getattr(self, '_fwd_raster_cache', None)

        def _worker():
            try:
                import forward_analysis
                importlib.reload(forward_analysis)
                forward_analysis.run_model(params, log_callback=self._log,
                                           precomputed_rasters=raster_cache)
                self._log('\n✓  Model run completed successfully.\n', tag='ok')
                self.root.after(0, self._plot_outputs)
            except Exception as exc:
                self._log(f'\n✗  ERROR: {exc}\n', tag='err')
                self._log(traceback.format_exc() + '\n', tag='err')
            finally:
                self._running = False
                self.root.after(0, lambda: self._run_btn.config(
                    state=tk.NORMAL,
                    bg=C['success'],
                    activebackground=C['success_dark']))

        threading.Thread(target=_worker, daemon=True).start()

    def _sync_inspector(self, *_):
        """Add/remove the Triggering Cell Inspector tab. Shown only in forward
        mode, rainfall-only triggering (Triggering_Indicator == 1), with more
        than one time point (nothing to inspect across time otherwise).
        Idempotent — safe to call as often as needed; also called at the end
        of _update_analysis_mode() so that its own unconditional forget of
        this tab (needed when switching away from forward mode) doesn't leave
        it hidden when it should still be shown (e.g. after a
        Triggering_Indicator change while remaining in forward mode)."""
        if self._get_raw_value('analysis_type') != 'forward':
            return
        try:
            trig = int(self._get_raw_value('Triggering_Indicator') or 1)
        except Exception:
            trig = 1
        try:
            n_tp = int(self._get_raw_value('time_points') or 1)
        except Exception:
            n_tp = 1
        should_show = (trig == 1 and n_tp > 1)
        if should_show and not self._inspector_in_notebook:
            self._notebook.add(self._inspector_frame,
                               text='  Triggering Cell Inspector  ')
            self._inspector_in_notebook = True
            def _on_tab_changed(event):
                try:
                    selected = self._notebook.tab(self._notebook.select(), 'text').strip()
                    if selected == 'Triggering Cell Inspector':
                        self._ci_load()
                except Exception:
                    pass
            self._notebook.bind('<<NotebookTabChanged>>', _on_tab_changed)
        elif not should_show and self._inspector_in_notebook:
            self._notebook.forget(self._inspector_frame)
            self._inspector_in_notebook = False

    def _update_analysis_mode(self, *_):
        if not hasattr(self, '_forward_tab_refs') or not hasattr(self, '_back_tab_refs'):
            return
        mode = self._get_raw_value('analysis_type')  # 'forward', 'back', or None

        self._run_btn.config(text='▶   Run Model')

        # Remember which tab the user was on (by title, since the frame
        # widgets themselves get forgotten and re-added below) so the
        # tab-rebuild needed to show/hide the Earthquake Inputs tab doesn't
        # silently reset the selection back to Project Details — the user
        # should stay exactly where they were (e.g. Triggering Event) and
        # choose the next tab themselves.
        try:
            prev_title = self._notebook.tab(self._notebook.select(), 'text')
        except Exception:
            prev_title = None

        # ── Step 1: strip everything back to just "Project Details" ──────
        for frame, _ in self._forward_tab_refs:
            try: self._notebook.forget(frame)
            except Exception: pass
        for frame, _ in self._back_tab_refs:
            try: self._notebook.forget(frame)
            except Exception: pass
        if hasattr(self, '_ba_outputs_tab'):
            try: self._notebook.forget(self._ba_outputs_tab)
            except Exception: pass
        if self._inspector_in_notebook:
            try: self._notebook.forget(self._inspector_frame)
            except Exception: pass
            self._inspector_in_notebook = False
        try: self._notebook.forget(self._log_tab)
        except Exception: pass

        # ── Step 2: add tabs for the chosen mode ─────────────────────────
        if mode == 'forward':
            # Earthquake Inputs is only relevant when Triggering Mechanism =
            # Rainfall + Earthquake — hide the whole tab otherwise
            show_eq = self._get_raw_value('Triggering_Indicator') == 2

            # Add all forward tabs except Outputs, then Running Log, then Outputs
            _fwd_outputs = [(f, t) for f, t in self._forward_tab_refs if t.strip() == 'Outputs']
            _fwd_other   = [(f, t) for f, t in self._forward_tab_refs
                            if t.strip() != 'Outputs'
                            and (show_eq or t.strip() != 'Earthquake Inputs')]
            for frame, title in _fwd_other:
                self._notebook.add(frame, text=title)
            self._notebook.add(self._log_tab, text='  Running Log  ')
            for frame, title in _fwd_outputs:
                self._notebook.add(frame, text=title)

        elif mode == 'back':
            for frame, title in self._back_tab_refs:
                self._notebook.add(frame, text=title)
            self._notebook.add(self._log_tab, text='  Running Log  ')
            self._notebook.add(self._ba_outputs_tab, text='  Outputs  ')

        # else None → only Project Details remains

        # Restore whichever tab the user was on before the rebuild (e.g.
        # Triggering Event), so changing Triggering_Indicator doesn't yank
        # the user elsewhere — they choose which tab to go to next. Matched
        # by title against the live notebook tabs (not the old frame
        # widgets, which were just forgotten and re-added above).
        restored = False
        if prev_title is not None:
            for tab_id in self._notebook.tabs():
                if self._notebook.tab(tab_id, 'text') == prev_title:
                    try:
                        self._notebook.select(tab_id)
                        restored = True
                    except Exception:
                        pass
                    break
        if not restored and mode == 'forward':
            # The previously active tab no longer exists in this mode (e.g.
            # it was Earthquake Inputs and the user just switched away from
            # earthquake triggering) — land on Triggering Event, since
            # that's the control that caused this rebuild.
            trig_frame = next(
                (f for f, t in self._forward_tab_refs if t.strip() == 'Triggering Event'), None)
            if trig_frame is not None:
                try:
                    self._notebook.select(trig_frame)
                except Exception:
                    pass

        # Step 1 above unconditionally forgot the Triggering Cell Inspector
        # tab (it's managed separately from _forward_tab_refs/_back_tab_refs)
        # — re-evaluate and re-add it now if it should still be shown, e.g.
        # after a Triggering_Indicator change while remaining in forward mode
        self._sync_inspector()

    def _run_back_analysis(self):
        if self._running:
            return

        all_params, errors = self._collect_params()
        if errors:
            messagebox.showerror('Validation Error', '\n'.join(errors))
            return

        # _validate_raster_sizes reads the DEM (and every other configured
        # raster) from disk — for a large regional DEM, or a slow/cloud-
        # synced drive, doing that synchronously right on the button click
        # would freeze the whole GUI before the actual run even starts. Run
        # it in a background thread and only proceed once it comes back.
        self._running = True
        self._run_btn.config(state=tk.DISABLED,
                             bg='#6b7280', activebackground='#6b7280')
        self._log('Checking input rasters against the DEM…\n', tag='dim')

        def _check_work():
            return self._validate_raster_sizes(is_ba=True)

        def _check_done(size_errors):
            if size_errors:
                self._running = False
                self._run_btn.config(state=tk.NORMAL,
                                     bg=C['success'], activebackground=C['success_dark'])
                messagebox.showerror(
                    'Raster Size Mismatch',
                    'The following inputs are incompatible with the DEM size '
                    'and cannot be used:\n\n' +
                    '\n'.join(f'  •  {e}' for e in size_errors))
                return
            self._start_back_analysis_run(all_params)

        self._run_in_bg(_check_work, _check_done)

    def _start_back_analysis_run(self, all_params):
        """Continuation of _run_back_analysis() once the pre-flight raster-
        size check has passed — actually starts the back-analysis run."""
        self._notebook.select(self._log_tab)
        self._log('─' * 64 + '\n', tag='dim')
        self._log('Starting CRISIS Back-Analysis…\n', tag='info')

        # Strip 'ba_' prefix so run_model() receives clean key names
        # (matching config_back_analysis_windows.json and the standalone CLI path)
        _BA_KEYS = {k for k in all_params if k.startswith('ba_')}
        ba_params = {k[3:] if k in _BA_KEYS else k: v for k, v in all_params.items()}

        # Reuse the Topography preview's Slope/Flow Direction/Flow Accumulation
        # if they were computed for this exact DEM + Cell_size — Flow
        # Direction alone can take about a minute, so this avoids paying that
        # cost twice for a single run
        raster_cache = getattr(self, '_ba_raster_cache', None)

        def _worker():
            _success = False
            try:
                import back_analysis
                back_analysis.run_model(ba_params, log_callback=self._log,
                                        precomputed_rasters=raster_cache)
                # Note: the "completed successfully" message is logged later,
                # by _run_ba_outputs's _done() callback below — after the
                # Excel summary and plots are actually finished, not here.
                _success = True
            except Exception as exc:
                self._log(f'\n✗  ERROR: {exc}\n', tag='err')
                self._log(traceback.format_exc() + '\n', tag='err')
            finally:
                self._running = False
                self.root.after(0, lambda: self._run_btn.config(
                    state=tk.NORMAL,
                    bg=C['success'],
                    activebackground=C['success_dark']))
                if _success:
                    self.root.after(500, self._run_ba_outputs)

        threading.Thread(target=_worker, daemon=True).start()

    # ------------------------------------------------------------------
    # Log
    # ------------------------------------------------------------------

    def _log(self, msg: str, tag: str = ''):
        # Buffer messages and flush at most ~20 times/sec to avoid flooding the
        # Tkinter event queue (which caused the GUI to appear frozen during runs).
        if not hasattr(self, '_log_buffer'):
            self._log_buffer = []
            self._log_flush_pending = False

        self._log_buffer.append((msg, tag))

        if not self._log_flush_pending:
            self._log_flush_pending = True
            self.root.after(50, self._flush_log)

    def _flush_log(self):
        if not hasattr(self, '_log_buffer') or not self._log_buffer:
            self._log_flush_pending = False
            return
        items, self._log_buffer = self._log_buffer, []
        self._log_flush_pending = False
        self._log_box.config(state=tk.NORMAL)
        for msg, tag in items:
            if tag:
                self._log_box.insert(tk.END, msg, tag)
            else:
                self._log_box.insert(tk.END, msg)
        self._log_box.see(tk.END)
        self._log_box.config(state=tk.DISABLED)




# =============================================================================
# Entry point
# =============================================================================

def _load_params():
    import crisis_params
    importlib.reload(crisis_params)
    return crisis_params.PARAMS, crisis_params.PARAM_CATEGORIES


def main():
    root = tk.Tk()
    CRISISGui(root)
    root.mainloop()


if __name__ == '__main__':
    from multiprocessing import freeze_support
    freeze_support()
    if not _IN_WORKER:
        main()
